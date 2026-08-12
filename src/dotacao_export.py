"""Exportação auditável, em memória, da análise de Dotação Orçamentária."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Mapping, Sequence

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.dotacao_analysis import (
    KNOWN_ITEM_INDICATORS,
    apply_dotacao_filters,
    build_detail_view,
    build_indicator_calculation_memory,
    build_item_indicators,
    summarize_value_states,
)
from src.tesouro_dotacao import CODE_COLUMNS


EXPORT_SHEET_NAMES = [
    "Resumo",
    "Memória de Cálculo",
    "Detalhamento",
    "Metadados",
]

FILTER_LABELS = {
    "exercicio": "Exercício",
    "iduso": "IDUSO",
    "resultado_primario": "Resultado Primário",
    "acao_governo": "Ação Governo",
    "ptres": "PTRES",
    "plano_orcamentario": "Plano Orçamentário",
    "grupo_despesa": "Grupo de Despesa",
    "fonte_recursos_detalhada": "Fonte de Recursos Detalhada",
    "periodo": "Período",
    "item_informacao": "Item de informação",
}

CURRENCY_FORMAT = '#,##0.00;[Red]-#,##0.00;0.00'
INTEGER_FORMAT = '#,##0'
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def build_dotacao_analysis_excel(
    dataframe: pd.DataFrame,
    active_filters: Mapping[str, Sequence[object]],
    validation_status: str,
    selected_indicator_code: str | None = None,
    generated_at: datetime | None = None,
) -> bytes:
    """Gera um arquivo Excel em memória para o recorte analítico atual."""

    filtered = apply_dotacao_filters(dataframe, active_filters)
    indicators = _summary_dataframe(filtered)
    memory = _memory_dataframe(filtered, selected_indicator_code)
    detail = build_detail_view(filtered)
    metadata = _metadata_dataframe(
        filtered,
        active_filters,
        validation_status,
        selected_indicator_code,
        generated_at or datetime.now(),
    )

    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        indicators.to_excel(writer, sheet_name="Resumo", index=False)
        memory.to_excel(writer, sheet_name="Memória de Cálculo", index=False)
        detail.to_excel(writer, sheet_name="Detalhamento", index=False)
        metadata.to_excel(writer, sheet_name="Metadados", index=False)

        _format_workbook(writer, indicators, memory, detail, metadata)

    return buffer.getvalue()


def _summary_dataframe(dataframe: pd.DataFrame) -> pd.DataFrame:
    summary = build_item_indicators(dataframe).rename(
        columns={
            "indicador": "Indicador",
            "item_informacao_codigo": "Item informação código",
            "quantidade_registros": "Quantidade de registros",
            "quantidade_nulos": "Quantidade de nulos",
            "quantidade_zeros": "Quantidade de zeros",
            "quantidade_negativos": "Quantidade de negativos",
            "valor_movimento_liquido": "Valor do movimento líquido",
        }
    )
    return summary


def _memory_dataframe(
    dataframe: pd.DataFrame,
    selected_indicator_code: str | None,
) -> pd.DataFrame:
    if selected_indicator_code is None:
        return pd.DataFrame(
            columns=[
                "item_informacao_origem",
                "item_informacao_codigo",
                "arquivo_origem",
                "aba_origem",
                "linha_origem",
                "coluna_origem",
                "valor_movimento_liquido",
            ]
        )

    memory = build_indicator_calculation_memory(
        dataframe,
        selected_indicator_code,
    )
    if not memory.reconciled:
        raise ValueError("A memória de cálculo não está reconciliada para exportação.")
    return memory.linhas.copy(deep=True)


def _metadata_dataframe(
    dataframe: pd.DataFrame,
    active_filters: Mapping[str, Sequence[object]],
    validation_status: str,
    selected_indicator_code: str | None,
    generated_at: datetime,
) -> pd.DataFrame:
    states = summarize_value_states(dataframe)
    files = _distinct_text(dataframe["arquivo_origem"])
    sheets = _distinct_text(dataframe["aba_origem"])
    selected_indicator = (
        f"{KNOWN_ITEM_INDICATORS[selected_indicator_code]} "
        f"[{selected_indicator_code}]"
        if selected_indicator_code in KNOWN_ITEM_INDICATORS
        else "Não aplicável"
    )
    records: list[dict[str, object]] = [
        {
            "Campo": "Status da validação da normalização",
            "Valor": validation_status,
        },
        {"Campo": "Arquivo(s) de origem", "Valor": files},
        {"Campo": "Aba(s) de origem", "Valor": sheets},
        {"Campo": "Indicador selecionado", "Valor": selected_indicator},
        {"Campo": "Data/hora de geração", "Valor": generated_at},
        {
            "Campo": "Quantidade de registros",
            "Valor": states["quantidade_registros"],
        },
        {
            "Campo": "Quantidade de valores nulos",
            "Valor": states["quantidade_nulos"],
        },
        {
            "Campo": "Quantidade de valores iguais a zero",
            "Valor": states["quantidade_zeros"],
        },
        {
            "Campo": "Quantidade de valores negativos",
            "Valor": states["quantidade_negativos"],
        },
        {
            "Campo": "Soma do recorte exportado",
            "Valor": states["valor_movimento_liquido"],
        },
        {
            "Campo": "Natureza da extração",
            "Valor": (
                "Extração derivada em memória; a fonte original não foi alterada."
            ),
        },
    ]

    active_filter_records = [
        {
            "Campo": f"Filtro ativo: {FILTER_LABELS.get(filter_name, filter_name)}",
            "Valor": _format_filter_values(values),
        }
        for filter_name, values in active_filters.items()
        if values
    ]
    if active_filter_records:
        records.extend(active_filter_records)
    else:
        records.append({"Campo": "Filtros ativos", "Valor": "Nenhum (todos os valores)"})

    return pd.DataFrame.from_records(records, columns=["Campo", "Valor"])


def _format_workbook(
    writer: pd.ExcelWriter,
    summary: pd.DataFrame,
    memory: pd.DataFrame,
    detail: pd.DataFrame,
    metadata: pd.DataFrame,
) -> None:
    dataframes = {
        "Resumo": summary,
        "Memória de Cálculo": memory,
        "Detalhamento": detail,
        "Metadados": metadata,
    }
    code_headers = {
        "Resumo": {"Item informação código"},
        "Memória de Cálculo": {"item_informacao_codigo"},
        "Detalhamento": set(CODE_COLUMNS),
        "Metadados": set(),
    }

    for sheet_name in EXPORT_SHEET_NAMES:
        worksheet = writer.sheets[sheet_name]
        dataframe = dataframes[sheet_name]
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        _style_header(worksheet)
        _format_columns(
            worksheet,
            dataframe,
            code_headers[sheet_name],
            is_metadata=sheet_name == "Metadados",
        )


def _style_header(worksheet) -> None:
    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _format_columns(
    worksheet,
    dataframe: pd.DataFrame,
    code_headers: set[str],
    is_metadata: bool,
) -> None:
    headers = list(dataframe.columns)
    for column_index, header in enumerate(headers, start=1):
        column_letter = get_column_letter(column_index)
        cells = list(worksheet.iter_cols(min_col=column_index, max_col=column_index))[
            0
        ]
        for cell in cells[1:]:
            if header in code_headers:
                cell.number_format = "@"
                if cell.value is not None:
                    cell.value = str(cell.value)
            elif header in {"valor_movimento_liquido", "Valor do movimento líquido"}:
                cell.number_format = CURRENCY_FORMAT
            elif header.startswith("Quantidade"):
                cell.number_format = INTEGER_FORMAT
            elif is_metadata and cell.row == 6 and column_index == 2:
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
            elif is_metadata and cell.row in {7, 8, 9, 10} and column_index == 2:
                cell.number_format = INTEGER_FORMAT
            elif is_metadata and cell.row == 11 and column_index == 2:
                cell.number_format = CURRENCY_FORMAT

        max_length = max(
            [len(str(header)), *[len(str(cell.value or "")) for cell in cells[1:]]]
        )
        worksheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 60)


def _distinct_text(values: pd.Series) -> str:
    return " | ".join(str(value) for value in values.dropna().drop_duplicates())


def _format_filter_values(values: Sequence[object]) -> str:
    return " | ".join(
        "(Valor nulo)" if pd.isna(value) else str(value) for value in values
    )
