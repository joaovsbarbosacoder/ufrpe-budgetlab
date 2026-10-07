"""Planos de referência da Captação: objetivos do PDI e metas do PLS (Etapa 3).

Separação de responsabilidades (como na importação, seção 6.3 do projeto de origem):

    leitor da planilha (abas PDI/PLS → linhas brutas)
      → validação (código e descrição obrigatórios, sem duplicidade)
        → carga transacional no banco (versionada)
          → busca por código ou palavra-chave

Regras:

- Códigos são IDENTIFICADORES (texto): zeros à esquerda e alfanuméricos são preservados.
  Célula numérica vira texto, com aviso, porque o Excel pode já ter perdido zeros.
- A carga é "tudo ou nada": havendo qualquer erro de validação, nada é gravado e cada
  erro é listado com a linha da planilha. Nada é descartado em silêncio.
- Recarregar uma lista é idempotente e nunca apaga: itens já existentes mantêm o `id`
  (vínculos de demandas continuam válidos) e têm a descrição atualizada; itens que
  sumiram da nova planilha são DESATIVADOS (`ativo = 0`), não excluídos. Descrições
  alteradas são devolvidas no resultado, para rastreabilidade.
- O arquivo original nunca é alterado: a leitura é em memória.
"""

from __future__ import annotations

import sqlite3
import unicodedata
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import IO, Any

import openpyxl

ABA_PDI = "PDI"
ABA_PLS = "PLS"

# campo interno → cabeçalho da planilha modelo (comparado sem asterisco, acento e caixa)
COLUNAS_PDI = {
    "codigo": "Código do objetivo",
    "dimensao": "Dimensão / eixo do PDI",
    "descricao": "Descrição do objetivo",
}
COLUNAS_PLS = {
    "eixo": "Eixo temático",
    "objetivo": "Objetivo",
    "codigo": "Código da meta",
    "descricao": "Descrição da meta",
}
OBRIGATORIAS_PDI = ("codigo", "descricao")
OBRIGATORIAS_PLS = ("codigo", "descricao")


@dataclass
class LeituraPlano:
    """Resultado de ler e validar uma aba. `erros` não vazio = a carga deve ser recusada."""

    linhas: list[dict[str, Any]] = field(default_factory=list)
    erros: list[str] = field(default_factory=list)
    avisos: list[str] = field(default_factory=list)


@dataclass
class ResultadoCarga:
    inseridos: int = 0
    atualizados: int = 0
    inalterados: int = 0
    desativados: int = 0
    reativados: int = 0
    descricoes_alteradas: list[tuple[str, str, str]] = field(default_factory=list)  # (código, antes, depois)


class ErroCarga(Exception):
    """A carga foi recusada; `erros` traz cada motivo (nada foi gravado)."""

    def __init__(self, erros: list[str]):
        self.erros = list(erros)
        super().__init__("; ".join(self.erros))


def _dobrar(texto: object) -> str:
    """Minúsculas, sem acento, sem asterisco e sem espaços repetidos — para comparar
    cabeçalhos e buscar texto."""

    base = unicodedata.normalize("NFKD", str(texto or ""))
    base = "".join(c for c in base if not unicodedata.combining(c))
    return " ".join(base.replace("*", " ").casefold().split())


def _texto_celula(valor: object) -> tuple[str, bool]:
    """(texto, veio_numerico). Inteiro vira "1"; decimal exato vira "1.1"."""

    if valor is None:
        return "", False
    if isinstance(valor, bool):
        return str(valor), False
    if isinstance(valor, int):
        return str(valor), True
    if isinstance(valor, float):
        return (str(int(valor)) if valor.is_integer() else repr(valor)), True
    return str(valor).strip(), False


def _ler_aba(
    linhas_brutas: list[tuple[Any, ...]],
    colunas: dict[str, str],
    obrigatorias: tuple[str, ...],
    rotulo_aba: str,
) -> LeituraPlano:
    leitura = LeituraPlano()
    if not linhas_brutas:
        leitura.erros.append(f"Aba {rotulo_aba}: planilha vazia.")
        return leitura

    cabecalho = [_dobrar(c) for c in linhas_brutas[0]]
    posicao: dict[str, int] = {}
    for campo, titulo in colunas.items():
        alvo = _dobrar(titulo)
        if alvo in cabecalho:
            posicao[campo] = cabecalho.index(alvo)
    faltando = [colunas[c] for c in obrigatorias if c not in posicao]
    if faltando:
        leitura.erros.append(
            f"Aba {rotulo_aba}: coluna obrigatória ausente ({', '.join(faltando)}). O cabeçalho é "
            "localizado pelo texto; não renomeie as colunas do modelo."
        )
        return leitura

    vistos: dict[str, int] = {}
    for numero, bruta in enumerate(linhas_brutas[1:], start=2):
        celulas = {campo: _texto_celula(bruta[i] if i < len(bruta) else None) for campo, i in posicao.items()}
        if not any(texto for texto, _ in celulas.values()):
            continue  # linha totalmente vazia
        registro: dict[str, Any] = {campo: texto for campo, (texto, _) in celulas.items()}
        registro["numero_linha"] = numero
        erros_linha: list[str] = []

        codigo = registro.get("codigo", "")
        if not codigo:
            erros_linha.append("código ausente")
        elif codigo.startswith("["):
            erros_linha.append(f"marcador {codigo!r} não substituído pelo dado oficial")
        elif codigo in vistos:
            erros_linha.append(f"código {codigo!r} duplicado (já aparece na linha {vistos[codigo]})")
        else:
            vistos[codigo] = numero
        if not registro.get("descricao", ""):
            erros_linha.append("descrição ausente")
        elif registro["descricao"].startswith("[") and not codigo.startswith("["):
            erros_linha.append("descrição é um marcador não substituído pelo dado oficial")
        if celulas.get("codigo", ("", False))[1]:
            leitura.avisos.append(
                f"Aba {rotulo_aba}, linha {numero}: código {codigo!r} veio como número; se o código "
                "tem zeros à esquerda ou casas decimais finais, formate a coluna como texto."
            )

        if erros_linha:
            leitura.erros.append(f"Aba {rotulo_aba}, linha {numero}: " + "; ".join(erros_linha) + ".")
        else:
            leitura.linhas.append(registro)

    if not leitura.linhas and not leitura.erros:
        leitura.erros.append(f"Aba {rotulo_aba}: nenhuma linha preenchida.")
    return leitura


def validar_linhas_pdi(linhas_brutas: list[tuple[Any, ...]]) -> LeituraPlano:
    """Valida as linhas brutas (primeira linha = cabeçalho) da aba PDI."""

    return _ler_aba(linhas_brutas, COLUNAS_PDI, OBRIGATORIAS_PDI, ABA_PDI)


def validar_linhas_pls(linhas_brutas: list[tuple[Any, ...]]) -> LeituraPlano:
    """Valida as linhas brutas (primeira linha = cabeçalho) da aba PLS."""

    return _ler_aba(linhas_brutas, COLUNAS_PLS, OBRIGATORIAS_PLS, ABA_PLS)


def ler_planilha(arquivo: str | Path | IO[bytes]) -> tuple[LeituraPlano, LeituraPlano]:
    """Lê as abas PDI e PLS de uma planilha .xlsx (caminho ou arquivo aberto), só em
    memória. Aba ausente vira erro, não exceção."""

    workbook = openpyxl.load_workbook(arquivo, read_only=True, data_only=True)
    try:
        resultados = []
        for aba, validar in ((ABA_PDI, validar_linhas_pdi), (ABA_PLS, validar_linhas_pls)):
            if aba not in workbook.sheetnames:
                leitura = LeituraPlano(erros=[f"A planilha não tem a aba {aba!r}."])
            else:
                leitura = validar(list(workbook[aba].iter_rows(values_only=True)))
            resultados.append(leitura)
    finally:
        workbook.close()
    return resultados[0], resultados[1]


# --- Trilha de auditoria ----------------------------------------------------------------

#: tipo lógico → tabela, coluna de versão e campos editáveis (além do código)
TIPOS = {
    "pdi": {"tabela": "objetivo_pdi", "versao": "versao_pdi", "campos": ("dimensao", "descricao")},
    "pls": {"tabela": "meta_pls", "versao": "versao_pls", "campos": ("eixo", "objetivo", "descricao")},
}


def _agora() -> str:
    return datetime.now(timezone.utc).isoformat()


def _registrar(
    conexao: sqlite3.Connection,
    tabela: str,
    item_id: int,
    acao: str,
    campo: str | None,
    antes: object,
    depois: object,
    usuario: str | None,
) -> None:
    """Uma linha de `historico_plano`. Não faz commit (quem chama controla a transação)."""

    conexao.execute(
        "INSERT INTO historico_plano (tabela, item_id, acao, campo, antes, depois, usuario, criado_em)"
        " VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
        (tabela, item_id, acao, campo, antes, depois, usuario, _agora()),
    )


# --- Carga ----------------------------------------------------------------------------


def _carregar(
    conexao: sqlite3.Connection,
    leitura: LeituraPlano,
    tabela: str,
    coluna_versao: str,
    versao: str,
    campos: tuple[str, ...],
    usuario: str | None,
) -> ResultadoCarga:
    if not str(versao or "").strip():
        raise ErroCarga(["Informe a versão da lista (ex.: PDI 2025–2030)."])
    if leitura.erros:
        raise ErroCarga(leitura.erros)
    versao = versao.strip()

    resultado = ResultadoCarga()
    existentes = {
        linha["codigo"]: linha
        for linha in conexao.execute(f"SELECT * FROM {tabela} WHERE {coluna_versao} = ?", (versao,))
    }
    novos = {linha["codigo"] for linha in leitura.linhas}

    with conexao:
        for linha in leitura.linhas:
            codigo = linha["codigo"]
            valores = {campo: (linha.get(campo) or None) for campo in campos}
            atual = existentes.get(codigo)
            if atual is None:
                colunas = ", ".join(("codigo", *campos, coluna_versao))
                marcas = ", ".join("?" * (len(campos) + 2))
                cursor = conexao.execute(
                    f"INSERT INTO {tabela} ({colunas}) VALUES ({marcas})",
                    (codigo, *[valores[c] for c in campos], versao),
                )
                _registrar(conexao, tabela, cursor.lastrowid, "CRIAR", "codigo", None, codigo, usuario)
                resultado.inseridos += 1
                continue
            mudancas = [(c, atual[c], valores[c]) for c in campos if atual[c] != valores[c]]
            if not mudancas and atual["ativo"]:
                resultado.inalterados += 1
                continue
            # A recarga sobrescreve edições manuais, mas o valor anterior fica no histórico.
            for campo, antes, depois in mudancas:
                _registrar(conexao, tabela, atual["id"], "CARGA", campo, antes, depois, usuario)
            if any(campo == "descricao" for campo, _, _ in mudancas):
                resultado.descricoes_alteradas.append((codigo, atual["descricao"], valores["descricao"]))
            if not atual["ativo"]:
                _registrar(conexao, tabela, atual["id"], "REATIVAR", "ativo", "0", "1", usuario)
                resultado.reativados += 1
            atribuicoes = ", ".join(f"{c} = ?" for c in campos)
            conexao.execute(
                f"UPDATE {tabela} SET {atribuicoes}, ativo = 1 WHERE id = ?",
                (*[valores[c] for c in campos], atual["id"]),
            )
            resultado.atualizados += 1
        for codigo, atual in existentes.items():
            if codigo not in novos and atual["ativo"]:
                conexao.execute(f"UPDATE {tabela} SET ativo = 0 WHERE id = ?", (atual["id"],))
                _registrar(conexao, tabela, atual["id"], "DESATIVAR", "ativo", "1", "0", usuario)
                resultado.desativados += 1
    return resultado


def carregar_pdi(
    conexao: sqlite3.Connection, leitura: LeituraPlano, versao_pdi: str, usuario: str | None = None
) -> ResultadoCarga:
    """Grava os objetivos do PDI da `versao_pdi` (idempotente; nunca exclui)."""

    return _carregar(conexao, leitura, "objetivo_pdi", "versao_pdi", versao_pdi, ("dimensao", "descricao"), usuario)


def carregar_pls(
    conexao: sqlite3.Connection, leitura: LeituraPlano, versao_pls: str, usuario: str | None = None
) -> ResultadoCarga:
    """Grava as metas do PLS da `versao_pls` (idempotente; nunca exclui)."""

    return _carregar(
        conexao, leitura, "meta_pls", "versao_pls", versao_pls, ("eixo", "objetivo", "descricao"), usuario
    )


# --- Edição manual --------------------------------------------------------------------


def _tipo(tipo: str) -> dict[str, Any]:
    if tipo not in TIPOS:
        raise ValueError(f"Tipo de plano desconhecido: {tipo!r} (use 'pdi' ou 'pls').")
    return TIPOS[tipo]


def _limpar(dados: dict[str, Any], campos: tuple[str, ...]) -> dict[str, str | None]:
    """Texto aparado; vazio vira None (campo opcional ausente, não string vazia)."""

    return {c: (str(dados.get(c) or "").strip() or None) for c in ("codigo", *campos)}


def _validar_edicao(
    conexao: sqlite3.Connection,
    config: dict[str, Any],
    valores: dict[str, str | None],
    versao: str,
    ignorar_id: int | None,
) -> None:
    erros = []
    codigo = valores["codigo"]
    if not codigo:
        erros.append("Informe o código.")
    elif codigo.startswith("["):
        erros.append(f"O código {codigo!r} parece um marcador não substituído.")
    else:
        repetido = conexao.execute(
            f"SELECT id FROM {config['tabela']} WHERE {config['versao']} = ? AND codigo = ? AND id IS NOT ?",
            (versao, codigo, ignorar_id),
        ).fetchone()
        if repetido:
            erros.append(f"Já existe o código {codigo!r} nesta versão ({versao}).")
    if not valores["descricao"]:
        erros.append("Informe a descrição.")
    if erros:
        raise ErroCarga(erros)


def _exigir_usuario(usuario: str | None) -> str:
    usuario = (usuario or "").strip()
    if not usuario:
        raise ErroCarga(["Informe quem está fazendo a alteração (fica no histórico)."])
    return usuario


def criar_item(
    conexao: sqlite3.Connection, tipo: str, versao: str, dados: dict[str, Any], usuario: str | None
) -> int:
    """Inclui um objetivo do PDI (`tipo='pdi'`) ou uma meta do PLS (`tipo='pls'`) à mão."""

    config = _tipo(tipo)
    usuario = _exigir_usuario(usuario)
    versao = str(versao or "").strip()
    if not versao:
        raise ErroCarga(["Informe a versão da lista."])
    valores = _limpar(dados, config["campos"])
    _validar_edicao(conexao, config, valores, versao, None)
    with conexao:
        colunas = ", ".join(("codigo", *config["campos"], config["versao"]))
        marcas = ", ".join("?" * (len(config["campos"]) + 2))
        cursor = conexao.execute(
            f"INSERT INTO {config['tabela']} ({colunas}) VALUES ({marcas})",
            (valores["codigo"], *[valores[c] for c in config["campos"]], versao),
        )
        item_id = int(cursor.lastrowid)
        _registrar(conexao, config["tabela"], item_id, "CRIAR", "codigo", None, valores["codigo"], usuario)
    return item_id


def editar_item(
    conexao: sqlite3.Connection, tipo: str, item_id: int, dados: dict[str, Any], usuario: str | None
) -> list[str]:
    """Altera código e campos de texto de um item (o `id` não muda, então vínculos de
    demandas continuam válidos). Devolve os campos alterados; sem mudança, não grava nada."""

    config = _tipo(tipo)
    usuario = _exigir_usuario(usuario)
    atual = conexao.execute(f"SELECT * FROM {config['tabela']} WHERE id = ?", (item_id,)).fetchone()
    if atual is None:
        raise ErroCarga([f"Item {item_id} não encontrado."])
    valores = _limpar(dados, config["campos"])
    _validar_edicao(conexao, config, valores, atual[config["versao"]], item_id)

    mudancas = [(c, atual[c], valores[c]) for c in ("codigo", *config["campos"]) if atual[c] != valores[c]]
    if not mudancas:
        return []
    with conexao:
        atribuicoes = ", ".join(f"{c} = ?" for c, _, _ in mudancas)
        conexao.execute(
            f"UPDATE {config['tabela']} SET {atribuicoes} WHERE id = ?",
            (*[depois for _, _, depois in mudancas], item_id),
        )
        for campo, antes, depois in mudancas:
            _registrar(conexao, config["tabela"], item_id, "EDITAR", campo, antes, depois, usuario)
    return [campo for campo, _, _ in mudancas]


def alterar_ativo(
    conexao: sqlite3.Connection, tipo: str, item_id: int, ativo: bool, usuario: str | None
) -> bool:
    """Desativa ou reativa (nunca exclui). Devolve se houve mudança."""

    config = _tipo(tipo)
    usuario = _exigir_usuario(usuario)
    atual = conexao.execute(f"SELECT ativo FROM {config['tabela']} WHERE id = ?", (item_id,)).fetchone()
    if atual is None:
        raise ErroCarga([f"Item {item_id} não encontrado."])
    if bool(atual["ativo"]) == ativo:
        return False
    with conexao:
        conexao.execute(f"UPDATE {config['tabela']} SET ativo = ? WHERE id = ?", (int(ativo), item_id))
        _registrar(
            conexao, config["tabela"], item_id, "REATIVAR" if ativo else "DESATIVAR",
            "ativo", str(int(not ativo)), str(int(ativo)), usuario,
        )
    return True


def historico_item(conexao: sqlite3.Connection, tipo: str, item_id: int) -> list[sqlite3.Row]:
    """Histórico do item, do mais antigo ao mais recente."""

    config = _tipo(tipo)
    return conexao.execute(
        "SELECT * FROM historico_plano WHERE tabela = ? AND item_id = ? ORDER BY id",
        (config["tabela"], item_id),
    ).fetchall()


# --- Busca ----------------------------------------------------------------------------


def _filtrar(
    linhas: list[sqlite3.Row], texto: str, campos_texto: tuple[str, ...]
) -> list[sqlite3.Row]:
    """Todas as palavras do texto precisam aparecer (sem acento/caixa) no código ou nos
    campos de texto. Código idêntico ao termo vem primeiro, depois código que começa com
    o termo, depois o resto na ordem do código."""

    termo = _dobrar(texto)
    if not termo:
        return list(linhas)
    palavras = termo.split()
    achadas = []
    for linha in linhas:
        alvo = _dobrar(" ".join(str(linha[c] or "") for c in ("codigo", *campos_texto)))
        if all(p in alvo for p in palavras):
            achadas.append(linha)

    def chave(linha: sqlite3.Row) -> tuple[int, str]:
        codigo = _dobrar(linha["codigo"])
        return (0 if codigo == termo else 1 if codigo.startswith(termo) else 2, str(linha["codigo"]))

    return sorted(achadas, key=chave)


def buscar_objetivos_pdi(
    conexao: sqlite3.Connection,
    texto: str = "",
    versao_pdi: str | None = None,
    apenas_ativos: bool = True,
) -> list[sqlite3.Row]:
    """Objetivos do PDI cujo código ou texto (dimensão/descrição) contém todas as palavras."""

    consulta, parametros = "SELECT * FROM objetivo_pdi WHERE 1 = 1", []
    if apenas_ativos:
        consulta += " AND ativo = 1"
    if versao_pdi:
        consulta += " AND versao_pdi = ?"
        parametros.append(versao_pdi)
    return _filtrar(conexao.execute(consulta, parametros).fetchall(), texto, ("dimensao", "descricao"))


def buscar_metas_pls(
    conexao: sqlite3.Connection,
    texto: str = "",
    versao_pls: str | None = None,
    eixo: str | None = None,
    apenas_ativas: bool = True,
) -> list[sqlite3.Row]:
    """Metas do PLS cujo código ou texto (eixo/objetivo/descrição) contém todas as palavras;
    `eixo` filtra por eixo exato (como o filtro por eixo da tela de nova demanda)."""

    consulta, parametros = "SELECT * FROM meta_pls WHERE 1 = 1", []
    if apenas_ativas:
        consulta += " AND ativo = 1"
    if versao_pls:
        consulta += " AND versao_pls = ?"
        parametros.append(versao_pls)
    if eixo:
        consulta += " AND eixo = ?"
        parametros.append(eixo)
    return _filtrar(conexao.execute(consulta, parametros).fetchall(), texto, ("eixo", "objetivo", "descricao"))
