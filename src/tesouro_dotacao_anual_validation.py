"""Validação de integridade da normalização de Dotação Anual (BI CPOC - Por Ano)."""

from __future__ import annotations

from dataclasses import dataclass

import pandas as pd
from openpyxl.worksheet.worksheet import Worksheet

from src.tesouro_dotacao_anual import FIRST_DATA_ROW, detect_dotacao_anual_structure


# A Dotação mensal usa tolerância de 0.000001 porque suas somas ficam na casa
# dos milhões. A Dotação Anual soma dezenas de milhares de células na casa dos
# bilhões: nessa magnitude, o próprio float64 acumula ruído de representação
# de ~0.000002 mesmo sem nenhuma divergência real de dado (verificado com a
# base real: contagens de nulos/zeros/negativos idênticas, soma por item e
# por ano idênticas). Por isso a tolerância aqui é em centavos, não em
# milionésimos — abaixo de um centavo não é uma diferença monetária real.
DEFAULT_TOLERANCE = 0.01
NULL_GROUP_LABEL = "<NULO>"

INCONSISTENCY_COLUMNS = [
    "tipo",
    "ano_lancamento",
    "item_informacao",
    "linha_origem",
    "coluna_origem",
    "valor_matriz_original",
    "valor_base_normalizada",
    "diferenca",
    "mensagem",
]


@dataclass
class DotacaoAnualNormalizationValidation:
    """Métricas e comparações de uma aba de Dotação Anual normalizada."""

    sheet_name: str
    tolerance: float
    raw_data_row_count: int
    monetary_column_count: int
    monetary_cell_count: int
    normalized_row_count: int
    raw_null_count: int
    normalized_null_count: int
    raw_zero_count: int
    normalized_zero_count: int
    raw_negative_count: int
    normalized_negative_count: int
    raw_sum: float
    normalized_sum: float
    difference: float
    ano_comparison: pd.DataFrame
    item_comparison: pd.DataFrame
    ano_item_comparison: pd.DataFrame
    inconsistencies: pd.DataFrame

    @property
    def approved(self) -> bool:
        return self.inconsistencies.empty


def validate_dotacao_anual_normalization(
    worksheet: Worksheet,
    normalized_dataframe: pd.DataFrame,
    sheet_name: str,
    tolerance: float = DEFAULT_TOLERANCE,
) -> DotacaoAnualNormalizationValidation:
    """Compara a matriz monetária bruta e sua representação em formato longo."""

    if tolerance < 0:
        raise ValueError("A tolerância não pode ser negativa.")

    structure, failures = detect_dotacao_anual_structure(worksheet)
    if structure is None:
        raise ValueError(
            "A validação só pode ser executada em uma aba reconhecida como "
            "Dotação Anual: " + "; ".join(failures)
        )

    required_columns = {
        "aba_origem",
        "linha_origem",
        "coluna_origem",
        "ano_lancamento",
        "item_informacao_origem",
        "valor_movimento_liquido",
    }
    missing_columns = sorted(required_columns.difference(normalized_dataframe.columns))
    if missing_columns:
        raise ValueError(
            "A base normalizada não contém as colunas necessárias: "
            + ", ".join(missing_columns)
        )

    normalized = normalized_dataframe[
        normalized_dataframe["aba_origem"] == sheet_name
    ].copy(deep=True)

    raw_long = _build_raw_monetary_long(worksheet, structure.monetary_columns)
    normalized_long = _build_normalized_monetary_long(normalized)

    raw_data_row_count = max(worksheet.max_row - FIRST_DATA_ROW + 1, 0)
    monetary_column_count = len(structure.monetary_columns)
    monetary_cell_count = raw_data_row_count * monetary_column_count
    normalized_row_count = len(normalized_long)

    raw_values = raw_long["valor"]
    normalized_values = normalized_long["valor"]

    raw_null_count = int(raw_values.isna().sum())
    normalized_null_count = int(normalized_values.isna().sum())
    raw_zero_count = int(raw_values.eq(0).fillna(False).sum())
    normalized_zero_count = int(normalized_values.eq(0).fillna(False).sum())
    raw_negative_count = int(raw_values.lt(0).fillna(False).sum())
    normalized_negative_count = int(normalized_values.lt(0).fillna(False).sum())

    raw_sum = float(raw_values.sum(skipna=True))
    normalized_sum = float(normalized_values.sum(skipna=True))
    difference = normalized_sum - raw_sum

    ano_comparison = _compare_aggregates(raw_long, normalized_long, ["ano_lancamento"], tolerance)
    item_comparison = _compare_aggregates(
        raw_long, normalized_long, ["item_informacao_origem"], tolerance
    )
    ano_item_comparison = _compare_aggregates(
        raw_long, normalized_long, ["ano_lancamento", "item_informacao_origem"], tolerance
    )

    inconsistency_records: list[dict[str, object]] = []
    _append_count_inconsistency(
        inconsistency_records,
        "quantidade_linhas_normalizadas",
        monetary_cell_count,
        normalized_row_count,
        "A quantidade de linhas normalizadas difere de linhas brutas × colunas "
        "monetárias.",
    )
    _append_count_inconsistency(
        inconsistency_records,
        "quantidade_nulos",
        raw_null_count,
        normalized_null_count,
        "A quantidade de valores nulos não foi preservada.",
    )
    _append_count_inconsistency(
        inconsistency_records,
        "quantidade_zeros",
        raw_zero_count,
        normalized_zero_count,
        "A quantidade de valores iguais a zero não foi preservada.",
    )
    _append_count_inconsistency(
        inconsistency_records,
        "quantidade_negativos",
        raw_negative_count,
        normalized_negative_count,
        "A quantidade de valores negativos não foi preservada.",
    )

    if not _values_equal(raw_sum, normalized_sum, tolerance):
        inconsistency_records.append(
            _inconsistency_record(
                inconsistency_type="soma_total",
                original_value=raw_sum,
                normalized_value=normalized_sum,
                difference=difference,
                message="A soma total difere acima da tolerância.",
            )
        )

    inconsistency_records.extend(
        _coordinate_inconsistencies(raw_long, normalized_long, tolerance)
    )
    inconsistency_records.extend(_aggregate_inconsistencies(ano_comparison, "soma_por_ano"))
    inconsistency_records.extend(_aggregate_inconsistencies(item_comparison, "soma_por_item"))
    inconsistency_records.extend(
        _aggregate_inconsistencies(ano_item_comparison, "soma_por_ano_item")
    )

    inconsistencies = pd.DataFrame.from_records(
        inconsistency_records, columns=INCONSISTENCY_COLUMNS
    )

    return DotacaoAnualNormalizationValidation(
        sheet_name=sheet_name,
        tolerance=tolerance,
        raw_data_row_count=raw_data_row_count,
        monetary_column_count=monetary_column_count,
        monetary_cell_count=monetary_cell_count,
        normalized_row_count=normalized_row_count,
        raw_null_count=raw_null_count,
        normalized_null_count=normalized_null_count,
        raw_zero_count=raw_zero_count,
        normalized_zero_count=normalized_zero_count,
        raw_negative_count=raw_negative_count,
        normalized_negative_count=normalized_negative_count,
        raw_sum=raw_sum,
        normalized_sum=normalized_sum,
        difference=difference,
        ano_comparison=ano_comparison,
        item_comparison=item_comparison,
        ano_item_comparison=ano_item_comparison,
        inconsistencies=inconsistencies,
    )


def _build_raw_monetary_long(worksheet: Worksheet, monetary_columns) -> pd.DataFrame:
    records: list[dict[str, object]] = []
    for column in monetary_columns:
        for row in range(FIRST_DATA_ROW, worksheet.max_row + 1):
            raw_value = worksheet.cell(row=row, column=column.column_index).value
            records.append(
                {
                    "linha_origem": row,
                    "coluna_origem": column.column_letter,
                    "ano_lancamento": _group_key(column.year),
                    "item_informacao_origem": _group_key(column.item_label),
                    "valor": _as_number(raw_value),
                }
            )
    dataframe = pd.DataFrame.from_records(
        records,
        columns=[
            "linha_origem",
            "coluna_origem",
            "ano_lancamento",
            "item_informacao_origem",
            "valor",
        ],
    )
    dataframe["valor"] = pd.array(dataframe["valor"], dtype="Float64")
    return dataframe


def _build_normalized_monetary_long(normalized: pd.DataFrame) -> pd.DataFrame:
    dataframe = normalized[
        [
            "linha_origem",
            "coluna_origem",
            "ano_lancamento",
            "item_informacao_origem",
            "valor_movimento_liquido",
        ]
    ].copy(deep=True)
    dataframe["ano_lancamento"] = dataframe["ano_lancamento"].map(_group_key)
    dataframe["item_informacao_origem"] = dataframe["item_informacao_origem"].map(_group_key)
    dataframe["valor"] = pd.array(
        pd.to_numeric(dataframe.pop("valor_movimento_liquido"), errors="coerce"),
        dtype="Float64",
    )
    return dataframe


def _compare_aggregates(
    raw_long: pd.DataFrame,
    normalized_long: pd.DataFrame,
    group_columns: list[str],
    tolerance: float,
) -> pd.DataFrame:
    raw_grouped = (
        raw_long.groupby(group_columns, dropna=False, sort=False)["valor"]
        .sum(min_count=1)
        .reset_index(name="soma_matriz_original")
    )
    normalized_grouped = (
        normalized_long.groupby(group_columns, dropna=False, sort=False)["valor"]
        .sum(min_count=1)
        .reset_index(name="soma_base_normalizada")
    )
    comparison = raw_grouped.merge(
        normalized_grouped, on=group_columns, how="outer", indicator=True
    )

    differences: list[object] = []
    consistent_values: list[bool] = []
    for _, row in comparison.iterrows():
        original_value = row["soma_matriz_original"]
        normalized_value = row["soma_base_normalizada"]
        both_present = row["_merge"] == "both"
        consistent = both_present and _values_equal(original_value, normalized_value, tolerance)
        consistent_values.append(consistent)
        differences.append(_numeric_difference(original_value, normalized_value))

    comparison["diferenca"] = pd.array(differences, dtype="Float64")
    comparison["consistente"] = consistent_values
    comparison = comparison.drop(columns="_merge")
    for column in group_columns:
        comparison[column] = comparison[column].replace(NULL_GROUP_LABEL, pd.NA)
    return comparison


def _coordinate_inconsistencies(
    raw_long: pd.DataFrame,
    normalized_long: pd.DataFrame,
    tolerance: float,
) -> list[dict[str, object]]:
    records: list[dict[str, object]] = []
    coordinate_columns = ["linha_origem", "coluna_origem"]

    duplicate_coordinates = normalized_long.duplicated(subset=coordinate_columns, keep=False)
    if duplicate_coordinates.any():
        records.append(
            _inconsistency_record(
                inconsistency_type="coordenadas_duplicadas",
                original_value=0,
                normalized_value=int(duplicate_coordinates.sum()),
                difference=int(duplicate_coordinates.sum()),
                message="A base normalizada possui coordenadas de origem duplicadas.",
            )
        )

    normalized_unique = normalized_long.drop_duplicates(subset=coordinate_columns, keep="first")
    comparison = raw_long.merge(
        normalized_unique,
        on=coordinate_columns,
        how="outer",
        suffixes=("_original", "_normalizada"),
        indicator=True,
    )
    for _, row in comparison.iterrows():
        original_value = row.get("valor_original", pd.NA)
        normalized_value = row.get("valor_normalizada", pd.NA)
        if row["_merge"] == "both" and _values_equal(original_value, normalized_value, tolerance):
            continue

        ano = row.get("ano_lancamento_original", pd.NA)
        item = row.get("item_informacao_origem_original", pd.NA)
        records.append(
            _inconsistency_record(
                inconsistency_type="valor_por_coordenada",
                ano=_display_group_value(ano),
                item=_display_group_value(item),
                source_row=row.get("linha_origem", pd.NA),
                source_column=row.get("coluna_origem", pd.NA),
                original_value=original_value,
                normalized_value=normalized_value,
                difference=_numeric_difference(original_value, normalized_value),
                message="O valor monetário difere na coordenada de origem.",
            )
        )
    return records


def _aggregate_inconsistencies(
    comparison: pd.DataFrame, inconsistency_type: str
) -> list[dict[str, object]]:
    records: list[dict[str, object]] = []
    for _, row in comparison[~comparison["consistente"]].iterrows():
        records.append(
            _inconsistency_record(
                inconsistency_type=inconsistency_type,
                ano=row.get("ano_lancamento", pd.NA),
                item=row.get("item_informacao_origem", pd.NA),
                original_value=row["soma_matriz_original"],
                normalized_value=row["soma_base_normalizada"],
                difference=row["diferenca"],
                message="A soma agregada difere acima da tolerância.",
            )
        )
    return records


def _append_count_inconsistency(
    records: list[dict[str, object]],
    inconsistency_type: str,
    expected: int,
    observed: int,
    message: str,
) -> None:
    if expected == observed:
        return
    records.append(
        _inconsistency_record(
            inconsistency_type=inconsistency_type,
            original_value=expected,
            normalized_value=observed,
            difference=observed - expected,
            message=message,
        )
    )


def _inconsistency_record(
    inconsistency_type: str,
    original_value: object,
    normalized_value: object,
    difference: object,
    message: str,
    ano: object = pd.NA,
    item: object = pd.NA,
    source_row: object = pd.NA,
    source_column: object = pd.NA,
) -> dict[str, object]:
    return {
        "tipo": inconsistency_type,
        "ano_lancamento": ano,
        "item_informacao": item,
        "linha_origem": source_row,
        "coluna_origem": source_column,
        "valor_matriz_original": original_value,
        "valor_base_normalizada": normalized_value,
        "diferenca": difference,
        "mensagem": message,
    }


def _values_equal(first: object, second: object, tolerance: float) -> bool:
    first_is_null = pd.isna(first)
    second_is_null = pd.isna(second)
    if first_is_null and second_is_null:
        return True
    if first_is_null or second_is_null:
        return False
    return abs(float(second) - float(first)) <= tolerance


def _numeric_difference(original: object, normalized: object) -> object:
    if pd.isna(original) and pd.isna(normalized):
        return 0.0
    if pd.isna(original) or pd.isna(normalized):
        return pd.NA
    return float(normalized) - float(original)


def _group_key(value: object) -> object:
    if value is None or pd.isna(value):
        return NULL_GROUP_LABEL
    text = str(value).strip()
    return text if text else NULL_GROUP_LABEL


def _display_group_value(value: object) -> object:
    if pd.isna(value) or value == NULL_GROUP_LABEL:
        return pd.NA
    return value


def _as_number(value: object) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text:
        return None
    try:
        return float(text.replace(",", "."))
    except ValueError:
        return None
