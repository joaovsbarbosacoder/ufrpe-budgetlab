"""Processamento multiaba de bases de Dotação Anual (BI CPOC - Por Ano)."""

from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from src.tesouro_dotacao_anual import (
    DOTACAO_ANUAL_OUTPUT_COLUMNS,
    read_dotacao_anual_sheet,
    recognize_dotacao_anual_base,
    transform_dotacao_anual_sheet,
)
from src.tesouro_dotacao_anual_validation import (
    DEFAULT_TOLERANCE,
    INCONSISTENCY_COLUMNS,
    validate_dotacao_anual_normalization,
)


TRACEABILITY_COLUMNS = [
    "arquivo_origem",
    "aba_origem",
    "linha_origem",
    "coluna_origem",
]

CONSOLIDATED_INCONSISTENCY_COLUMNS = ["aba_origem", *INCONSISTENCY_COLUMNS]


@dataclass
class DotacaoAnualWorkbookSheetResult:
    """Resumo rastreável do processamento de uma aba do arquivo."""

    sheet_name: str
    recognized: bool
    integrity_approved: bool | None = None
    raw_data_row_count: int = 0
    monetary_column_count: int = 0
    monetary_cell_count: int = 0
    normalized_row_count: int = 0
    null_count: int = 0
    zero_count: int = 0
    negative_count: int = 0
    raw_sum: float = 0.0
    normalized_sum: float = 0.0
    difference: float = 0.0
    anos_encontrados: list[int] = field(default_factory=list)
    itens_encontrados: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    validation_failures: list[str] = field(default_factory=list)
    structural_failures: list[str] = field(default_factory=list)
    processing_error: str | None = None
    inconsistencies: pd.DataFrame = field(
        default_factory=lambda: pd.DataFrame(columns=INCONSISTENCY_COLUMNS)
    )

    @property
    def status_label(self) -> str:
        if self.processing_error:
            return "Falha no processamento"
        if not self.recognized:
            return "Não reconhecida"
        return "Aprovada" if self.integrity_approved else "Com inconsistências"


@dataclass
class DotacaoAnualWorkbookResult:
    """Resultado consolidado, em memória, das abas de Dotação Anual reconhecidas."""

    filename: str
    sheet_results: list[DotacaoAnualWorkbookSheetResult]
    consolidated_data: pd.DataFrame
    consolidated_inconsistencies: pd.DataFrame
    consolidated_duplicate_count: int = 0

    @property
    def recognized_sheets(self) -> list[str]:
        return [result.sheet_name for result in self.sheet_results if result.recognized]

    @property
    def unrecognized_sheets(self) -> list[str]:
        return [
            result.sheet_name
            for result in self.sheet_results
            if not result.recognized and result.processing_error is None
        ]

    @property
    def processing_errors(self) -> list[str]:
        return [
            f"{result.sheet_name}: {result.processing_error}"
            for result in self.sheet_results
            if result.processing_error
        ]

    @property
    def integrity_approved(self) -> bool:
        recognized = [result for result in self.sheet_results if result.recognized]
        return (
            bool(recognized)
            and not self.processing_errors
            and all(
                result.processing_error is None and result.integrity_approved
                for result in recognized
            )
            and self.consolidated_inconsistencies.empty
        )

    @property
    def raw_sum(self) -> float:
        return sum(result.raw_sum for result in self._validated_results)

    @property
    def normalized_sum(self) -> float:
        return sum(result.normalized_sum for result in self._validated_results)

    @property
    def difference(self) -> float:
        return self.normalized_sum - self.raw_sum

    @property
    def null_count(self) -> int:
        return sum(result.null_count for result in self._validated_results)

    @property
    def zero_count(self) -> int:
        return sum(result.zero_count for result in self._validated_results)

    @property
    def negative_count(self) -> int:
        return sum(result.negative_count for result in self._validated_results)

    @property
    def normalized_row_count(self) -> int:
        return sum(result.normalized_row_count for result in self._validated_results)

    @property
    def _validated_results(self) -> list[DotacaoAnualWorkbookSheetResult]:
        return [
            result
            for result in self.sheet_results
            if result.recognized and result.processing_error is None
        ]


def process_dotacao_anual_workbook(
    file_content: bytes,
    filename: str,
    tolerance: float = DEFAULT_TOLERANCE,
) -> DotacaoAnualWorkbookResult:
    """Normaliza e valida todas as abas estruturalmente reconhecidas."""

    sheet_results: list[DotacaoAnualWorkbookSheetResult] = []
    normalized_frames: list[pd.DataFrame] = []
    consolidated_inconsistency_frames: list[pd.DataFrame] = []

    try:
        sheet_names = load_workbook(BytesIO(file_content), read_only=True).sheetnames
    except Exception as error:
        return DotacaoAnualWorkbookResult(
            filename=filename,
            sheet_results=[
                DotacaoAnualWorkbookSheetResult(
                    sheet_name="",
                    recognized=False,
                    processing_error=_format_error(error),
                )
            ],
            consolidated_data=_empty_output_dataframe(),
            consolidated_inconsistencies=pd.DataFrame(
                columns=CONSOLIDATED_INCONSISTENCY_COLUMNS
            ),
        )

    for sheet_name in sheet_names:
        try:
            worksheet = read_dotacao_anual_sheet(file_content, filename, sheet_name)
        except Exception as error:
            sheet_results.append(
                DotacaoAnualWorkbookSheetResult(
                    sheet_name=sheet_name,
                    recognized=False,
                    processing_error=_format_error(error),
                )
            )
            continue

        recognition = recognize_dotacao_anual_base(worksheet)
        if not recognition.matched:
            sheet_results.append(
                DotacaoAnualWorkbookSheetResult(
                    sheet_name=sheet_name,
                    recognized=False,
                    structural_failures=recognition.validation_failures,
                )
            )
            continue

        try:
            treatment = transform_dotacao_anual_sheet(worksheet, filename, sheet_name)
            validation = validate_dotacao_anual_normalization(
                worksheet,
                treatment.normalized_data,
                sheet_name,
                tolerance=tolerance,
            )
        except Exception as error:
            sheet_results.append(
                DotacaoAnualWorkbookSheetResult(
                    sheet_name=sheet_name,
                    recognized=True,
                    processing_error=_format_error(error),
                )
            )
            continue

        normalized_frames.append(treatment.normalized_data)
        if not validation.inconsistencies.empty:
            sheet_inconsistencies = validation.inconsistencies.copy(deep=True)
            sheet_inconsistencies.insert(0, "aba_origem", sheet_name)
            consolidated_inconsistency_frames.append(sheet_inconsistencies)

        sheet_results.append(
            DotacaoAnualWorkbookSheetResult(
                sheet_name=sheet_name,
                recognized=True,
                integrity_approved=(
                    validation.approved and not treatment.validation_failures
                ),
                raw_data_row_count=validation.raw_data_row_count,
                monetary_column_count=validation.monetary_column_count,
                monetary_cell_count=validation.monetary_cell_count,
                normalized_row_count=validation.normalized_row_count,
                null_count=validation.normalized_null_count,
                zero_count=validation.normalized_zero_count,
                negative_count=validation.normalized_negative_count,
                raw_sum=validation.raw_sum,
                normalized_sum=validation.normalized_sum,
                difference=validation.difference,
                anos_encontrados=treatment.anos_encontrados,
                itens_encontrados=treatment.itens_encontrados,
                warnings=treatment.warnings,
                validation_failures=treatment.validation_failures,
                inconsistencies=validation.inconsistencies.copy(deep=True),
            )
        )

    consolidated_data = _consolidate_normalized_frames(normalized_frames)
    consolidated_inconsistencies = _consolidate_inconsistency_frames(
        consolidated_inconsistency_frames
    )

    duplicate_count = int(
        consolidated_data.duplicated(subset=TRACEABILITY_COLUMNS, keep=False).sum()
    )
    if duplicate_count:
        duplicate_inconsistency = pd.DataFrame.from_records(
            [
                {
                    "aba_origem": pd.NA,
                    "tipo": "coordenadas_duplicadas_consolidado",
                    "ano_lancamento": pd.NA,
                    "item_informacao": pd.NA,
                    "linha_origem": pd.NA,
                    "coluna_origem": pd.NA,
                    "valor_matriz_original": 0,
                    "valor_base_normalizada": duplicate_count,
                    "diferenca": duplicate_count,
                    "mensagem": (
                        "O consolidado possui coordenadas de origem duplicadas."
                    ),
                }
            ],
            columns=CONSOLIDATED_INCONSISTENCY_COLUMNS,
        )
        consolidated_inconsistencies = pd.concat(
            [consolidated_inconsistencies, duplicate_inconsistency],
            ignore_index=True,
        )

    raw_sum = sum(
        result.raw_sum
        for result in sheet_results
        if result.recognized and result.processing_error is None
    )
    normalized_sum = sum(
        result.normalized_sum
        for result in sheet_results
        if result.recognized and result.processing_error is None
    )
    consolidated_difference = normalized_sum - raw_sum
    if abs(consolidated_difference) > tolerance:
        sum_inconsistency = pd.DataFrame.from_records(
            [
                {
                    "aba_origem": pd.NA,
                    "tipo": "soma_total_consolidada",
                    "ano_lancamento": pd.NA,
                    "item_informacao": pd.NA,
                    "linha_origem": pd.NA,
                    "coluna_origem": pd.NA,
                    "valor_matriz_original": raw_sum,
                    "valor_base_normalizada": normalized_sum,
                    "diferenca": consolidated_difference,
                    "mensagem": "A soma consolidada difere acima da tolerância.",
                }
            ],
            columns=CONSOLIDATED_INCONSISTENCY_COLUMNS,
        )
        consolidated_inconsistencies = pd.concat(
            [consolidated_inconsistencies, sum_inconsistency],
            ignore_index=True,
        )

    return DotacaoAnualWorkbookResult(
        filename=filename,
        sheet_results=sheet_results,
        consolidated_data=consolidated_data,
        consolidated_inconsistencies=consolidated_inconsistencies,
        consolidated_duplicate_count=duplicate_count,
    )


def build_sheet_summary(result: DotacaoAnualWorkbookResult) -> pd.DataFrame:
    """Monta uma tabela leve para inspeção dos resultados por aba."""

    records = []
    for sheet in result.sheet_results:
        records.append(
            {
                "Aba": sheet.sheet_name,
                "Reconhecida": "Sim" if sheet.recognized else "Não",
                "Status de integridade": sheet.status_label,
                "Linhas brutas de dados": sheet.raw_data_row_count,
                "Colunas monetárias": sheet.monetary_column_count,
                "Linhas normalizadas": sheet.normalized_row_count,
                "Nulos": sheet.null_count,
                "Zeros": sheet.zero_count,
                "Negativos": sheet.negative_count,
                "Soma original": sheet.raw_sum,
                "Soma normalizada": sheet.normalized_sum,
                "Diferença": sheet.difference,
                "Anos encontrados": ", ".join(str(ano) for ano in sheet.anos_encontrados),
                "Avisos": len(sheet.warnings),
                "Falhas de validação": len(sheet.validation_failures),
            }
        )
    return pd.DataFrame.from_records(records)


def _consolidate_normalized_frames(frames: list[pd.DataFrame]) -> pd.DataFrame:
    if not frames:
        return _empty_output_dataframe()
    return pd.concat(frames, ignore_index=True)


def _consolidate_inconsistency_frames(frames: list[pd.DataFrame]) -> pd.DataFrame:
    if not frames:
        return pd.DataFrame(columns=CONSOLIDATED_INCONSISTENCY_COLUMNS)
    return pd.concat(frames, ignore_index=True)


def _empty_output_dataframe() -> pd.DataFrame:
    return pd.DataFrame(columns=DOTACAO_ANUAL_OUTPUT_COLUMNS)


def _format_error(error: Exception) -> str:
    return f"{type(error).__name__}: {error}"
