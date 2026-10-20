"""Relatório "Projeção pela Execução" — Contratos Contínuos (tela, PDF e Excel).

Camada: regra de relatório. Não lê planilha nem importa Streamlit. Relatório SEPARADO (pedido
explícito, 06/10/2026): não altera a Necessidade de Empenho, o Registro de contratos nem o Relatório
de Reforço — só os LÊ. Recebe a saída de `relatorio_necessidade_empenho.necessidade_por_ne` (custo,
saldo e necessidade contratual por NE, a mesma conta do Resumo Consolidado) e a Liquidação por
Competência por (NE, mês); a regra do fator está na primeira seção deste arquivo (abaixo).

Por NE: fator de execução (origem, peso, meses usados), grade Jan–Dez com o realizado, o restante dos
meses em aberto e a despesa projetada até dezembro, e:
  * Despesa projetada pela execução (restante do exercício) e a mesma pelo valor contratado cheio
    (fator 1, mesmos meses) — a diferença é o efeito isolado do fator;
  * Necessidade pela execução = max(0, despesa projetada pela execução − saldo do empenho), ao lado
    da Necessidade até dezembro contratual (Resumo Consolidado), para comparação.
Contrato sem NE fica de fora (não há execução para medir) e é contado nos avisos.

Contrato antecessor (herança do fator para NE sem histórico): informado por quem emite, a cada
emissão (`antecessores`); não é gravado no cadastro.

Pedido (06/10/2026): projetar a despesa levando em conta o que foi de fato liquidado — grande parte
dos contratos liquida menos que o valor mensal. Decisão do usuário: a diferença é EXECUÇÃO ABAIXO DO
CONTRATADO (economia real), não valor retido a pagar depois — por isso o fator abaixo de 1 é usado
como está. Contrato novo repete o valor mensal (fator 1).

REGRA (calibrada no teste retroativo sobre os meses já realizados de 2026, ver histórico):
  * Razão de cada mês = liquidado por competência ÷ custo contratado do mês (`custo_mensal`: valor
    em vigor, aditivos e meses proporcionais já embutidos).
  * Só meses FECHADOS entram: até `ultimo_mes_fechado` (o relatório usa 3 meses antes do mês da
    extração — os dois meses anteriores à extração ainda recebem liquidação). Dos meses fechados com
    registro e custo > 0, os últimos `JANELA_MESES` (6): tira janeiro–março (recesso e pico de
    retroativo de abril) da conta na maior parte do ano.
  * Execução observada = soma do liquidado ÷ soma do custo nesses meses (acumulada: atraso
    compensado no mês seguinte se anula, ao contrário da mediana).
  * Peso da execução = n·τ² ÷ (n·τ² + σ²), com n meses e σ² a variância das razões mensais:
    cresce com o número de meses e cai quanto mais a execução oscila. τ² = `TAU2_PRUDENCIA_ALTA`
    (0,01 — a mais prudente das testadas; a calibrada pelos dados, ~0,046, projeta menos e erra mais
    para baixo). Razões todas iguais (σ² = 0): peso 1.
  * Fator = 1 + (execução observada − 1) × peso. Menos de `MIN_MESES` (3) meses: sem histórico,
    fator 1 (valor mensal cheio) — ou o fator do contrato ANTECESSOR, quando informado.
  * Mês em aberto (os dois anteriores ao da extração): registro abaixo de `LIMIAR_MES_ABERTO` (50%)
    do esperado (custo × fator), ou sem registro → projeta o restante (esperado − liquidado);
    com 50% ou mais, conta como completo. Do mês da extração em diante: esperado − o que já houver
    de liquidado no mês (nunca negativo).

  * Teto para NE SEM histórico (14/10/2026, só Contratos Contínuos): sem fator medido, o mês projetado (valor cheio)
    não passa do MAIOR mês já liquidado por competência na NE no exercício — a NE ainda não liquidou o contratado
    cheio. Sem nenhum mês liquidado positivo, não há teto. NE com fator próprio ou herdado do antecessor: sem teto.
    A projeção pelo valor cheio (coluna de comparação) NÃO leva o teto.

Valores: nulo ≠ zero. Liquidação negativa (estorno) entra na conta como veio; custo nulo (valor
desconhecido) não projeta (nulo); custo zero (fora da execução) não projeta nem conta como mês.
"""

from __future__ import annotations

import math
from dataclasses import dataclass, field
from io import BytesIO
from xml.sax.saxutils import escape

import pandas as pd
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from src.contratos_aditivos import aditivos_do_registro, custo_mensal
from src.relatorio_necessidade_empenho import EXERCICIO_ENCERRADO

# ------------------------------------------------------------------ fator de execução (regra)
JANELA_MESES = 6
MIN_MESES = 3
TAU2_PRUDENCIA_ALTA = 0.01
LIMIAR_MES_ABERTO = 0.5
#: meses entre o último mês fechado e o mês da extração (extração em outubro → fechados até julho)
DEFASAGEM_MESES_ABERTOS = 2

ORIGEM_EXECUCAO = "Execução observada"
ORIGEM_HERDADO = "Herdado do antecessor"
ORIGEM_SEM_HISTORICO = "Sem histórico (valor cheio)"

#: de onde vem o liquidado da NE (13/10/2026, `src/liquidacao_consumo.py`)
BASE_COMPETENCIA = "Competência"
BASE_LANCAMENTO = "Lançamento (consumo do orçamento)"

TIPO_REALIZADO = "Realizado"
TIPO_RESTANTE_ABERTO = "Restante do mês em aberto"
TIPO_PROJETADO = "Projetado"


def _numero(valor: object) -> float | None:
    if valor is None:
        return None
    try:
        numero = float(valor)
    except (TypeError, ValueError):
        return None
    return None if math.isnan(numero) else numero


@dataclass(frozen=True)
class FatorExecucao:
    fator: float
    peso: float
    #: liquidado ÷ custo acumulado nos meses usados; nulo sem histórico
    execucao_observada: float | None
    meses_usados: tuple[int, ...]
    origem: str
    ne_antecessora: str | None = None


def ultimo_mes_fechado(mes_referencia: int) -> int:
    """Último mês fechado do exercício para o mês da extração (0 a 12): três meses antes dele."""

    return mes_referencia - DEFASAGEM_MESES_ABERTOS - 1


def fator_de_execucao(
    realizado: dict[int, float], custo: list[float], ultimo_fechado: int, tau2: float = TAU2_PRUDENCIA_ALTA,
) -> FatorExecucao:
    """Fator de execução de uma NE (ver docstring do módulo). `realizado`: mês (1-12) → liquidado por
    competência no exercício (só meses com registro). `custo`: 12 custos contratados do mês."""

    meses = [
        mes for mes in range(1, 13)
        if mes <= ultimo_fechado and _numero(realizado.get(mes)) is not None
        and (_numero(custo[mes - 1]) or 0.0) > 0
    ][-JANELA_MESES:]
    if len(meses) < MIN_MESES:
        return FatorExecucao(1.0, 0.0, None, tuple(meses), ORIGEM_SEM_HISTORICO)

    liquidados = [float(realizado[mes]) for mes in meses]
    custos = [float(custo[mes - 1]) for mes in meses]
    observada = sum(liquidados) / sum(custos)
    razoes = [l / c for l, c in zip(liquidados, custos)]
    media = sum(razoes) / len(razoes)
    variancia = sum((r - media) ** 2 for r in razoes) / (len(razoes) - 1)
    n = len(meses)
    peso = 1.0 if variancia <= 1e-12 else n * tau2 / (n * tau2 + variancia)
    return FatorExecucao(1.0 + (observada - 1.0) * peso, peso, observada, tuple(meses), ORIGEM_EXECUCAO)


def herdar_fator(proprio: FatorExecucao, antecessor: FatorExecucao | None, ne_antecessora: str | None) -> FatorExecucao:
    """O fator do antecessor substitui o próprio só quando a NE não tem histórico e o antecessor tem.
    Caso contrário devolve o próprio (sem alteração)."""

    if proprio.origem != ORIGEM_SEM_HISTORICO or antecessor is None or antecessor.origem != ORIGEM_EXECUCAO:
        return proprio
    return FatorExecucao(
        antecessor.fator, antecessor.peso, antecessor.execucao_observada, antecessor.meses_usados,
        ORIGEM_HERDADO, ne_antecessora,
    )


@dataclass(frozen=True)
class ProjecaoMensal:
    #: 12 valores: liquidado por competência do mês (nulo sem registro)
    realizado: list[float | None]
    #: 12 valores: despesa projetada pela execução (nulo quando o mês não é projetado)
    projetado: list[float | None]
    #: 12 valores: a mesma projeção pelo valor contratado cheio (fator 1), nos mesmos meses
    projetado_valor_cheio: list[float | None]
    #: 12 tipos (`TIPO_*`) ou nulo
    tipos: list[str | None]
    #: algum mês projetado foi reduzido pelo teto (maior mês liquidado, NE sem histórico)
    limitado_pelo_teto: bool = False
    #: 12 valores: esperado do mês (custo × fator, com o teto) nos meses projetados — o projetado é esperado − liquidado
    esperado: list[float | None] = field(default_factory=lambda: [None] * 12)
    #: 12 booleanos: o mês tem ajuste do usuário em vigor (substituiu o esperado)
    ajustado: list[bool] = field(default_factory=lambda: [False] * 12)
    #: 12 valores: esperado que a METODOLOGIA daria (custo × fator, com teto) nos meses editáveis; nulo onde ela não projeta
    metodologia: list[float | None] = field(default_factory=lambda: [None] * 12)
    #: meses (1-12) com ajuste gravado que não valem mais (mês já fechado ou exercício encerrado)
    superados: list[int] = field(default_factory=list)

    @property
    def total_projetado(self) -> float:
        return sum(v for v in self.projetado if v is not None)

    @property
    def total_valor_cheio(self) -> float:
        return sum(v for v in self.projetado_valor_cheio if v is not None)

    def restante_em_aberto(self) -> list[tuple[int, float]]:
        return [
            (mes, self.projetado[mes - 1]) for mes in range(1, 13)
            if self.tipos[mes - 1] == TIPO_RESTANTE_ABERTO and (self.projetado[mes - 1] or 0.0) > 0
        ]


#: tolerância (em bolsas) para tratar liquidado ÷ valor da bolsa como número inteiro
TOLERANCIA_BOLSA_INTEIRA = 0.01


def arredondar_em_unidades(valor: float, unidade: float | None) -> float:
    """`valor` arredondado para o múltiplo mais próximo de `unidade` (metade para cima). Sem unidade válida
    (nula ou não positiva), devolve `valor` como veio. Bolsa é paga por inteiro ou não é paga (16/10/2026)."""

    if unidade is None or unidade <= 0:
        return valor
    return math.floor(valor / unidade + 0.5 + 1e-9) * unidade


def quantidade_nao_inteira(realizado: dict[int, float], meses: tuple[int, ...], unidade: float) -> bool:
    """Algum dos `meses` tem liquidado que não é múltiplo de `unidade` (quantidade de bolsas fracionária)."""

    for mes in meses:
        valor = _numero(realizado.get(mes))
        if valor is None:
            continue
        quantidade = valor / unidade
        if abs(quantidade - round(quantidade)) > TOLERANCIA_BOLSA_INTEIRA:
            return True
    return False


def mes_sem_execucao(liquidado: float | None) -> bool:
    """Mês sem liquidação no exercício: sem registro, ou registro zerado (16/10/2026). Mês fechado nessa situação aceita
    ajuste manual do esperado (fatura ainda não liquidada); com liquidação, vale o realizado."""

    return liquidado is None or abs(liquidado) < 0.005


def projetar_meses(
    realizado: dict[int, float], custo: list[float], fator: float, mes_referencia: int,
    teto_mensal: float | None = None, ajustes: dict[int, float] | None = None, unidade: float | None = None,
) -> ProjecaoMensal:
    """Grade Jan–Dez de uma NE (ver docstring do módulo). Exercício encerrado (`mes_referencia` ≥ 13,
    `EXERCICIO_ENCERRADO`, mesma regra do relatório de Necessidade): nada é projetado. `teto_mensal` (opcional): o esperado do mês
    (custo × fator) não passa dele. `ajustes` (opcional): mês (1-12) → ESPERADO TOTAL do mês gravado pelo usuário
    (14/10/2026); substitui custo × fator e o teto, em qualquer mês a partir do primeiro em aberto — inclusive nos que
    a metodologia deixa vazios. O liquidado do mês continua descontado. Mês já FECHADO sem execução (`mes_sem_execucao`,
    16/10/2026) também aceita ajuste: entra como restante a liquidar (tipo `TIPO_RESTANTE_ABERTO`, fatura pendente).
    Ajuste de mês fechado COM liquidação, ou de exercício encerrado, não vale: vai para `superados`. `unidade` (opcional, Bolsas e Auxílios): valor de uma bolsa — o
    esperado do mês (custo × fator) é arredondado para um número inteiro de bolsas antes do teto e do abatimento do
    liquidado; a projeção pelo valor cheio não é arredondada."""

    ajustes = {int(mes): float(valor) for mes, valor in (ajustes or {}).items()}
    # A comparação "valor cheio" não leva ajuste: com ajustes, ela vem da mesma projeção SEM eles.
    cheio_sem_ajuste = (
        projetar_meses(realizado, custo, fator, mes_referencia, teto_mensal).projetado_valor_cheio if ajustes else None
    )
    fechado = ultimo_mes_fechado(mes_referencia)
    abertos = {mes for mes in range(fechado + 1, mes_referencia) if 1 <= mes <= 12}
    realizado_lista: list[float | None] = [_numero(realizado.get(mes)) for mes in range(1, 13)]
    projetado: list[float | None] = [None] * 12
    cheio: list[float | None] = [None] * 12
    esperados: list[float | None] = [None] * 12
    tipos: list[str | None] = ["Realizado" if v is not None else None for v in realizado_lista]
    metodologias: list[float | None] = [None] * 12
    ajustados = [False] * 12
    limitado = False
    if mes_referencia >= EXERCICIO_ENCERRADO:
        return ProjecaoMensal(realizado_lista, projetado, cheio, tipos, superados=sorted(ajustes))
    superados = sorted(
        mes for mes in ajustes
        if not 1 <= mes <= 12 or (mes <= fechado and not mes_sem_execucao(realizado_lista[mes - 1]))
    )
    # Mês fechado sem execução com esperado informado (16/10/2026): fatura pendente, projetada por inteiro.
    for mes in sorted(m for m in ajustes if 1 <= m <= min(fechado, 12) and m not in superados):
        esperado = ajustes[mes]
        ja = realizado_lista[mes - 1] or 0.0
        ajustados[mes - 1] = True
        projetado[mes - 1] = max(0.0, esperado - ja)
        esperados[mes - 1] = esperado
        tipos[mes - 1] = TIPO_RESTANTE_ABERTO

    for mes in range(max(fechado + 1, 1), 13):
        custo_mes = _numero(custo[mes - 1])
        com_custo = custo_mes is not None and custo_mes > 0
        liquidado = realizado_lista[mes - 1]
        metodo = arredondar_em_unidades(custo_mes * fator, unidade) if com_custo else None
        teto_aplicado = metodo is not None and teto_mensal is not None and metodo > teto_mensal
        if teto_aplicado:
            metodo = teto_mensal
        metodologias[mes - 1] = metodo
        if mes in ajustes:
            esperado = ajustes[mes]
            ajustados[mes - 1] = True
        elif metodo is None:
            continue
        else:
            esperado = metodo
            limitado = limitado or teto_aplicado
        if mes in abertos:
            if mes not in ajustes and liquidado is not None and liquidado >= LIMIAR_MES_ABERTO * esperado:
                continue  # mês em aberto que já recebeu a maior parte: completo (não vale para o número do usuário)
            tipo = TIPO_RESTANTE_ABERTO
        elif mes < mes_referencia:
            continue
        else:
            tipo = TIPO_PROJETADO
        ja = liquidado or 0.0
        if com_custo:
            cheio[mes - 1] = max(0.0, custo_mes - ja)
        if liquidado is not None and liquidado > 0 and esperado - ja <= 0:
            # Mês (da extração em diante) já liquidado por inteiro (13/10/2026, ex.: contrato pago antecipado): nada
            # a projetar; fica como Realizado, igual aos meses anteriores, em vez de um "Projetado" de 0,00 que
            # escondia o liquidado na grade. A projeção pelo valor cheio (`cheio`) segue como antes.
            if mes in ajustes:
                # O ajuste não entrou na conta (20/10/2026): o mês vale o realizado, como em qualquer mês que fecha.
                # Sem isto o mês seguia marcado como ajustado e fora de `superados`, e o aviso de ajuste listava um
                # mês que a grade mostra como Realizado.
                ajustados[mes - 1] = False
                superados.append(mes)
            continue
        projetado[mes - 1] = max(0.0, esperado - ja)
        esperados[mes - 1] = esperado
        tipos[mes - 1] = tipo
    return ProjecaoMensal(
        realizado_lista, projetado, cheio if cheio_sem_ajuste is None else cheio_sem_ajuste, tipos, limitado, esperados,
        ajustados, metodologias, sorted(superados),
    )


# ------------------------------------------------------------------ relatório
TITULO = "PROJEÇÃO DA DESPESA PELA EXECUÇÃO — CONTRATOS CONTÍNUOS"
MESES = ["Jan", "Fev", "Mar", "Abr", "Mai", "Jun", "Jul", "Ago", "Set", "Out", "Nov", "Dez"]

#: fator abaixo disso: liquidação muito menor que o cadastrado — provável despesa mensal a revisar
LIMIAR_ABAIXO_DO_CADASTRO = 0.5
#: execução observada acima disso: liquida mais que o contratado — provável aditivo não cadastrado
LIMIAR_ACIMA_DO_CONTRATO = 1.03

NOTA_REGRA = (
    f"Fator de execução = 1 + (execução observada − 1) × peso. Execução observada = liquidado por "
    f"competência ÷ custo contratado, acumulados nos últimos {JANELA_MESES} meses fechados (os dois meses "
    f"anteriores ao da extração ainda estão em aberto). Peso = n·τ² ÷ (n·τ² + variância das razões "
    f"mensais), τ² = {TAU2_PRUDENCIA_ALTA} (prudência alta): cresce com o nº de meses e cai com a oscilação. "
    f"Menos de {MIN_MESES} meses: fator 1 (valor mensal cheio) ou o do antecessor informado. Mês em aberto "
    f"com menos de {int(LIMIAR_MES_ABERTO * 100)}% do esperado (ou sem registro): projeta o restante. Do mês "
    "da extração a dezembro: custo contratado do mês × fator. Necessidade pela execução = despesa projetada "
    "− saldo do empenho (nunca negativa). A diferença para o valor contratado é tratada como execução abaixo "
    "do contratado (decisão de 06/10/2026)."
)


#: legenda da grade mensal (14/10/2026): o que cada cor significa e o que NÃO está descontado
LEGENDA_GRADE = [
    "Sem cor: realizado (liquidado por competência, ou por lançamento nas NEs de competência futura).",
    "Azul (projetado): do mês da extração a dezembro — custo contratado do mês × fator de execução, MENOS o que "
    "já foi liquidado nesse mês. Sem liquidação no mês, é o esperado inteiro.",
    "Laranja (restante de mês em aberto): os dois meses anteriores ao da extração, e mês já fechado sem liquidação "
    "com esperado informado por você (fatura pendente) — o mesmo cálculo do azul "
    "(esperado − já liquidado no mês); mês com 50% ou mais do esperado já liquidado conta como completo e fica sem cor.",
    "Valor com asterisco (*): o liquidado desse mês JÁ foi abatido (esperado − liquidado); a aba Detalhe mensal traz "
    "o esperado, o liquidado e o projetado de cada mês.",
    "Valor com cruz (†): o esperado desse mês foi AJUSTADO por você (no lugar de custo × fator); a coluna Ajustado "
    "manualmente do Excel e o aviso do relatório listam as NEs.",
    "O SALDO DO EMPENHO NÃO é descontado na grade: as cores mostram a despesa esperada. O saldo só entra na "
    "coluna Necessidade pela execução (total projetado até dezembro − saldo).",
]


#: linha extra da grade (PDF) e aba "Quantidade de bolsas" (Excel) quando o relatório tem `unidades` (Bolsas)
LEGENDA_BOLSAS = (
    "Linha 'bolsas' (itálico cinza, abaixo do valor em reais): valor do mês ÷ valor de uma bolsa. O projetado é "
    "sempre número inteiro de bolsas, exceto quando o liquidado do mês já abatido não é múltiplo da bolsa; NE paga "
    "em valor variável (cadastro) não tem essa linha."
)


@dataclass(frozen=True)
class RotulosProjecao:
    """Textos que mudam conforme a base do relatório (pedido de 07/10/2026: o mesmo relatório em Bolsas e
    Auxílios). A regra do fator e da projeção é a mesma; só os nomes exibidos mudam. As colunas de
    `linhas` continuam com os nomes de contrato (`fornecedor`, `contrato_numero`, `status_contrato`) —
    quem monta a entrada de outra base preenche nelas o equivalente (ex.: programa, processo, situação)."""

    titulo: str
    nome: str
    sem_nome: str
    documento: str
    status: str
    valor_referencia: str
    necessidade_referencia: str
    sem_ne: str
    onde_ver_sem_ne: str
    acima_do_cadastro: str
    nota_regra: str


CONTRATOS = RotulosProjecao(
    titulo=TITULO,
    nome="Fornecedor",
    sem_nome="(sem fornecedor)",
    documento="Contrato",
    status="Status",
    valor_referencia="valor contratado",
    necessidade_referencia="contratual",
    sem_ne="item(ns) de contrato sem NE",
    onde_ver_sem_ne="o relatório de Necessidade de Empenho",
    acima_do_cadastro="Liquida acima do contratado — provável reajuste/aditivo não cadastrado",
    nota_regra=NOTA_REGRA + (
        " NE sem histórico (menos de 3 meses fechados): o mês projetado não passa do maior mês já liquidado da NE "
        "no exercício (decisão de 14/10/2026); sem nenhum mês liquidado, não há teto. O usuário pode gravar o esperado "
        "total de um mês (†): ele substitui custo × fator e o teto, o liquidado do mês continua descontado e, ao "
        "fechar o mês, vale o realizado."
    ),
)

BOLSAS_AUXILIOS = RotulosProjecao(
    titulo="PROJEÇÃO DA DESPESA PELA EXECUÇÃO — BOLSAS E AUXÍLIOS",
    nome="Programa",
    sem_nome="(sem programa)",
    documento="Processo",
    status="Situação",
    valor_referencia="valor cadastrado",
    necessidade_referencia="Resumo Consolidado",
    sem_ne="programa(s) sem NE",
    onde_ver_sem_ne="o Resumo Consolidado",
    acima_do_cadastro="Liquida acima do cadastrado — conferir quantidade e valor unitário do cadastro",
    nota_regra=(
        NOTA_REGRA.replace("custo contratado", "custo cadastrado").replace("valor contratado", "valor cadastrado")
        .replace("do contratado (decisão de 06/10/2026)", "do cadastrado (mesma decisão de 06/10/2026 dos contratos)")
        + " Bolsa é paga por inteiro: o esperado de cada mês projetado (custo × fator) é arredondado para o "
        "número inteiro de bolsas mais próximo (metade para cima), nunca fracionado (decisão de 16/10/2026). "
        "Custo cadastrado da bolsa: quantidade efetiva × valor unitário em cada um dos meses do ano "
        "cadastrados, contados a partir do início da execução, até dezembro. Início, nesta ordem: o informado "
        "no cadastro; o primeiro mês do exercício com liquidação positiva por competência da NE; o mês do "
        "primeiro empenho; janeiro (decisão de 07/10/2026)."
    ),
)


@dataclass(frozen=True)
class ContextoProjecaoExecucao:
    exercicio: int
    data_extracao: str
    hash_manifesto: str
    data_emissao: str
    origem_competencia: str


@dataclass(frozen=True)
class RelatorioProjecaoExecucao:
    #: uma linha por NE, em ordem decrescente de necessidade pela execução
    linhas: pd.DataFrame
    #: mesmo índice: r1..r12 (realizado), p1..p12 (projetado pela execução), t1..t12 (tipo do mês)
    mensal: pd.DataFrame
    exercicio: int
    mes_referencia: int
    ultimo_fechado: int
    qtd_sem_ne: int
    avisos: list[str] = field(default_factory=list)
    rotulos: RotulosProjecao = CONTRATOS
    #: NE → meses com ajuste gravado que não vale mais (mês já fechado / exercício encerrado)
    superados: dict[str, list[int]] = field(default_factory=dict)
    #: NEs com ajuste gravado que não estão no relatório (sem NE, ou NE que saiu da base)
    ajustes_sem_ne: list[str] = field(default_factory=list)
    #: NE → valor de uma bolsa (só Bolsas e Auxílios, ver `montar_relatorio(unidades=)`); vazio nos demais
    unidades: dict[str, float] = field(default_factory=dict)

    def _soma(self, coluna: str) -> float:
        return float(self.linhas[coluna].sum(min_count=1)) if not self.linhas.empty else 0.0

    @property
    def total_projetado_execucao(self) -> float:
        return self._soma("projetado_execucao")

    @property
    def total_projetado_valor_cheio(self) -> float:
        return self._soma("projetado_valor_cheio")

    @property
    def total_restante_em_aberto(self) -> float:
        return self._soma("restante_em_aberto")

    @property
    def total_necessidade_execucao(self) -> float:
        return self._soma("necessidade_execucao")

    @property
    def total_necessidade_contratual(self) -> float:
        return self._soma("necessidade_contratual")

    @property
    def total_saldo(self) -> float:
        return self._soma("saldo")


def _realizado_por_ne(liquidacao_mensal: pd.DataFrame | None, exercicio: int) -> dict[str, dict[int, float]]:
    if liquidacao_mensal is None or liquidacao_mensal.empty:
        return {}
    ano_mes = liquidacao_mensal["ano_mes"].astype("int64")
    serie = liquidacao_mensal.assign(ano=ano_mes // 100, mes=ano_mes % 100)
    serie = serie[(serie["ano"] == exercicio) & serie["mes"].between(1, 12)]
    agrupado = serie.groupby(["ne_curta", "mes"])["valor"].sum(min_count=1)
    resultado: dict[str, dict[int, float]] = {}
    for (ne, mes), valor in agrupado.items():
        if not pd.isna(valor):
            resultado.setdefault(ne, {})[int(mes)] = float(valor)
    return resultado


def _custo_da_ne(linha: pd.Series, exercicio: int) -> list[float]:
    """Custo contratado mês a mês da NE — a MESMA chamada de `_aplicar_necessidade` (Resumo)."""

    return custo_mensal(
        linha["despesa_mensal"], aditivos_do_registro(linha["aditivos"]), exercicio,
        status=linha["status_contrato"], vigencia_fim=linha["vigencia_fim"],
        inicio=linha["inicio_execucao_considerado"], data_suspensao=linha["data_suspensao"],
        meses_no_ano=linha["meses_no_ano"],
    )


def _custo(linha: pd.Series, exercicio: int, custos: dict[str, list[float]] | None) -> list[float]:
    """Custo da NE: o informado em `custos` (base com regra própria de custo, ex.: Bolsas e Auxílios) ou,
    sem ele, o de contrato (`_custo_da_ne`)."""

    return _custo_da_ne(linha, exercicio) if custos is None else custos[linha["ne_curta"]]


def fatores_por_ne(
    por_ne: pd.DataFrame, liquidacao_mensal: pd.DataFrame | None, exercicio: int, mes_referencia: int,
    custos: dict[str, list[float]] | None = None,
) -> dict[str, object]:
    """Fator próprio de cada NE (sem herança) — a tela usa para listar quem pode ser antecessor.
    `custos` (opcional): NE → 12 custos do mês, no lugar do custo de contrato."""

    realizados = _realizado_por_ne(liquidacao_mensal, exercicio)
    fechado = ultimo_mes_fechado(mes_referencia)
    return {
        linha["ne_curta"]: fator_de_execucao(realizados.get(linha["ne_curta"], {}), _custo(linha, exercicio, custos), fechado)
        for _, linha in por_ne.iterrows()
    }


#: colunas de `RelatorioProjecaoExecucao.linhas`, na ordem em que cada NE é registrada em `montar_relatorio`.
COLUNAS_LINHAS = [
    "ne_curta", "fornecedor", "contrato_numero", "status_contrato", "vigencia_fim", "despesa_mensal",
    "valor_empenhado", "saldo", "origem_fator", "ne_antecessora", "execucao_observada", "peso", "fator",
    "meses_usados", "restante_em_aberto", "projetado_execucao", "projetado_valor_cheio",
    "necessidade_execucao", "necessidade_contratual", "base_liquidacao", "limitada_ao_maior_mes",
    "ajustada_manualmente",
]


def montar_relatorio(
    por_ne: pd.DataFrame,
    sem_ne: pd.DataFrame,
    liquidacao_mensal: pd.DataFrame | None,
    exercicio: int,
    mes_referencia: int,
    antecessores: dict[str, str] | None = None,
    *,
    custos: dict[str, list[float]] | None = None,
    rotulos: RotulosProjecao = CONTRATOS,
    nes_pelo_lancamento: set[str] | frozenset[str] | list[str] | None = None,
    teto_sem_historico: bool = False,
    ajustes: dict[str, dict[int, float]] | None = None,
    unidades: dict[str, float] | None = None,
    nes_valor_variavel: set[str] | frozenset[str] | None = None,
) -> RelatorioProjecaoExecucao:
    """`por_ne`/`sem_ne`: saída de `necessidade_por_ne(..., exercicio)`. `liquidacao_mensal`:
    `ne_curta`, `ano_mes`, `valor` (Liquidação por Competência). `antecessores`: NE sem histórico →
    NE cujo fator ela herda. `custos`/`rotulos`: para outra base que não Contratos Contínuos (ex.:
    `src.projecao_execucao_bolsas`) — o custo do mês de cada NE já calculado pela regra da base e os
    textos exibidos; sem eles, custo e textos de contrato. `nes_pelo_lancamento`: NEs cujo `liquidacao_mensal`
    já vem pelo mês do lançamento (`liquidacao_consumo.competencia_pelo_consumo`) — só rotula a base de cada
    linha, não muda conta alguma. `teto_sem_historico`: liga o teto do maior mês liquidado para NE sem histórico
    (ver docstring do módulo); desligado por padrão (Bolsas e Auxílios). `ajustes`: NE → mês → ESPERADO TOTAL do mês
    gravado pelo usuário (`projecao_ajustes.reconstruir`); ver `projetar_meses`. NE com ajuste que não está em
    `por_ne` fica em `ajustes_sem_ne`. `unidades`: NE → valor de uma bolsa
    (`projecao_execucao_bolsas.unidades_por_ne`); liga a projeção em bolsas inteiras (esperado do mês
    arredondado ao múltiplo mais próximo da bolsa, decisão de 16/10/2026) e os avisos de quantidade
    fracionária e de NE sem unidade. Nulo: sem arredondamento (Contratos Contínuos). `nes_valor_variavel`: NEs
    marcadas no cadastro como pagas em valor variável — projetadas sem arredondar, com aviso próprio (não
    entram nos avisos de unidade nem de quantidade fracionária). Não altera as entradas."""

    antecessores = antecessores or {}
    pelo_lancamento = set(nes_pelo_lancamento or ())
    ajustes = ajustes or {}
    superados: dict[str, list[int]] = {}
    realizados = _realizado_por_ne(liquidacao_mensal, exercicio)
    fechado = ultimo_mes_fechado(mes_referencia)
    proprios = fatores_por_ne(por_ne, liquidacao_mensal, exercicio, mes_referencia, custos)

    registros, mensais = [], []
    avisos_antecessor = []
    fracionarias, sem_unidade, variaveis = [], [], []
    marcadas = set(nes_valor_variavel or ())
    for _, linha in por_ne.iterrows():
        ne = linha["ne_curta"]
        custo = _custo(linha, exercicio, custos)
        fator = proprios[ne]
        ne_ant = antecessores.get(ne)
        if ne_ant:
            herdado = herdar_fator(fator, proprios.get(ne_ant), ne_ant)
            if herdado is fator:
                motivo = (
                    "ela já tem histórico próprio" if fator.origem != ORIGEM_SEM_HISTORICO
                    else f"{ne_ant} não tem histórico suficiente"
                )
                avisos_antecessor.append(f"{ne}: antecessor {ne_ant} ignorado — {motivo}.")
            fator = herdado
        teto = None
        if teto_sem_historico and fator.origem == ORIGEM_SEM_HISTORICO:
            positivos = [v for v in realizados.get(ne, {}).values() if v > 0]
            teto = max(positivos) if positivos else None
        unidade = None if unidades is None else unidades.get(ne)
        if unidades is not None:
            if ne in marcadas:
                variaveis.append(ne)
            elif unidade is None:
                sem_unidade.append(ne)
            elif quantidade_nao_inteira(realizados.get(ne, {}), fator.meses_usados, unidade):
                fracionarias.append(ne)
        projecao = projetar_meses(
            realizados.get(ne, {}), custo, fator.fator, mes_referencia, teto, ajustes=ajustes.get(ne), unidade=unidade,
        )
        if projecao.superados:
            superados[ne] = projecao.superados
        saldo = linha["saldo_para_necessidade"]
        saldo_num = 0.0 if pd.isna(saldo) else float(saldo)
        restante = sum(v for _, v in projecao.restante_em_aberto())
        projetado_execucao = projecao.total_projetado
        projetado_valor_cheio = projecao.total_valor_cheio
        necessidade_execucao = max(0.0, projecao.total_projetado - saldo_num)
        # Valor mensal desconhecido (08/10/2026, pedido do usuário): as regras de custo devolvem os 12
        # meses nulos (`custo_mensal`, `custo_mensal_bolsa`) e nenhum mês era projetado — a necessidade
        # saía 0 e o Resultado Orçamentário herdava esse zero. Agora fica nula ("—"), fora dos totais e
        # avisada; valor mensal zero continua zero. Exercício encerrado não projeta nada de qualquer forma.
        if mes_referencia < EXERCICIO_ENCERRADO and all(_numero(valor) is None for valor in custo) and not any(projecao.ajustado):
            restante = projetado_execucao = projetado_valor_cheio = necessidade_execucao = None
        registros.append(
            {
                "ne_curta": ne,
                "fornecedor": linha["fornecedor"],
                "contrato_numero": linha["contrato_numero"],
                "status_contrato": linha["status_contrato"],
                "vigencia_fim": linha["vigencia_fim"],
                "despesa_mensal": linha["despesa_mensal"],
                "valor_empenhado": linha["valor_empenhado_exibido"],
                "saldo": saldo,
                "origem_fator": fator.origem,
                "ne_antecessora": fator.ne_antecessora,
                "execucao_observada": fator.execucao_observada,
                "peso": fator.peso,
                "fator": fator.fator,
                "meses_usados": ", ".join(MESES[m - 1] for m in fator.meses_usados),
                "restante_em_aberto": restante,
                "projetado_execucao": projetado_execucao,
                "projetado_valor_cheio": projetado_valor_cheio,
                "necessidade_execucao": necessidade_execucao,
                "necessidade_contratual": linha["necessidade"],
                "base_liquidacao": BASE_LANCAMENTO if ne in pelo_lancamento else BASE_COMPETENCIA,
                "limitada_ao_maior_mes": projecao.limitado_pelo_teto,
                "ajustada_manualmente": any(projecao.ajustado),
            }
        )
        mensais.append(
            {
                **{f"r{m}": projecao.realizado[m - 1] for m in range(1, 13)},
                **{f"p{m}": projecao.projetado[m - 1] for m in range(1, 13)},
                **{f"t{m}": projecao.tipos[m - 1] for m in range(1, 13)},
                **{f"e{m}": projecao.esperado[m - 1] for m in range(1, 13)},
                **{f"a{m}": projecao.ajustado[m - 1] for m in range(1, 13)},
                **{f"m{m}": projecao.metodologia[m - 1] for m in range(1, 13)},
            }
        )

    colunas_mensal = [f"r{m}" for m in range(1, 13)] + [f"p{m}" for m in range(1, 13)] + [f"t{m}" for m in range(1, 13)] + [f"e{m}" for m in range(1, 13)]
    colunas_mensal += [f"a{m}" for m in range(1, 13)] + [f"m{m}" for m in range(1, 13)]
    # Colunas fixas mesmo sem NE (13/10/2026): `pd.DataFrame([])` nascia sem coluna alguma e quebrava quem lê
    # `linhas["ne_curta"]` (ex.: busca em Contratos Contínuos que isola só contratos sem NE).
    linhas = pd.DataFrame(registros, columns=COLUNAS_LINHAS)
    mensal = pd.DataFrame(mensais, columns=colunas_mensal)
    if not linhas.empty:
        ordem = linhas["necessidade_execucao"].sort_values(ascending=False, kind="stable").index
        linhas = linhas.loc[ordem].reset_index(drop=True)
        mensal = mensal.loc[ordem].reset_index(drop=True)
        for coluna in ("execucao_observada", "peso", "fator", "saldo", "necessidade_contratual", "valor_empenhado"):
            linhas[coluna] = pd.to_numeric(linhas[coluna], errors="coerce").astype("float64")

    relatorio = RelatorioProjecaoExecucao(
        linhas=linhas, mensal=mensal, exercicio=exercicio, mes_referencia=mes_referencia,
        ultimo_fechado=fechado, qtd_sem_ne=len(sem_ne), rotulos=rotulos, superados=superados,
        ajustes_sem_ne=sorted(ne for ne in ajustes if ne not in set(por_ne["ne_curta"])), unidades=dict(unidades or {}),
    )
    avisos_bolsas = []
    if fracionarias:
        avisos_bolsas.append(
            "Liquidado por competência que não é múltiplo do valor da bolsa (quantidade de bolsas fracionária) nos "
            "meses usados no fator — o fator usa o liquidado como veio, sem descarte, e só a quantidade projetada é "
            "arredondada: " + ", ".join(fracionarias) + "."
        )
    if sem_unidade:
        avisos_bolsas.append(
            "Valor unitário diferente entre os programas da NE, ou valor mensal diferente de quantidade × valor "
            "unitário: projetada sem arredondamento em bolsas inteiras: " + ", ".join(sem_unidade) + "."
        )
    if variaveis:
        avisos_bolsas.append(
            "Marcadas no cadastro como pagas em valor variável (auxílio por demanda, sem bolsa inteira): projetadas "
            "pelo fator, sem arredondamento: " + ", ".join(variaveis) + "."
        )
    relatorio.avisos.extend(_avisos(relatorio) + avisos_antecessor + avisos_bolsas)
    return relatorio


def _avisos(relatorio: RelatorioProjecaoExecucao) -> list[str]:
    linhas = relatorio.linhas
    rotulos = relatorio.rotulos

    def _rotulo_ne(linha: pd.Series) -> str:
        documento = linha["contrato_numero"]
        return f"{linha['ne_curta']}" + (f" ({rotulos.documento.lower()} {documento})" if pd.notna(documento) else "")

    avisos = []
    if relatorio.mes_referencia >= EXERCICIO_ENCERRADO:
        avisos.append("Extração em ano posterior ao exercício: exercício encerrado, nada é projetado.")
    if linhas.empty:
        return avisos
    sem_valor = linhas[linhas["necessidade_execucao"].isna()]
    if not sem_valor.empty:
        avisos.append(
            "Sem valor mensal cadastrado — necessidade pela execução não calculada (—) e fora dos totais: "
            + ", ".join(_rotulo_ne(l) for _, l in sem_valor.iterrows()) + "."
        )
    abaixo =linhas[(linhas["origem_fator"] == ORIGEM_EXECUCAO) & (linhas["fator"] < LIMIAR_ABAIXO_DO_CADASTRO)]
    if not abaixo.empty:
        avisos.append(
            "Liquidação muito abaixo do cadastrado (fator < "
            f"{LIMIAR_ABAIXO_DO_CADASTRO:.2f}) — conferir a despesa mensal do cadastro: "
            + ", ".join(f"{_rotulo_ne(l)} fator {l['fator']:.2f}" for _, l in abaixo.iterrows()) + "."
        )
    acima = linhas[(linhas["origem_fator"] == ORIGEM_EXECUCAO) & (linhas["execucao_observada"] > LIMIAR_ACIMA_DO_CONTRATO)]
    if not acima.empty:
        avisos.append(
            f"{rotulos.acima_do_cadastro}: "
            + ", ".join(f"{_rotulo_ne(l)} execução {l['execucao_observada']:.2f}" for _, l in acima.iterrows()) + "."
        )
    herdados = linhas[linhas["origem_fator"] == ORIGEM_HERDADO]
    if not herdados.empty:
        avisos.append(
            "Fator herdado do contrato antecessor informado: "
            + ", ".join(f"{_rotulo_ne(l)} ← {l['ne_antecessora']} ({l['fator']:.2f})" for _, l in herdados.iterrows()) + "."
        )
    sem_historico = linhas[(linhas["origem_fator"] == ORIGEM_SEM_HISTORICO) & (linhas["projetado_execucao"] > 0)]
    if not sem_historico.empty:
        avisos.append(
            f"Menos de {MIN_MESES} meses fechados de execução — projetado pelo valor mensal cheio: "
            + ", ".join(_rotulo_ne(l) for _, l in sem_historico.iterrows()) + "."
        )
    limitadas = linhas[linhas["limitada_ao_maior_mes"].fillna(False).astype(bool)]
    if not limitadas.empty:
        avisos.append(
            "Sem histórico de execução e projeção limitada ao maior mês já liquidado da NE (o valor cheio do cadastro "
            "não foi liquidado até aqui): " + ", ".join(_rotulo_ne(l) for _, l in limitadas.iterrows()) + "."
        )
    def _titulo_ne(linha: pd.Series) -> str:
        """NE com o título do contrato: número e fornecedor (como na conciliação por competência)."""
        partes = [f"{rotulos.documento} {_texto(linha['contrato_numero']) or '—'}", _texto(linha["fornecedor"]) or rotulos.sem_nome]
        return f"{linha['ne_curta']} ({' — '.join(partes)})"

    por_ne_curta = linhas.set_index("ne_curta", drop=False)
    ajustadas = linhas[linhas["ajustada_manualmente"].fillna(False).astype(bool)]
    if not ajustadas.empty:
        avisos.append(
            "Projeção ajustada manualmente (esperado do mês gravado pelo usuário): "
            + "; ".join(
                f"{_titulo_ne(l)} — "
                + ", ".join(MESES[m - 1] for m in range(1, 13) if relatorio.mensal.at[i, f"a{m}"])
                for i, l in ajustadas.iterrows()
            ) + "."
        )
    if relatorio.superados:
        avisos.append(
            "Ajuste superado (o mês já fechou: vale o realizado; o ajuste continua gravado): "
            + "; ".join(
                f"{_titulo_ne(por_ne_curta.loc[ne])} — " + ", ".join(MESES[m - 1] for m in meses if 1 <= m <= 12)
                for ne, meses in relatorio.superados.items() if ne in por_ne_curta.index
            ) + "."
        )
    if relatorio.ajustes_sem_ne:
        avisos.append(
            "Ajuste gravado para NE que não está no relatório (ignorado): " + ", ".join(relatorio.ajustes_sem_ne) + "."
        )
    em_aberto = linhas[linhas["restante_em_aberto"] > 0].sort_values("restante_em_aberto", ascending=False)
    if not em_aberto.empty:
        maiores = ", ".join(
            f"{l['ne_curta']} R$ {_formatar_brl(l['restante_em_aberto'])}" for _, l in em_aberto.head(5).iterrows()
        )
        avisos.append(
            f"Meses em aberto (liquidação ainda incompleta ou sem registro): {len(em_aberto)} NE(s), restante "
            f"projetado R$ {_formatar_brl(em_aberto['restante_em_aberto'].sum())} — maiores: {maiores} "
            "(detalhe mês a mês na grade, em laranja)."
        )
    if relatorio.qtd_sem_ne:
        avisos.append(
            f"{relatorio.qtd_sem_ne} {rotulos.sem_ne} fora deste relatório (sem execução para medir) — "
            f"ver {rotulos.onde_ver_sem_ne}."
        )
    return avisos


# ------------------------------------------------------------------------------ formatação
def _formatar_brl(valor: object) -> str:
    if valor is None or pd.isna(valor):
        return "—"
    texto = f"{float(valor):,.2f}"
    return texto.replace(",", "X").replace(".", ",").replace("X", ".")


def _formatar_bolsas(valor: object) -> str:
    """Quantidade de bolsas: inteira sem casas (8), fracionária com até 2 (8,3); nulo vira "—"."""

    if valor is None or pd.isna(valor):
        return "—"
    return f"{float(valor):.2f}".rstrip("0").rstrip(".").replace(".", ",")


def _formatar_fator(valor: object) -> str:
    return "—" if valor is None or pd.isna(valor) else f"{float(valor):.2f}".replace(".", ",")


def _texto(valor: object) -> str:
    return "" if valor is None or pd.isna(valor) else str(valor)


def linhas_parametros(relatorio: RelatorioProjecaoExecucao, contexto: ContextoProjecaoExecucao) -> list[tuple[str, str]]:
    fechado = relatorio.ultimo_fechado
    rotulos = relatorio.rotulos
    return [
        ("Exercício", str(contexto.exercicio)),
        ("Extração da Execução Mensal", f"{contexto.data_extracao} (manifesto {contexto.hash_manifesto})"),
        ("Liquidação por Competência", contexto.origem_competencia),
        ("Meses fechados (base do fator)", f"até {MESES[fechado - 1]}" if 1 <= fechado <= 12 else "nenhum"),
        ("Emitido em", contexto.data_emissao),
        ("NEs no relatório", str(len(relatorio.linhas))),
        ("Saldo dos empenhos (R$)", _formatar_brl(relatorio.total_saldo)),
        ("Despesa projetada pela execução (R$)", _formatar_brl(relatorio.total_projetado_execucao)),
        ("  dos quais restante de meses em aberto (R$)", _formatar_brl(relatorio.total_restante_em_aberto)),
        (f"Despesa projetada pelo {rotulos.valor_referencia} (R$)", _formatar_brl(relatorio.total_projetado_valor_cheio)),
        ("Necessidade pela execução (R$)", _formatar_brl(relatorio.total_necessidade_execucao)),
        (f"Necessidade até dezembro — {rotulos.necessidade_referencia} (R$)", _formatar_brl(relatorio.total_necessidade_contratual)),
        ("Regra", rotulos.nota_regra),
    ]


def _ou_nulo(valor: object) -> float | None:
    return None if valor is None or pd.isna(valor) else float(valor)


def liquidado_abatido(relatorio: RelatorioProjecaoExecucao, indice: int, mes: int) -> bool:
    """Mês projetado (azul/laranja) cujo valor JÁ descontou liquidado do próprio mês (esperado − liquidado)."""

    if relatorio.mensal.at[indice, f"t{mes}"] not in (TIPO_PROJETADO, TIPO_RESTANTE_ABERTO):
        return False
    realizado = relatorio.mensal.at[indice, f"r{mes}"]
    return realizado is not None and not pd.isna(realizado) and realizado != 0


def _legenda(relatorio: RelatorioProjecaoExecucao) -> list[str]:
    """Legenda da grade. A edição do esperado vale em Contratos Contínuos e, desde 15/10/2026, em Bolsas e Auxílios."""

    return LEGENDA_GRADE


def ajustado_no_mes(relatorio: RelatorioProjecaoExecucao, indice: int, mes: int) -> bool:
    """Mês projetado (azul/laranja) cujo esperado é o AJUSTE gravado pelo usuário, e não o da metodologia."""

    return (
        relatorio.mensal.at[indice, f"t{mes}"] in (TIPO_PROJETADO, TIPO_RESTANTE_ABERTO)
        and bool(relatorio.mensal.at[indice, f"a{mes}"])
    )


def _marcas(relatorio: RelatorioProjecaoExecucao, indice: int, mes: int) -> str:
    """Marcas do valor na grade do Excel e do PDF: `*` liquidado abatido, `†` esperado ajustado pelo usuário."""

    return ("*" if liquidado_abatido(relatorio, indice, mes) else "") + ("†" if ajustado_no_mes(relatorio, indice, mes) else "")


COLUNAS_QUANTIDADE_DETALHE = ["Esperado (bolsas)", "Realizado (bolsas)", "Projetado (bolsas)"]
ABA_QUANTIDADE = "Quantidade de bolsas"


def _em_bolsas(valor: object, unidade: float | None) -> float | None:
    """`valor` ÷ valor da bolsa (4 casas, sem ruído de ponto flutuante); nulo continua nulo."""

    numero = _ou_nulo(valor)
    return None if numero is None or unidade is None else round(numero / unidade, 4)


def _resumo(relatorio: RelatorioProjecaoExecucao) -> pd.DataFrame:
    linhas = relatorio.linhas
    rotulos = relatorio.rotulos
    resumo = pd.DataFrame(
        {
            "NE": linhas["ne_curta"].map(_texto).astype(object),
            rotulos.nome: linhas["fornecedor"].map(_texto).astype(object),
            rotulos.documento: linhas["contrato_numero"].map(_texto).astype(object),
            rotulos.status: linhas["status_contrato"].map(_texto).astype(object),
            "Despesa mensal": linhas["despesa_mensal"],
            "Empenhado": linhas["valor_empenhado"],
            "Saldo do empenho": linhas["saldo"],
            "Origem do fator": linhas["origem_fator"].astype(object),
            "Base da liquidação": linhas["base_liquidacao"].astype(object),
            "Limitada ao maior mês liquidado": linhas["limitada_ao_maior_mes"].map({True: "Sim", False: "Não"}).astype(object),
            "Ajustado manualmente": linhas["ajustada_manualmente"].map({True: "Sim", False: "Não"}).astype(object),
            "NE antecessora": linhas["ne_antecessora"].map(_texto).replace("", None).astype(object),
            "Meses usados": linhas["meses_usados"].replace("", None).astype(object),
            "Execução observada": linhas["execucao_observada"],
            "Peso da execução": linhas["peso"],
            "Fator de execução": linhas["fator"],
            "Restante de meses em aberto": linhas["restante_em_aberto"],
            "Despesa projetada (execução)": linhas["projetado_execucao"],
            f"Despesa projetada ({rotulos.valor_referencia})": linhas["projetado_valor_cheio"],
            "Necessidade pela execução": linhas["necessidade_execucao"],
            f"Necessidade até dezembro ({rotulos.necessidade_referencia})": linhas["necessidade_contratual"],
        }
    )
    return resumo


_COLUNAS_FATOR_XLSX = {"Execução observada", "Peso da execução", "Fator de execução"}
_COLUNAS_TEXTO_XLSX = {"NE", "Origem do fator", "Base da liquidação", "Limitada ao maior mês liquidado", "Ajustado manualmente", "NE antecessora", "Meses usados", "Tipo", "Liquidado abatido", "Ajustado"}
_PREENCHIMENTO_PROJETADO = "DCE6F2"
_PREENCHIMENTO_ABERTO = "FCE4C4"


def _formatar_aba(planilha, dados: pd.DataFrame, rotulos: RotulosProjecao) -> None:
    colunas_texto = _COLUNAS_TEXTO_XLSX | {rotulos.nome, rotulos.documento, rotulos.status}
    for indice, cabecalho in enumerate(dados.columns, start=1):
        letra = planilha.cell(row=1, column=indice).column_letter
        formato = "@" if cabecalho in colunas_texto else ("0.00" if cabecalho in _COLUNAS_FATOR_XLSX else "#,##0.00")
        for celula in planilha[letra][1:]:
            celula.number_format = formato
        planilha.column_dimensions[letra].width = 45 if cabecalho == rotulos.nome else min(max(len(cabecalho), 12) + 2, 34)
    planilha.freeze_panes = "B2"


def gerar_xlsx(relatorio: RelatorioProjecaoExecucao, contexto: ContextoProjecaoExecucao) -> bytes:
    """Abas: "Resumo por NE", "Grade mensal" (Jan–Dez: realizado ou projetado; projetado sombreado,
    restante de mês em aberto em laranja), "Detalhe mensal" (formato longo com o tipo de cada mês) e
    "Parâmetros" (avisos, totais e regra). Nulo é célula vazia, nunca zero."""

    from openpyxl.styles import Font, PatternFill

    rotulos = relatorio.rotulos
    resumo = _resumo(relatorio)
    grade_linhas, marcas, detalhe = [], [], []
    marcas_quantidade: list[list] = []
    quantidade_linhas = []
    for indice, linha in relatorio.linhas.iterrows():
        valores, tipos = {}, []
        unidade = relatorio.unidades.get(linha["ne_curta"])
        quantidades = {}
        for mes in range(1, 13):
            tipo = relatorio.mensal.at[indice, f"t{mes}"]
            projetado = relatorio.mensal.at[indice, f"p{mes}"]
            realizado = relatorio.mensal.at[indice, f"r{mes}"]
            if tipo in (TIPO_PROJETADO, TIPO_RESTANTE_ABERTO):
                valores[MESES[mes - 1]] = projetado
            else:
                valores[MESES[mes - 1]] = realizado
            tipos.append(tipo)
            quantidades[MESES[mes - 1]] = _em_bolsas(valores[MESES[mes - 1]], unidade)
            if tipo is not None:
                detalhe.append(
                    {
                        "NE": linha["ne_curta"], rotulos.nome: _texto(linha["fornecedor"]), "Nº do mês": mes,
                        "Mês": MESES[mes - 1], "Tipo": tipo,
                        "Esperado do mês (custo × fator)": _ou_nulo(relatorio.mensal.at[indice, f"e{mes}"]),
                        "Realizado": None if pd.isna(realizado) else realizado,
                        "Projetado": None if projetado is None or pd.isna(projetado) else projetado,
                        "Liquidado abatido": "Sim" if liquidado_abatido(relatorio, indice, mes) else None,
                        "Ajustado": "Sim" if ajustado_no_mes(relatorio, indice, mes) else None,
                        COLUNAS_QUANTIDADE_DETALHE[0]: _em_bolsas(relatorio.mensal.at[indice, f"e{mes}"], unidade),
                        COLUNAS_QUANTIDADE_DETALHE[1]: _em_bolsas(realizado, unidade),
                        COLUNAS_QUANTIDADE_DETALHE[2]: _em_bolsas(projetado, unidade),
                    }
                )
        grade_linhas.append({"NE": linha["ne_curta"], rotulos.nome: _texto(linha["fornecedor"]), "Fator de execução": linha["fator"], **valores,
                             "Despesa projetada (execução)": linha["projetado_execucao"]})
        marcas.append(tipos)
        if unidade is not None:
            quantidade_linhas.append(
                {"NE": linha["ne_curta"], rotulos.nome: _texto(linha["fornecedor"]), "Valor da bolsa (R$)": unidade, **quantidades}
            )
            marcas_quantidade.append(tipos)
    grade = pd.DataFrame(grade_linhas, columns=["NE", rotulos.nome, "Fator de execução", *MESES, "Despesa projetada (execução)"])
    detalhe_df = pd.DataFrame(
        detalhe, columns=[
            "NE", rotulos.nome, "Nº do mês", "Mês", "Tipo", "Esperado do mês (custo × fator)", "Realizado", "Projetado",
            "Liquidado abatido", *COLUNAS_QUANTIDADE_DETALHE, "Ajustado",
        ],
    )
    if not relatorio.unidades:
        detalhe_df = detalhe_df.drop(columns=COLUNAS_QUANTIDADE_DETALHE)
    quantidade = pd.DataFrame(quantidade_linhas, columns=["NE", rotulos.nome, "Valor da bolsa (R$)", *MESES])
    parametros = (
        [("Avisos", aviso) for aviso in relatorio.avisos] + linhas_parametros(relatorio, contexto)
        + [("Legenda da grade", texto) for texto in _legenda(relatorio)]
    )
    coluna_jan = list(grade.columns).index(MESES[0]) + 1

    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        resumo.to_excel(writer, sheet_name="Resumo por NE", index=False)
        _formatar_aba(writer.sheets["Resumo por NE"], resumo, rotulos)
        grade.to_excel(writer, sheet_name="Grade mensal", index=False)
        planilha = writer.sheets["Grade mensal"]
        _formatar_aba(planilha, grade, rotulos)
        for numero_linha, tipos in enumerate(marcas, start=2):
            for mes, tipo in enumerate(tipos):
                if tipo in (TIPO_PROJETADO, TIPO_RESTANTE_ABERTO):
                    celula = planilha.cell(row=numero_linha, column=coluna_jan + mes)
                    cor = _PREENCHIMENTO_ABERTO if tipo == TIPO_RESTANTE_ABERTO else _PREENCHIMENTO_PROJETADO
                    celula.fill = PatternFill("solid", fgColor=cor)
                    celula.font = Font(italic=True)
                    marcas = _marcas(relatorio, relatorio.linhas.index[numero_linha - 2], mes + 1)
                    if marcas:
                        celula.number_format = f'#,##0.00"{marcas}"'
        primeira_legenda = len(grade) + 3
        planilha.cell(row=primeira_legenda - 1, column=1, value="Legenda").font = Font(bold=True)
        for deslocamento, texto in enumerate(_legenda(relatorio)):
            planilha.cell(row=primeira_legenda + deslocamento, column=1, value=texto)
        detalhe_df.to_excel(writer, sheet_name="Detalhe mensal", index=False)
        _formatar_aba(writer.sheets["Detalhe mensal"], detalhe_df, rotulos)
        writer.sheets["Detalhe mensal"].column_dimensions["C"].width = 10
        for celula in writer.sheets["Detalhe mensal"]["C"][1:]:
            celula.number_format = "0"
        if relatorio.unidades:
            for cabecalho in COLUNAS_QUANTIDADE_DETALHE:
                letra = writer.sheets["Detalhe mensal"].cell(row=1, column=list(detalhe_df.columns).index(cabecalho) + 1).column_letter
                for celula in writer.sheets["Detalhe mensal"][letra][1:]:
                    celula.number_format = "General"
        if relatorio.unidades:
            quantidade.to_excel(writer, sheet_name=ABA_QUANTIDADE, index=False)
            planilha_q = writer.sheets[ABA_QUANTIDADE]
            _formatar_aba(planilha_q, quantidade, rotulos)
            for celula in planilha_q["C"][1:]:
                celula.number_format = "#,##0.00"
            for coluna in range(4, 4 + 12):  # Jan–Dez: bolsas (inteiras no projetado; fração só no liquidado)
                for celula in planilha_q[planilha_q.cell(row=1, column=coluna).column_letter][1:]:
                    celula.number_format = "General"
            for numero_linha, tipos in enumerate(marcas_quantidade, start=2):
                for mes, tipo in enumerate(tipos):
                    if tipo in (TIPO_PROJETADO, TIPO_RESTANTE_ABERTO):
                        celula = planilha_q.cell(row=numero_linha, column=4 + mes)
                        cor = _PREENCHIMENTO_ABERTO if tipo == TIPO_RESTANTE_ABERTO else _PREENCHIMENTO_PROJETADO
                        celula.fill = PatternFill("solid", fgColor=cor)
                        celula.font = Font(italic=True)
            planilha_q.cell(row=len(quantidade) + 3, column=1, value=(
                "Quantidade de bolsas = valor do mês ÷ valor de uma bolsa. Sombreado: projetado (azul) ou restante de "
                "mês em aberto (laranja), mesmas cores da grade em reais. NE fora desta aba: valor unitário não "
                "utilizável (ver avisos)."
            ))
        pd.DataFrame(parametros, columns=["Parâmetro", "Valor"]).to_excel(writer, sheet_name="Parâmetros", index=False)
        writer.sheets["Parâmetros"].column_dimensions["A"].width = 46
        writer.sheets["Parâmetros"].column_dimensions["B"].width = 120
    return buffer.getvalue()


# -------------------------------------------------------------------------------- PDF
_CABECALHO_GRADE_PDF = ["NE", "FATOR", *[m.upper() for m in MESES], "PROJETADO"]
_LARGURAS_GRADE_PDF = [62, 36] + [48] * 12 + [66]
_CABECALHO_RESUMO_PDF = [
    "NE", "{nome} / {documento}", "ORIGEM DO FATOR", "EXEC. OBS.", "PESO", "FATOR", "SALDO (R$)",
    "PROJ. EXECUÇÃO (R$)", "PROJ. {valor} (R$)", "NEC. EXECUÇÃO (R$)", "NEC. {necessidade} (R$)",
]
_LARGURAS_RESUMO_PDF = [58, 170, 78, 38, 36, 40, 66, 72, 72, 72, 72]  # 774 pt de 784 úteis (A4 paisagem, margens de 10 mm)
_COR_TITULO_NE = colors.HexColor("#E8EEF7")
_COR_PROJETADO = colors.HexColor(f"#{_PREENCHIMENTO_PROJETADO}")
_COR_ABERTO = colors.HexColor(f"#{_PREENCHIMENTO_ABERTO}")


def _estilo_base() -> list[tuple]:
    return [
        ("FONTSIZE", (0, 0), (-1, -1), 6.5),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DDDDDD")),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]


def _cabecalho_resumo_pdf(rotulos: RotulosProjecao) -> list[str]:
    referencia = rotulos.valor_referencia.split()[-1]  # "contratado" / "cadastrado"
    necessidade = "CONTRATUAL" if rotulos.necessidade_referencia == "contratual" else "RESUMO"
    return [
        celula.format(nome=rotulos.nome.upper(), documento=rotulos.documento.upper(), valor=referencia.upper(),
                      necessidade=necessidade)
        for celula in _CABECALHO_RESUMO_PDF
    ]


def gerar_pdf(relatorio: RelatorioProjecaoExecucao, contexto: ContextoProjecaoExecucao) -> bytes:
    rotulos = relatorio.rotulos
    buffer = BytesIO()
    documento = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        leftMargin=10 * mm, rightMargin=10 * mm, topMargin=10 * mm, bottomMargin=10 * mm,
        title=f"{rotulos.titulo} — {contexto.exercicio}",
    )
    estilos = getSampleStyleSheet()
    estilo_parametro = estilos["Normal"].clone("parametro_projecao")
    estilo_parametro.fontSize = 8
    estilo_parametro.leading = 10
    estilo_aviso = estilo_parametro.clone("aviso_projecao")
    estilo_aviso.textColor = colors.HexColor("#8A4B00")
    estilo_celula = estilos["Normal"].clone("celula_projecao")
    estilo_celula.fontSize = 6.5
    estilo_celula.leading = 7.5
    estilo_secao = estilos["Heading4"]
    # Cabeçalho do resumo quebra em linhas dentro da coluna (13/10/2026): em Paragraph o rótulo longo
    # ("PROJ. EXECUÇÃO (R$)") não invade a coluna vizinha como acontecia com texto simples.
    estilo_cabecalho = estilos["Normal"].clone("cabecalho_projecao")
    estilo_cabecalho.fontName = "Helvetica-Bold"
    estilo_cabecalho.fontSize = 6.5
    estilo_cabecalho.leading = 7.5
    estilo_cabecalho.alignment = 1  # centralizado

    elementos: list = [Paragraph(escape(f"{rotulos.titulo} — EXERCÍCIO {contexto.exercicio}"), estilos["Heading3"])]
    for aviso in relatorio.avisos:
        elementos.append(Paragraph(f"<b>ATENÇÃO:</b> {escape(aviso)}", estilo_aviso))
    for rotulo, valor in linhas_parametros(relatorio, contexto):
        elementos.append(Paragraph(f"<b>{escape(rotulo.strip())}:</b> {escape(valor)}", estilo_parametro))
    elementos.append(Spacer(1, 8))

    # ---- resumo por NE
    dados = [[Paragraph(escape(rotulo), estilo_cabecalho) for rotulo in _cabecalho_resumo_pdf(rotulos)]]
    for _, linha in relatorio.linhas.iterrows():
        origem = linha["origem_fator"] + (f" ({linha['ne_antecessora']})" if pd.notna(linha["ne_antecessora"]) and linha["ne_antecessora"] else "")
        dados.append(
            [
                linha["ne_curta"],
                Paragraph(f"{escape(_texto(linha['fornecedor']) or rotulos.sem_nome)}<br/>{rotulos.documento} {escape(_texto(linha['contrato_numero']) or '—')}", estilo_celula),
                Paragraph(escape(origem), estilo_celula),
                _formatar_fator(linha["execucao_observada"]), _formatar_fator(linha["peso"]), _formatar_fator(linha["fator"]),
                _formatar_brl(linha["saldo"]), _formatar_brl(linha["projetado_execucao"]),
                _formatar_brl(linha["projetado_valor_cheio"]), _formatar_brl(linha["necessidade_execucao"]),
                _formatar_brl(linha["necessidade_contratual"]),
            ]
        )
    dados.append(
        [
            "TOTAL", "", "", "", "", "", _formatar_brl(relatorio.total_saldo),
            _formatar_brl(relatorio.total_projetado_execucao), _formatar_brl(relatorio.total_projetado_valor_cheio),
            _formatar_brl(relatorio.total_necessidade_execucao), _formatar_brl(relatorio.total_necessidade_contratual),
        ]
    )
    tabela = Table(dados, colWidths=_LARGURAS_RESUMO_PDF, repeatRows=1)
    tabela.setStyle(TableStyle(_estilo_base() + [
        ("ALIGN", (3, 1), (-1, -1), "RIGHT"),
        ("LEFTPADDING", (0, 0), (-1, 0), 3), ("RIGHTPADDING", (0, 0), (-1, 0), 3),  # cabeçalho: mais espaço útil
        ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#F5F5F5")]),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
    ]))
    elementos.append(Paragraph("Resumo por NE — fator de execução, despesa projetada e necessidade", estilo_secao))
    elementos.append(tabela)
    elementos.append(Spacer(1, 10))

    # ---- grade mensal
    dados = [_CABECALHO_GRADE_PDF]
    comandos: list[tuple] = []
    for indice, linha in relatorio.linhas.iterrows():
        posicao = len(dados)
        titulo = (
            f"<b>{escape(linha['ne_curta'])}</b> — {escape(_texto(linha['fornecedor']) or rotulos.sem_nome)}"
            f" — {rotulos.documento} {escape(_texto(linha['contrato_numero']) or '—')} — {escape(linha['origem_fator'])}"
            f" — Meses usados: {escape(_texto(linha['meses_usados']) or '—')}"
            + (f" — Liquidação: {escape(linha['base_liquidacao'])}" if linha["base_liquidacao"] == BASE_LANCAMENTO else "")
        )
        dados.append([Paragraph(titulo, estilo_celula)] + [""] * (len(_CABECALHO_GRADE_PDF) - 1))
        unidade = relatorio.unidades.get(linha["ne_curta"])
        celulas, bolsas = [], []
        for mes in range(1, 13):
            tipo = relatorio.mensal.at[indice, f"t{mes}"]
            if tipo in (TIPO_PROJETADO, TIPO_RESTANTE_ABERTO):
                celulas.append(
                    _formatar_brl(relatorio.mensal.at[indice, f"p{mes}"]) + _marcas(relatorio, indice, mes)
                )
                bolsas.append(_formatar_bolsas(_em_bolsas(relatorio.mensal.at[indice, f"p{mes}"], unidade)))
                cor = _COR_ABERTO if tipo == TIPO_RESTANTE_ABERTO else _COR_PROJETADO
                comandos += [
                    ("BACKGROUND", (mes + 1, posicao + 1), (mes + 1, posicao + 1), cor),
                    ("FONTNAME", (mes + 1, posicao + 1), (mes + 1, posicao + 1), "Helvetica-Oblique"),
                ]
            else:
                celulas.append(_formatar_brl(relatorio.mensal.at[indice, f"r{mes}"]))
                bolsas.append(_formatar_bolsas(_em_bolsas(relatorio.mensal.at[indice, f"r{mes}"], unidade)))
        dados.append([linha["ne_curta"], _formatar_fator(linha["fator"]), *celulas, _formatar_brl(linha["projetado_execucao"])])
        comandos += [
            ("SPAN", (0, posicao), (-1, posicao)),
            ("BACKGROUND", (0, posicao), (-1, posicao), _COR_TITULO_NE),
            ("NOSPLIT", (0, posicao), (-1, posicao + (2 if unidade is not None else 1))),
        ]
        if unidade is not None:
            # linha de quantidade de bolsas (valor do mês ÷ valor da bolsa), logo abaixo da linha em reais
            dados.append(["", f"bolsas (R$ {_formatar_brl(unidade)})", *bolsas, ""])
            comandos += [
                ("TEXTCOLOR", (0, posicao + 2), (-1, posicao + 2), colors.HexColor("#555555")),
                ("FONTNAME", (0, posicao + 2), (-1, posicao + 2), "Helvetica-Oblique"),
            ]
    tabela = Table(dados, colWidths=_LARGURAS_GRADE_PDF, repeatRows=1)
    tabela.setStyle(TableStyle(_estilo_base() + [("ALIGN", (1, 0), (-1, -1), "RIGHT")] + comandos))
    elementos.append(Paragraph(
        "Realizado × projeção mês a mês (projetado sombreado em azul; restante de mês em aberto em laranja)", estilo_secao
    ))
    for texto in _legenda(relatorio):
        elementos.append(Paragraph(escape(texto), estilo_celula))
    if relatorio.unidades:
        elementos.append(Paragraph(escape(LEGENDA_BOLSAS), estilo_celula))
    elementos.append(Spacer(1, 6))
    elementos.append(tabela)
    documento.build(elementos)
    return buffer.getvalue()


def aplicar_projecao_na_necessidade(por_ne: pd.DataFrame, relatorio: RelatorioProjecaoExecucao) -> pd.DataFrame:
    """Cópia de `por_ne` (saída de `necessidade_por_ne`) com a Necessidade pela PROJEÇÃO PELA EXECUÇÃO: `necessidade`
    passa a ser a `necessidade_execucao` do relatório de cada NE, e a contratual fica em `necessidade_contratual`.
    É a escolha "base da necessidade" da página de Contratos Contínuos (12/10/2026). NE que o relatório não projetou,
    ou com valor mensal desconhecido (necessidade nula), fica NULA — nunca vira zero; zero projetado continua zero.
    Não altera as entradas."""

    por_execucao = relatorio.linhas.set_index("ne_curta")["necessidade_execucao"]
    saida = por_ne.copy()
    saida["necessidade_contratual"] = saida["necessidade"]
    saida["necessidade"] = pd.to_numeric(saida["ne_curta"].map(por_execucao), errors="coerce").astype("float64")
    return saida
