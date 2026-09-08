"""Leitura do relatório anual de acompanhamento de Emendas Parlamentares.

O arquivo observado segue o padrão do Tesouro Gerencial: três linhas de
cabeçalho, dimensões representadas por células realmente mescladas e quatro
medidas monetárias em colunas. A leitura usa ``openpyxl`` diretamente para
resolver somente mesclagens confirmadas pelo arquivo; não há ``forward fill``
heurístico.

Granularidade normalizada: uma linha por exercício × emenda × PTRES × GND.
Uma emenda pode ter mais de um PTRES e um PTRES pode ocupar mais de uma linha
quando houver mais de um GND. Códigos permanecem texto e nulo, zero e negativo
são preservados como estados distintos.
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
import re
import unicodedata

import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet


HEADER_ROWS = 3
FIRST_DATA_ROW = HEADER_ROWS + 1

COLUNAS_SAIDA = [
    "arquivo_origem",
    "aba_origem",
    "linha_origem",
    "resultado_primario_cod",
    "ptres",
    "emenda_numero",
    "autor_emenda",
    "parlamentar",
    "gnd_cod",
    "ano",
    "dotacao_atualizada",
    "empenhada_relatorio",
    "liquidada_relatorio",
    "paga_relatorio",
]

COLUNAS_IDENTIFICADORAS = [
    "arquivo_origem",
    "aba_origem",
    "resultado_primario_cod",
    "ptres",
    "emenda_numero",
    "autor_emenda",
    "parlamentar",
    "gnd_cod",
]

COLUNAS_FINANCEIRAS = [
    "dotacao_atualizada",
    "empenhada_relatorio",
    "liquidada_relatorio",
    "paga_relatorio",
]

_CABECALHOS_FIXOS = {
    "A1": "RESULTADO PRIMARIO LEI",
    "B1": "PTRES",
    "C1": "AUTOR EMENDAS ORCAMENTO",
    "E1": "GRUPO DESPESA",
    "F1": "ITEM INFORMACAO",
    "F3": "ANO LANCAMENTO",
}

_COLUNAS_MEDIDAS = {
    7: ("13", "DOTACAO ATUALIZADA", "dotacao_atualizada"),
    8: ("23", "DESPESAS EMPENHADAS", "empenhada_relatorio"),
    9: ("25", "DESPESAS LIQUIDADAS", "liquidada_relatorio"),
    10: ("28", "DESPESAS PAGAS", "paga_relatorio"),
}

_MESCLAGENS_CABECALHO = {"A1:A3", "B1:B3", "C1:D3", "E1:E3", "F1:F2"}
_RPS_EMENDA = {"6", "7", "8"}


class ErroLayoutBase(ValueError):
    """O arquivo não possui o contrato estrutural esperado."""


class ErroDadosBase(ValueError):
    """O layout foi reconhecido, mas um valor obrigatório é inválido."""


class _MergeResolver:
    """Resolve somente células pertencentes a intervalos mesclados reais."""

    def __init__(self, worksheet: Worksheet) -> None:
        self._worksheet = worksheet
        self._owners: dict[tuple[int, int], object] = {}
        for merged_range in worksheet.merged_cells.ranges:
            owner = worksheet.cell(merged_range.min_row, merged_range.min_col).value
            for row in range(merged_range.min_row, merged_range.max_row + 1):
                for column in range(merged_range.min_col, merged_range.max_col + 1):
                    self._owners[(row, column)] = owner

    def value(self, row: int, column: int) -> object:
        value = self._worksheet.cell(row, column).value
        if value is not None:
            return value
        return self._owners.get((row, column))


def ler_emendas_acompanhamento(caminho: str | Path) -> pd.DataFrame:
    """Lê um ``.xlsx`` de acompanhamento sem alterar o arquivo original."""

    caminho = Path(caminho)
    if not caminho.exists():
        raise FileNotFoundError(caminho)
    if caminho.suffix.lower() != ".xlsx":
        raise ErroLayoutBase(
            "O relatório de Emendas deve ser .xlsx, pois a leitura depende "
            "das mesclagens reais do arquivo."
        )
    return ler_emendas_acompanhamento_bytes(caminho.read_bytes(), caminho.name)


def ler_emendas_acompanhamento_bytes(
    conteudo: bytes,
    nome_arquivo: str,
) -> pd.DataFrame:
    """Versão em memória do leitor, adequada a uploads futuros do Streamlit."""

    if Path(nome_arquivo).suffix.lower() != ".xlsx":
        raise ErroLayoutBase(
            "O relatório de Emendas deve ser .xlsx, pois a leitura depende "
            "das mesclagens reais do arquivo."
        )

    try:
        workbook = load_workbook(BytesIO(conteudo), data_only=False, read_only=False)
    except Exception as error:
        raise ErroLayoutBase(f"Não foi possível abrir {nome_arquivo!r}: {error}") from error

    reconhecidas: list[Worksheet] = []
    falhas_por_aba: list[str] = []
    for worksheet in workbook.worksheets:
        falhas = _falhas_layout(worksheet)
        if falhas:
            falhas_por_aba.append(f"{worksheet.title}: {'; '.join(falhas)}")
        else:
            reconhecidas.append(worksheet)

    if not reconhecidas:
        detalhe = " | ".join(falhas_por_aba)
        raise ErroLayoutBase(
            f"Nenhuma aba reconhecida em {nome_arquivo!r}. {detalhe}"
        )
    if len(reconhecidas) > 1:
        nomes = [worksheet.title for worksheet in reconhecidas]
        raise ErroLayoutBase(
            f"Mais de uma aba corresponde ao relatório de Emendas: {nomes}."
        )

    return _normalizar_aba(reconhecidas[0], nome_arquivo)


def reconciliar_emendas_acompanhamento(df: pd.DataFrame) -> dict[str, object]:
    """Resume contagens, totais e estados financeiros sem preencher nulos."""

    _exigir_colunas(df, COLUNAS_SAIDA)
    totais = {
        coluna: _soma_ou_none(df[coluna])
        for coluna in COLUNAS_FINANCEIRAS
    }
    totais_por_ano: dict[int, dict[str, float | None]] = {}
    for ano, grupo in df.groupby("ano", dropna=False):
        if pd.isna(ano):
            continue
        totais_por_ano[int(ano)] = {
            coluna: _soma_ou_none(grupo[coluna])
            for coluna in COLUNAS_FINANCEIRAS
        }

    return {
        "linhas": int(len(df)),
        "anos": sorted(int(ano) for ano in df["ano"].dropna().unique()),
        "emendas_distintas": int(
            df[["ano", "emenda_numero"]].drop_duplicates().shape[0]
        ),
        "ptres_distintos": int(df["ptres"].nunique(dropna=True)),
        "totais": totais,
        "totais_por_ano": totais_por_ano,
        "nulos": {
            coluna: int(df[coluna].isna().sum())
            for coluna in COLUNAS_FINANCEIRAS
        },
        "zeros": {
            coluna: int(df[coluna].eq(0).sum())
            for coluna in COLUNAS_FINANCEIRAS
        },
        "negativos": {
            coluna: int(df[coluna].lt(0).sum())
            for coluna in COLUNAS_FINANCEIRAS
        },
    }


def _falhas_layout(worksheet: Worksheet) -> list[str]:
    falhas: list[str] = []
    if worksheet.max_row < FIRST_DATA_ROW:
        falhas.append("menos de três linhas de cabeçalho e uma linha de dados")
    if worksheet.max_column != 10:
        falhas.append(f"{worksheet.max_column} colunas; esperadas 10")

    for coordenada, esperado in _CABECALHOS_FIXOS.items():
        encontrado = _normalizar_rotulo(worksheet[coordenada].value)
        if encontrado != esperado:
            falhas.append(
                f"{coordenada}: esperado {esperado!r}, encontrado {encontrado!r}"
            )

    mesclagens = {str(intervalo) for intervalo in worksheet.merged_cells.ranges}
    ausentes = sorted(_MESCLAGENS_CABECALHO - mesclagens)
    if ausentes:
        falhas.append(f"mesclagens de cabeçalho ausentes: {ausentes}")

    for coluna, (codigo, rotulo, _) in _COLUNAS_MEDIDAS.items():
        letra = worksheet.cell(1, coluna).column_letter
        codigo_encontrado = _as_code(worksheet.cell(1, coluna).value)
        rotulo_encontrado = _normalizar_rotulo(worksheet.cell(2, coluna).value)
        metrica = _normalizar_rotulo(worksheet.cell(3, coluna).value)
        if codigo_encontrado != codigo:
            falhas.append(
                f"{letra}1: código {codigo_encontrado!r}; esperado {codigo!r}"
            )
        if rotulo_encontrado != rotulo:
            falhas.append(
                f"{letra}2: esperado {rotulo!r}, encontrado {rotulo_encontrado!r}"
            )
        if "MOVIM LIQUIDO" not in metrica:
            falhas.append(f"{letra}3: métrica inesperada {metrica!r}")

    return falhas


def _normalizar_aba(worksheet: Worksheet, nome_arquivo: str) -> pd.DataFrame:
    resolver = _MergeResolver(worksheet)
    registros: list[dict[str, object]] = []

    for row in range(FIRST_DATA_ROW, worksheet.max_row + 1):
        valores = [resolver.value(row, column) for column in range(1, 11)]
        if all(_vazio(value) for value in valores):
            continue

        rp = _campo_obrigatorio_codigo(valores[0], worksheet, row, 1)
        ptres = _campo_obrigatorio_codigo(valores[1], worksheet, row, 2)
        emenda_numero = _campo_obrigatorio_codigo(valores[2], worksheet, row, 3)
        autor_emenda = _campo_obrigatorio_texto(valores[3], worksheet, row, 4)
        gnd = _campo_obrigatorio_codigo(valores[4], worksheet, row, 5)
        ano = _parse_ano(valores[5])
        if ano is None:
            raise ErroDadosBase(
                f"{worksheet.cell(row, 6).coordinate}: exercício inválido {valores[5]!r}."
            )
        if rp not in _RPS_EMENDA:
            raise ErroDadosBase(
                f"{worksheet.cell(row, 1).coordinate}: RP {rp!r} não identifica "
                "emenda parlamentar (esperado 6, 7 ou 8)."
            )

        registro: dict[str, object] = {
            "arquivo_origem": nome_arquivo,
            "aba_origem": worksheet.title,
            "linha_origem": row,
            "resultado_primario_cod": rp,
            "ptres": ptres,
            "emenda_numero": emenda_numero,
            "autor_emenda": autor_emenda,
            "parlamentar": _extrair_parlamentar(autor_emenda),
            "gnd_cod": gnd,
            "ano": ano,
        }
        for coluna, (_, _, nome_saida) in _COLUNAS_MEDIDAS.items():
            registro[nome_saida] = _as_number_strict(
                valores[coluna - 1], worksheet.cell(row, coluna).coordinate
            )
        registros.append(registro)

    df = _aplicar_tipos(pd.DataFrame.from_records(registros, columns=COLUNAS_SAIDA))
    if df.empty:
        raise ErroDadosBase("A aba reconhecida não possui linhas de dados.")

    chave_linha = [
        "ano",
        "resultado_primario_cod",
        "emenda_numero",
        "ptres",
        "gnd_cod",
    ]
    duplicadas = df.duplicated(subset=chave_linha, keep=False)
    if duplicadas.any():
        linhas = df.loc[duplicadas, "linha_origem"].astype(int).tolist()
        raise ErroDadosBase(
            "Há linhas duplicadas na chave exercício × RP × emenda × PTRES × GND: "
            f"linhas {linhas}."
        )
    return df


def _aplicar_tipos(df: pd.DataFrame) -> pd.DataFrame:
    typed = df.copy(deep=True)
    for coluna in COLUNAS_IDENTIFICADORAS:
        typed[coluna] = pd.array(typed[coluna], dtype="string")
    for coluna in ["linha_origem", "ano"]:
        typed[coluna] = pd.array(typed[coluna], dtype="Int64")
    for coluna in COLUNAS_FINANCEIRAS:
        typed[coluna] = pd.array(typed[coluna], dtype="Float64")
    return typed


def _campo_obrigatorio_codigo(
    value: object,
    worksheet: Worksheet,
    row: int,
    column: int,
) -> str:
    codigo = _as_code(value)
    if codigo is None:
        raise ErroDadosBase(
            f"{worksheet.cell(row, column).coordinate}: identificador obrigatório ausente."
        )
    return codigo


def _campo_obrigatorio_texto(
    value: object,
    worksheet: Worksheet,
    row: int,
    column: int,
) -> str:
    texto = _as_text(value)
    if texto is None:
        raise ErroDadosBase(
            f"{worksheet.cell(row, column).coordinate}: texto obrigatório ausente."
        )
    return texto


def _as_code(value: object) -> str | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(value)
    text = str(value).strip()
    return text or None


def _as_text(value: object) -> str | None:
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value)).strip()
    return text or None


def _as_number_strict(value: object, coordinate: str) -> object:
    if value is None or (isinstance(value, str) and not value.strip()):
        return pd.NA
    if isinstance(value, bool):
        raise ErroDadosBase(f"{coordinate}: booleano não é valor financeiro válido.")
    if isinstance(value, (int, float)):
        return float(value)

    text = str(value).strip()
    if text.startswith("="):
        raise ErroDadosBase(
            f"{coordinate}: fórmula encontrada; o relatório deve trazer valores materializados."
        )
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    try:
        return float(text)
    except ValueError as error:
        raise ErroDadosBase(
            f"{coordinate}: valor financeiro inválido {value!r}."
        ) from error


def _parse_ano(value: object) -> int | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    text = str(value).strip()
    return int(text) if text.isdigit() else None


def _extrair_parlamentar(autor_emenda: str) -> str:
    partes = re.split(r"\s*/\s*EMENDA\s+\d+\s*$", autor_emenda, flags=re.IGNORECASE)
    parlamentar = partes[0].strip()
    return parlamentar or autor_emenda


def _normalizar_rotulo(value: object) -> str:
    if value is None:
        return ""
    sem_acento = "".join(
        character
        for character in unicodedata.normalize("NFKD", str(value))
        if not unicodedata.combining(character)
    )
    return re.sub(r"[^A-Z0-9]+", " ", sem_acento.upper()).strip()


def _vazio(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _soma_ou_none(series: pd.Series) -> float | None:
    soma = series.sum(min_count=1)
    return None if pd.isna(soma) else float(soma)


def _exigir_colunas(df: pd.DataFrame, colunas: list[str]) -> None:
    faltando = [coluna for coluna in colunas if coluna not in df.columns]
    if faltando:
        raise ValueError(f"Colunas ausentes: {faltando}.")
