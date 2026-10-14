"""Testes da entrada de Bolsas e Auxílios no relatório "Projeção pela Execução"
(`src/projecao_execucao_bolsas.py` + `src/relatorio_projecao_execucao.py` com `custos`/`rotulos`).
Valores esperados calculados à mão em cada caso. O caso real usa a fixture congelada de Liquidação por
Competência (`tests/fixtures/liquidacao_competencia_2026-09-11.xlsx`), não `data/raw/`."""

from __future__ import annotations

import math
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook

from src.execucao_ne_utils import ne_curta
from src.liquidacao_competencia import ler_liquidacao_competencia, liquidado_por_ne_e_mes
from src.projecao_execucao_bolsas import (
    custo_mensal_bolsa,
    entrada_relatorio,
    nes_valor_variavel,
    primeiro_mes_liquidado_por_ne,
    unidades_por_ne,
)
from src.relatorio_projecao_execucao import (
    BOLSAS_AUXILIOS,
    ORIGEM_EXECUCAO,
    TIPO_PROJETADO,
    ContextoProjecaoExecucao,
    gerar_pdf,
    LEGENDA_BOLSAS,
    _formatar_bolsas,
    arredondar_em_unidades,
    gerar_xlsx,
    montar_relatorio,
)

FIXTURE_COMPETENCIA = Path(__file__).parent / "fixtures" / "liquidacao_competencia_2026-09-11.xlsx"


class TestCustoMensalBolsa(unittest.TestCase):
    def test_meses_a_partir_do_inicio(self):
        # 10 meses a partir de fevereiro: fev–nov; jan e dez fora
        custo = custo_mensal_bolsa(70000.0, 10, 2)
        self.assertEqual(custo, [0.0] + [70000.0] * 10 + [0.0])

    def test_para_em_dezembro(self):
        # 10 meses a partir de maio: mai–dez (8 meses); o resto cairia em 2027, fora do exercício
        custo = custo_mensal_bolsa(5600.0, 10, 5)
        self.assertEqual(custo, [0.0] * 4 + [5600.0] * 8)

    def test_sem_inicio_comeca_em_janeiro(self):
        self.assertEqual(custo_mensal_bolsa(1000.0, 3, None), [1000.0] * 3 + [0.0] * 9)
        self.assertEqual(custo_mensal_bolsa(1000.0, 3, float("nan")), [1000.0] * 3 + [0.0] * 9)

    def test_meses_no_ano_vazio_vale_12(self):
        self.assertEqual(custo_mensal_bolsa(1000.0, None, None), [1000.0] * 12)

    def test_fracao_de_mes(self):
        # 2,5 meses a partir de outubro: out 1.000, nov 1.000, dez 500
        self.assertEqual(custo_mensal_bolsa(1000.0, 2.5, 10), [0.0] * 9 + [1000.0, 1000.0, 500.0])

    def test_valor_nulo_continua_nulo_e_zero_continua_zero(self):
        self.assertTrue(all(math.isnan(v) for v in custo_mensal_bolsa(None, 12, 1)))
        self.assertEqual(custo_mensal_bolsa(0.0, 12, 1), [0.0] * 12)

    def test_inicio_invalido(self):
        with self.assertRaises(ValueError):
            custo_mensal_bolsa(1000.0, 12, 13)


def _programa(**campos: object) -> dict:
    base = {
        "ne_curta": None, "programa_bolsa": "P", "processo": "1/2026", "situacao_tg": "ATUALIZADO",
        "valor_mensal": 1000.0, "meses_no_ano": 12, "inicio_execucao_mes": None, "inicio_execucao_efetivo": None,
        "valor_empenhado_execucao": None, "valor_empenhado_tg": None,
        "saldo_execucao": None, "saldo_colado_planilha": None,
    }
    base.update(campos)
    return base


class TestEntradaRelatorio(unittest.TestCase):
    def test_uma_linha_por_ne_e_sem_ne_a_parte(self):
        filtrado = pd.DataFrame([
            # dois programas na mesma NE: valor mensal e custo somam; saldo/empenhado da NE entram uma vez
            _programa(ne_curta="2026NE000001", programa_bolsa="A", processo="10/2026", valor_mensal=1000.0,
                      meses_no_ano=12, valor_empenhado_execucao=9000.0, saldo_execucao=2000.0),
            _programa(ne_curta="2026NE000001", programa_bolsa="B", processo="10/2026", valor_mensal=500.0,
                      meses_no_ano=6, inicio_execucao_efetivo=7, valor_empenhado_execucao=9000.0,
                      saldo_execucao=2000.0),
            # sem saldo na Execução Mensal: usa o colado na planilha
            _programa(ne_curta="2026NE000002", programa_bolsa="C", valor_mensal=200.0,
                      valor_empenhado_tg=600.0, saldo_colado_planilha=150.0),
            _programa(ne_curta=None, programa_bolsa="SEM NE"),
        ])
        por_ne, sem_ne, custos = entrada_relatorio(filtrado)

        self.assertEqual(por_ne["ne_curta"].tolist(), ["2026NE000001", "2026NE000002"])
        self.assertEqual(len(sem_ne), 1)
        ne1 = por_ne.set_index("ne_curta").loc["2026NE000001"]
        self.assertEqual(ne1["fornecedor"], "A; B")
        self.assertEqual(ne1["contrato_numero"], "10/2026")
        self.assertEqual(ne1["despesa_mensal"], 1500.0)
        self.assertEqual(ne1["valor_empenhado_exibido"], 9000.0)
        self.assertEqual(ne1["saldo_para_necessidade"], 2000.0)
        # necessidade do Resumo: A = 1.000 × (12 − 9) = 3.000; B = 500 × max(0, 6 − 18) = 0
        self.assertEqual(ne1["necessidade"], 3000.0)
        self.assertEqual(custos["2026NE000001"], [1000.0] * 6 + [1500.0] * 6)
        ne2 = por_ne.set_index("ne_curta").loc["2026NE000002"]
        self.assertEqual(ne2["valor_empenhado_exibido"], 600.0)
        self.assertEqual(ne2["saldo_para_necessidade"], 150.0)
        self.assertEqual(ne2["necessidade"], 200.0 * 9)  # 12 − 600/200

    def test_inicio_pelo_primeiro_mes_liquidado(self):
        # NE 1: empenho em jan (efetivo 1), liquidação a partir de fev → início fev: fev–nov.
        # NE 2: início manual em mar prevalece sobre a liquidação (fev).
        # NE 3: sem liquidação no exercício → primeiro empenho (abr).
        # NE 4: só estorno/zero no exercício e liquidação em 2025 → não contam; sem empenho → janeiro.
        filtrado = pd.DataFrame([
            _programa(ne_curta="2026NE000001", meses_no_ano=10, inicio_execucao_efetivo=1),
            _programa(ne_curta="2026NE000002", meses_no_ano=10, inicio_execucao_mes=3, inicio_execucao_efetivo=3),
            _programa(ne_curta="2026NE000003", meses_no_ano=2, inicio_execucao_efetivo=4),
            _programa(ne_curta="2026NE000004", meses_no_ano=2),
        ])
        competencia = pd.DataFrame({
            "ne_curta": ["2026NE000001", "2026NE000001", "2026NE000002", "2026NE000004", "2026NE000004", "2026NE000004"],
            "ano_mes": [202602, 202603, 202602, 202505, 202606, 202607],
            "valor": [1000.0, 1000.0, 1000.0, 1000.0, -50.0, 0.0],
        })
        _, _, custos = entrada_relatorio(filtrado, competencia, 2026)
        self.assertEqual(custos["2026NE000001"], [0.0] + [1000.0] * 10 + [0.0])
        self.assertEqual(custos["2026NE000002"], [0.0] * 2 + [1000.0] * 10)
        self.assertEqual(custos["2026NE000003"], [0.0] * 3 + [1000.0] * 2 + [0.0] * 7)
        self.assertEqual(custos["2026NE000004"], [1000.0] * 2 + [0.0] * 10)

    def test_primeiro_mes_liquidado_por_ne(self):
        competencia = pd.DataFrame({
            "ne_curta": ["A", "A", "B", "C"], "ano_mes": [202605, 202603, 202512, 202604], "valor": [1.0, 2.0, 5.0, -1.0],
        })
        self.assertEqual(primeiro_mes_liquidado_por_ne(competencia, 2026).to_dict(), {"A": 3})
        self.assertTrue(primeiro_mes_liquidado_por_ne(None, 2026).empty)

    def test_nao_altera_a_entrada(self):
        filtrado = pd.DataFrame([_programa(ne_curta="2026NE000001")])
        copia = filtrado.copy()
        entrada_relatorio(filtrado)
        pd.testing.assert_frame_equal(filtrado, copia)


def _competencia_fixture() -> pd.DataFrame:
    tempo = liquidado_por_ne_e_mes(ler_liquidacao_competencia(FIXTURE_COMPETENCIA))
    tempo["ne_curta"] = tempo["ne_ccor"].apply(ne_curta)
    return tempo


NE_REAL = "2026NE000022"  # bolsa real: 70.000/mês, 10 meses, início em fevereiro


def _relatorio_real():
    filtrado = pd.DataFrame([
        _programa(ne_curta=NE_REAL, programa_bolsa="PROGRAMA REAL", processo="23082.000000/2026-00",
                  valor_mensal=70000.0, meses_no_ano=10, inicio_execucao_efetivo=1,
                  valor_empenhado_execucao=560000.0, saldo_execucao=10000.0),
        _programa(ne_curta=None, programa_bolsa="SEM NE"),
    ])
    competencia = _competencia_fixture()
    # início pela competência (fev), não pelo primeiro empenho (jan) — caso real que motivou a regra
    por_ne, sem_ne, custos = entrada_relatorio(filtrado, competencia, 2026)
    # extração da fixture: 11/09/2026 → mês de referência 9, fechados até junho
    return montar_relatorio(por_ne, sem_ne, competencia, 2026, 9, custos=custos, rotulos=BOLSAS_AUXILIOS)


class TestRelatorioBolsasFixture(unittest.TestCase):
    def test_projecao_calculada_a_mao(self):
        # Fixture (competência de 2026): fev–abr 70.000, mai 69.300, jun–ago 70.000.
        # Meses fechados com registro: fev–jun (5). Observada = 349.300 ÷ 350.000 = 0,998.
        # Razões 1; 1; 1; 0,99; 1 → média 0,998; variância amostral = (4 × 0,002² + 0,008²) ÷ 4 = 0,00002.
        # Peso = 5 × 0,01 ÷ (0,05 + 0,00002) = 0,9996002; fator = 1 − 0,002 × 0,9996002 = 0,9980008.
        # Jul e ago (em aberto) com 70.000 ≥ 50% do esperado → completos. Set, out, nov: 70.000 × fator
        # cada (dez fora: 10 meses a partir de fevereiro). Projetado = 210.000 × 0,9980008 = 209.580,17.
        # Valor cheio = 210.000. Necessidade pela execução = 209.580,17 − saldo 10.000 = 199.580,17.
        # Necessidade do Resumo = 70.000 × (10 − 560.000/70.000) = 140.000.
        relatorio = _relatorio_real()
        linha = relatorio.linhas.set_index("ne_curta").loc[NE_REAL]
        self.assertEqual(linha["origem_fator"], ORIGEM_EXECUCAO)
        self.assertEqual(linha["meses_usados"], "Fev, Mar, Abr, Mai, Jun")
        self.assertAlmostEqual(linha["execucao_observada"], 0.998, places=9)
        self.assertAlmostEqual(linha["peso"], 0.05 / 0.05002, places=9)
        self.assertAlmostEqual(linha["fator"], 1 - 0.002 * 0.05 / 0.05002, places=9)
        self.assertAlmostEqual(linha["restante_em_aberto"], 0.0)
        self.assertAlmostEqual(linha["projetado_execucao"], 209580.17, places=2)
        self.assertAlmostEqual(linha["projetado_valor_cheio"], 210000.0)
        self.assertAlmostEqual(linha["necessidade_execucao"], 199580.17, places=2)
        self.assertAlmostEqual(linha["necessidade_contratual"], 140000.0)
        indice = relatorio.linhas.index[relatorio.linhas["ne_curta"] == NE_REAL][0]
        self.assertEqual([relatorio.mensal.at[indice, f"t{m}"] for m in (9, 10, 11)], [TIPO_PROJETADO] * 3)
        self.assertIsNone(relatorio.mensal.at[indice, "t12"])
        self.assertEqual(relatorio.qtd_sem_ne, 1)
        self.assertTrue(any("programa(s) sem NE" in a and "Resumo Consolidado" in a for a in relatorio.avisos))

    def test_arquivos_com_rotulos_de_bolsas(self):
        relatorio = _relatorio_real()
        contexto = ContextoProjecaoExecucao(2026, "11/09/2026", "abcd1234", "07/10/2026 12:00", FIXTURE_COMPETENCIA.name)
        self.assertTrue(gerar_pdf(relatorio, contexto).startswith(b"%PDF"))
        livro = load_workbook(BytesIO(gerar_xlsx(relatorio, contexto)))
        cabecalho = [c.value for c in livro["Resumo por NE"][1]]
        for coluna in ("Programa", "Processo", "Situação", "Despesa projetada (valor cadastrado)",
                       "Necessidade até dezembro (Resumo Consolidado)"):
            self.assertIn(coluna, cabecalho)
        self.assertNotIn("Fornecedor", cabecalho)
        linha = next(r for r in livro["Resumo por NE"].iter_rows(min_row=2, values_only=True) if r[0] == NE_REAL)
        self.assertEqual(linha[cabecalho.index("Programa")], "PROGRAMA REAL")
        self.assertEqual([c.value for c in livro["Grade mensal"][1]][1], "Programa")
        parametros = {r[0]: r[1] for r in livro["Parâmetros"].iter_rows(min_row=2, values_only=True)}
        self.assertIn("Despesa projetada pelo valor cadastrado (R$)", parametros)
        self.assertIn("início da execução", parametros["Regra"])

    def test_nota_descreve_a_ordem_do_inicio(self):
        # a nota impressa segue a ordem de `primeiro_mes_liquidado_por_ne`/`entrada_relatorio`:
        # cadastro → liquidação por competência → primeiro empenho → janeiro
        nota = BOLSAS_AUXILIOS.nota_regra
        posicoes = [nota.index(t) for t in ("informado no cadastro", "liquidação positiva por competência",
                                            "primeiro empenho", "janeiro")]
        self.assertEqual(posicoes, sorted(posicoes))


def _bolsa(ne: str, qtd: float, unitario: float, **extra: object) -> dict:
    return _programa(
        ne_curta=ne, valor_mensal=qtd * unitario, meses_no_ano=12, inicio_execucao_mes=1,
        qtd_efetiva=qtd, valor_unitario=unitario, **extra,
    )


class TestBolsasInteiras(unittest.TestCase):
    """Decisão de 16/10/2026: bolsa é paga por inteiro — o esperado do mês projetado vira número inteiro de
    bolsas (mais próximo, metade para cima). Valores calculados à mão."""

    def test_arredondar_em_unidades(self):
        self.assertEqual(arredondar_em_unidades(3320.0, 400.0), 3200.0)   # 8,3 → 8
        self.assertEqual(arredondar_em_unidades(3400.0, 400.0), 3600.0)   # 8,5 → 9 (metade para cima)
        self.assertEqual(arredondar_em_unidades(1049.0, 700.0), 700.0)    # 1,499 → 1
        self.assertEqual(arredondar_em_unidades(0.0, 700.0), 0.0)
        self.assertEqual(arredondar_em_unidades(123.45, None), 123.45)    # sem unidade: como veio
        self.assertEqual(arredondar_em_unidades(123.45, 0.0), 123.45)

    def test_unidades_por_ne(self):
        filtrado = pd.DataFrame([
            _bolsa("2026NE000001", 10, 400.0),
            _bolsa("2026NE000002", 2, 700.0), _bolsa("2026NE000002", 3, 700.0),   # mesmo unitário: soma
            _bolsa("2026NE000003", 2, 700.0), _bolsa("2026NE000003", 1, 500.0),   # unitários diferentes: fora
            {**_bolsa("2026NE000004", 2, 700.0), "valor_mensal": 1500.0},          # mensal excepcional: fora
            _bolsa("2026NE000005", 2, float("nan")),                                # unitário nulo: fora
            _programa(ne_curta=None, qtd_efetiva=1, valor_unitario=100.0),          # sem NE: fora
        ])
        copia = filtrado.copy()
        self.assertEqual(unidades_por_ne(filtrado), {"2026NE000001": 400.0, "2026NE000002": 700.0})
        pd.testing.assert_frame_equal(filtrado, copia)

    def _montar(self, unidades: bool):
        # NE 1: 10 bolsas de R$ 400 (R$ 4.000/mês), liquidado jan–jun = R$ 3.320 (8,3 bolsas), jul–ago = R$ 3.200.
        # NE 2: igual, mas com unitários diferentes no cadastro → sem unidade, sem arredondamento.
        filtrado = pd.DataFrame([
            _bolsa("2026NE000001", 10, 400.0),
            _bolsa("2026NE000002", 5, 400.0), _bolsa("2026NE000002", 5, 400.01),
        ])
        filtrado.loc[filtrado["ne_curta"] == "2026NE000002", "valor_mensal"] = [2000.0, 2000.05]
        meses = [202601, 202602, 202603, 202604, 202605, 202606, 202607, 202608]
        valores = [3320.0] * 6 + [3200.0] * 2
        competencia = pd.concat([
            pd.DataFrame({"ne_curta": ne, "ano_mes": meses, "valor": valores})
            for ne in ("2026NE000001", "2026NE000002")
        ], ignore_index=True)
        por_ne, sem_ne, custos = entrada_relatorio(filtrado, competencia, 2026)
        return montar_relatorio(
            por_ne, sem_ne, competencia, 2026, 9, custos=custos, rotulos=BOLSAS_AUXILIOS,
            unidades=unidades_por_ne(filtrado) if unidades else None,
        )

    def test_projecao_em_bolsas_inteiras(self):
        # NE 1: razões todas 0,83 → variância 0, peso 1, fator 0,83. Jul e ago (em aberto): 3.200 ≥ 50% de
        # 3.320 → completos. Set–dez: esperado 4.000 × 0,83 = 3.320 = 8,3 bolsas → 8 → R$ 3.200 cada.
        # Projetado = 4 × 3.200 = 12.800 (fracionado seria 13.280).
        linha = self._montar(True).linhas.set_index("ne_curta").loc["2026NE000001"]
        self.assertAlmostEqual(linha["fator"], 0.83, places=9)
        self.assertAlmostEqual(linha["projetado_execucao"], 12800.0)
        self.assertAlmostEqual(linha["projetado_valor_cheio"], 16000.0)  # valor cheio não arredonda
        self.assertEqual(linha["projetado_execucao"] % 400.0, 0.0)

    def test_sem_unidades_continua_fracionado(self):
        linha = self._montar(False).linhas.set_index("ne_curta").loc["2026NE000001"]
        self.assertAlmostEqual(linha["projetado_execucao"], 13280.0)  # 4 × 3.320

    def test_avisos_quantidade_fracionaria_e_ne_sem_unidade(self):
        relatorio = self._montar(True)
        fracionaria = [a for a in relatorio.avisos if "não é múltiplo do valor da bolsa" in a]
        self.assertEqual(len(fracionaria), 1)
        self.assertIn("2026NE000001", fracionaria[0])
        self.assertNotIn("2026NE000002", fracionaria[0])  # sem unidade: não dá para medir a quantidade
        sem_unidade = [a for a in relatorio.avisos if "sem arredondamento em bolsas inteiras" in a]
        self.assertEqual(len(sem_unidade), 1)
        self.assertIn("2026NE000002", sem_unidade[0])
        # NE sem unidade: projetada como antes (fator × custo, sem arredondar)
        linha = relatorio.linhas.set_index("ne_curta").loc["2026NE000002"]
        self.assertNotEqual(round(linha["projetado_execucao"], 2) % 400.0, 0.0)

    def test_sem_unidades_nao_ha_avisos_de_bolsa(self):
        avisos = self._montar(False).avisos
        self.assertFalse(any("múltiplo do valor da bolsa" in a or "bolsas inteiras" in a for a in avisos))

    def _livro(self, unidades: bool):
        contexto = ContextoProjecaoExecucao(2026, "13/10/2026", "abcd1234", "16/10/2026 12:00", "competencia.xlsx")
        return load_workbook(BytesIO(gerar_xlsx(self._montar(unidades), contexto)))

    def test_excel_traz_quantidade_de_bolsas(self):
        # NE 1 (R$ 400 a bolsa): jan–jun liquidado 3.320 = 8,3 bolsas; jul–ago 3.200 = 8; set–dez projetado 3.200 = 8
        livro = self._livro(True)
        aba = livro["Quantidade de bolsas"]
        cabecalho = [c.value for c in aba[1]]
        self.assertEqual(cabecalho[:3], ["NE", "Programa", "Valor da bolsa (R$)"])
        linhas = {r[0]: r for r in aba.iter_rows(min_row=2, max_row=2, values_only=True)}
        self.assertEqual(list(linhas), ["2026NE000001"])  # NE 2 não tem unidade: fora da aba
        ne1 = dict(zip(cabecalho, linhas["2026NE000001"]))
        self.assertEqual(ne1["Valor da bolsa (R$)"], 400.0)
        self.assertAlmostEqual(ne1["Jan"], 8.3)
        self.assertEqual([ne1[m] for m in ("Jul", "Ago", "Set", "Out", "Nov", "Dez")], [8.0] * 6)
        detalhe = livro["Detalhe mensal"]
        cab_detalhe = [c.value for c in detalhe[1]]
        for coluna in ("Esperado (bolsas)", "Realizado (bolsas)", "Projetado (bolsas)"):
            self.assertIn(coluna, cab_detalhe)
        set_ne1 = next(
            dict(zip(cab_detalhe, r)) for r in detalhe.iter_rows(min_row=2, values_only=True)
            if r[0] == "2026NE000001" and r[cab_detalhe.index("Mês")] == "Set"
        )
        self.assertEqual(set_ne1["Esperado (bolsas)"], 8.0)
        self.assertEqual(set_ne1["Projetado (bolsas)"], 8.0)
        self.assertIsNone(set_ne1["Realizado (bolsas)"])

    def test_excel_sem_unidades_nao_muda(self):
        livro = self._livro(False)
        self.assertNotIn("Quantidade de bolsas", livro.sheetnames)
        self.assertNotIn("Esperado (bolsas)", [c.value for c in livro["Detalhe mensal"][1]])

    def test_totais_nao_mudam_com_a_aba_de_quantidade(self):
        # a aba é só de exibição: o total projetado da NE 1 segue 12.800
        relatorio = self._montar(True)
        self.assertAlmostEqual(relatorio.linhas.set_index("ne_curta").loc["2026NE000001", "projetado_execucao"], 12800.0)
        self.assertEqual(relatorio.unidades, {"2026NE000001": 400.0})

    def test_valor_variavel_fica_fora_do_arredondamento(self):
        filtrado = pd.DataFrame([
            _bolsa("2026NE000001", 10, 400.0),
            {**_bolsa("2026NE000002", 1, 30000.0), "pago_valor_variavel": True},
            {**_bolsa("2026NE000003", 2, 700.0), "pago_valor_variavel": None},
        ])
        self.assertEqual(nes_valor_variavel(filtrado), {"2026NE000002"})
        self.assertEqual(unidades_por_ne(filtrado), {"2026NE000001": 400.0, "2026NE000003": 700.0})
        self.assertEqual(nes_valor_variavel(filtrado.drop(columns="pago_valor_variavel")), set())

    def test_ne_valor_variavel_projeta_pelo_fator_com_aviso_proprio(self):
        # mesma NE da `_montar`, agora marcada como valor variável: projetado = 4 × 3.320 = 13.280 (sem arredondar),
        # fora da aba de quantidade e dos avisos de unidade/fração; aviso próprio a lista
        filtrado = pd.DataFrame([{**_bolsa("2026NE000001", 10, 400.0), "pago_valor_variavel": True}])
        meses = [202601, 202602, 202603, 202604, 202605, 202606, 202607, 202608]
        competencia = pd.DataFrame({"ne_curta": "2026NE000001", "ano_mes": meses, "valor": [3320.0] * 6 + [3200.0] * 2})
        por_ne, sem_ne, custos = entrada_relatorio(filtrado, competencia, 2026)
        relatorio = montar_relatorio(
            por_ne, sem_ne, competencia, 2026, 9, custos=custos, rotulos=BOLSAS_AUXILIOS,
            unidades=unidades_por_ne(filtrado), nes_valor_variavel=nes_valor_variavel(filtrado),
        )
        linha = relatorio.linhas.set_index("ne_curta").loc["2026NE000001"]
        self.assertAlmostEqual(linha["projetado_execucao"], 13280.0)
        self.assertEqual(relatorio.unidades, {})
        proprio = [a for a in relatorio.avisos if "pagas em valor variável" in a]
        self.assertEqual(len(proprio), 1)
        self.assertIn("2026NE000001", proprio[0])
        self.assertFalse(any("múltiplo do valor da bolsa" in a or "bolsas inteiras" in a for a in relatorio.avisos))

    def test_formatar_bolsas(self):
        self.assertEqual(_formatar_bolsas(8.0), "8")
        self.assertEqual(_formatar_bolsas(8.3), "8,3")
        self.assertEqual(_formatar_bolsas(213.5), "213,5")
        self.assertEqual(_formatar_bolsas(0.0), "0")
        self.assertEqual(_formatar_bolsas(10.0), "10")
        self.assertEqual(_formatar_bolsas(None), "—")

    def _grade_do_pdf(self, unidades: bool) -> list[list]:
        """Linhas da última tabela do PDF (grade mensal), capturadas na montagem."""
        import src.relatorio_projecao_execucao as modulo

        capturadas = []
        original = modulo.Table

        def espiao(dados, *args, **kwargs):
            capturadas.append(dados)
            return original(dados, *args, **kwargs)

        contexto = ContextoProjecaoExecucao(2026, "13/10/2026", "abcd1234", "16/10/2026 12:00", "competencia.xlsx")
        with patch.object(modulo, "Table", espiao):
            self.assertTrue(gerar_pdf(self._montar(unidades), contexto).startswith(b"%PDF"))
        return capturadas[-1]

    def test_pdf_traz_linha_de_bolsas_abaixo_do_valor_em_reais(self):
        grade = self._grade_do_pdf(True)
        # NE 1 (400 por bolsa): título, linha em reais e linha de bolsas, nessa ordem
        indice = next(i for i, linha in enumerate(grade) if linha[0] == "2026NE000001")
        bolsas = grade[indice + 1]
        self.assertEqual(bolsas[0], "")
        self.assertIn("bolsas (R$ 400,00)", bolsas[1])
        # jan–jun 8,3; jul–ago 8; set–dez projetado 8
        self.assertEqual(bolsas[2:14], ["8,3"] * 6 + ["8"] * 6)
        # NE 2 (sem unidade): sem linha de bolsas — a linha seguinte já é o título da próxima NE ou o fim
        self.assertFalse(any("bolsas (R$" in str(l[1]) for l in grade if l is not grade[indice + 1]))

    def test_pdf_sem_unidades_nao_tem_linha_de_bolsas(self):
        grade = self._grade_do_pdf(False)
        self.assertFalse(any("bolsas (R$" in str(l[1]) for l in grade))
        self.assertIn("Linha 'bolsas'", LEGENDA_BOLSAS)

    def test_nota_da_regra_menciona_bolsa_inteira(self):
        self.assertIn("arredondado para o número inteiro de bolsas", BOLSAS_AUXILIOS.nota_regra)


if __name__ == "__main__":
    unittest.main()
