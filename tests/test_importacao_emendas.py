"""Importação versionada de Emendas: histórico inicial + atualizações 2026+."""

from __future__ import annotations

import shutil
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from src.emendas_parlamentares import ErroPoliticaImportacao
from src.importacao_emendas import (
    Manifesto,
    carregar_atual,
    importar,
    instalar_carga_inicial,
    sha256_arquivo,
)


CAMINHO_BASE = Path("tests/fixtures/emendas_acompanhamento_2026-08-28.xlsx")


def _relatorio_de_um_ano(
    destino: Path,
    *,
    ano: int,
    ptres: str = "000123",
    emenda: str = "202600000001",
    dotacao: float = 123.45,
) -> Path:
    referencia = load_workbook(CAMINHO_BASE, data_only=False)
    origem = referencia.active
    workbook = Workbook()
    aba = workbook.active
    aba.title = origem.title
    for linha in range(1, 4):
        for coluna in range(1, 11):
            aba.cell(linha, coluna).value = origem.cell(linha, coluna).value
    for intervalo in ("A1:A3", "B1:B3", "C1:D3", "E1:E3", "F1:F2"):
        aba.merge_cells(intervalo)

    valores = [
        "6",
        ptres,
        emenda,
        "PARLAMENTAR TESTE / Emenda 1",
        "4",
        ano,
        dotacao,
        None,
        None,
        None,
    ]
    for coluna, valor in enumerate(valores, start=1):
        aba.cell(4, coluna).value = valor
    workbook.save(destino)
    return destino


def _aplicar_atualizacao(resultado, caminho: Path, raw: Path, manifestos: Path) -> None:
    destino = raw / resultado.manifesto.arquivo
    shutil.copy2(caminho, destino)
    resultado.manifesto.salvar(manifestos)


@unittest.skipUnless(CAMINHO_BASE.exists(), f"Fixture ausente em {CAMINHO_BASE}")
class ImportacaoEmendasTests(unittest.TestCase):
    def test_carga_inicial_preserva_todo_o_historico(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            raiz = Path(tmp)
            raw = raiz / "raw"
            manifestos = raiz / "manifestos"
            hash_antes = sha256_arquivo(CAMINHO_BASE)

            resultado = instalar_carga_inicial(
                CAMINHO_BASE,
                raw,
                manifestos,
            )
            composto = carregar_atual(raw, manifestos)

            self.assertTrue(resultado.ok)
            self.assertEqual(
                sorted(int(ano) for ano in composto["ano"].unique()),
                [2016, 2019, 2020, 2021, 2022, 2023, 2024, 2025, 2026],
            )
            self.assertEqual(sha256_arquivo(CAMINHO_BASE), hash_antes)
            self.assertIn(resultado.manifesto.sha256[:8], resultado.manifesto.arquivo)
            self.assertTrue((raw / resultado.manifesto.arquivo).exists())

    def test_atualizacao_2026_substitui_so_2026_e_mantem_anos_anteriores(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            raiz = Path(tmp)
            raw = raiz / "raw"
            manifestos = raiz / "manifestos"
            instalar_carga_inicial(CAMINHO_BASE, raw, manifestos)
            atualizacao = _relatorio_de_um_ano(raiz / "atualizacao.xlsx", ano=2026)

            resultado = importar(atualizacao, manifestos, registrar=False)
            self.assertTrue(resultado.ok)
            _aplicar_atualizacao(resultado, atualizacao, raw, manifestos)
            composto = carregar_atual(raw, manifestos)

            self.assertIn(2016, composto["ano"].tolist())
            linhas_2026 = composto[composto["ano"].eq(2026)]
            self.assertEqual(len(linhas_2026), 1)
            self.assertEqual(linhas_2026.iloc[0]["ptres"], "000123")
            self.assertEqual(linhas_2026.iloc[0]["dotacao_atualizada"], 123.45)

    def test_upload_com_ano_anterior_a_2026_e_bloqueado_por_inteiro(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            raiz = Path(tmp)
            raw = raiz / "raw"
            manifestos = raiz / "manifestos"
            inicial = instalar_carga_inicial(CAMINHO_BASE, raw, manifestos)
            atualizacao = _relatorio_de_um_ano(
                raiz / "historico.xlsx",
                ano=2025,
                emenda="202500000001",
            )

            resultado = importar(atualizacao, manifestos, registrar=True)

            self.assertFalse(resultado.ok)
            self.assertTrue(any("2025" in erro for erro in resultado.validacao.erros))
            self.assertEqual(Manifesto.atual(manifestos).sha256, inicial.manifesto.sha256)
            composto = carregar_atual(raw, manifestos)
            self.assertGreater(len(composto[composto["ano"].eq(2025)]), 0)

    def test_carga_inicial_nao_pode_ser_registrada_duas_vezes(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            raiz = Path(tmp)
            instalar_carga_inicial(CAMINHO_BASE, raiz / "raw", raiz / "manifestos")
            with self.assertRaisesRegex(ErroPoliticaImportacao, "já foi registrada"):
                instalar_carga_inicial(
                    CAMINHO_BASE,
                    raiz / "raw",
                    raiz / "manifestos",
                )

    def test_mesmo_nome_com_conteudo_diferente_recebe_hash_no_nome(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            raiz = Path(tmp)
            raw = raiz / "raw"
            manifestos = raiz / "manifestos"
            instalar_carga_inicial(CAMINHO_BASE, raw, manifestos)
            pasta_a = raiz / "a"
            pasta_b = raiz / "b"
            pasta_a.mkdir()
            pasta_b.mkdir()
            arquivo_a = _relatorio_de_um_ano(pasta_a / "relatorio.xlsx", ano=2026)
            arquivo_b = _relatorio_de_um_ano(
                pasta_b / "relatorio.xlsx",
                ano=2027,
                emenda="202700000001",
                dotacao=999.0,
            )

            resultado_a = importar(arquivo_a, manifestos, registrar=False)
            resultado_b = importar(arquivo_b, manifestos, registrar=False)

            self.assertNotEqual(resultado_a.manifesto.sha256, resultado_b.manifesto.sha256)
            self.assertNotEqual(resultado_a.manifesto.arquivo, resultado_b.manifesto.arquivo)
            self.assertTrue(resultado_a.manifesto.arquivo.endswith(".xlsx"))


if __name__ == "__main__":
    unittest.main()
