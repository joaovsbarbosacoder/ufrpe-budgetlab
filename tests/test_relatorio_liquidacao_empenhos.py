"""Testes do relatório (PDF/Excel) de liquidação mensal das NEs marcadas na Consulta de
Empenhos — `src/relatorio_liquidacao_empenhos.py`. Séries sintéticas montadas à mão, no mesmo
formato (`ne_ccor`, `ano_mes`, `valor`) de `liquidacao_competencia.liquidado_por_ne_e_mes` e
`tesouro_execucao_mensal.liquidado_por_ne_e_mes`."""

from __future__ import annotations

import math
import unittest
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from src.relatorio_liquidacao_empenhos import (
    BASE_COMPETENCIA,
    BASE_LANCAMENTO,
    BASE_SEM_DADO,
    ContextoRelatorioLiquidacao,
    consolidar_por_dimensao,
    gerar_pdf,
    gerar_xlsx,
    montar_relatorio,
)
from src.tesouro_execucao_mensal import liquidado_por_ne_e_mes

NE_COMP = "153165152662025NE000010"   # tem competência (2025 e 2026 — restos a pagar)
NE_LANC = "153165152662026NE000056"   # só na Execução Mensal -> fallback por lançamento
NE_VAZIA = "153165152662026NE000099"  # nenhuma liquidação
NE_FORA = "153165152662026NE000777"   # não marcada — não pode entrar


def _nes() -> pd.DataFrame:
    return pd.DataFrame(
        {
            "ne_ccor": [NE_LANC, NE_COMP, NE_VAZIA],
            "ano": [2026, 2025, 2026],
            "ne_favorecido": ["EMPRESA B & CIA <LTDA>", "EMPRESA A", None],
            "ne_descricao": ["BOLSAS <PIBIC> & AUXÍLIOS", "LIMPEZA E CONSERVAÇÃO", None],
        }
    )


def _lancamento() -> pd.DataFrame:
    return pd.DataFrame(
        {
            "ne_ccor": [NE_COMP, NE_COMP, NE_LANC, NE_LANC, NE_LANC, NE_FORA],
            "ano_mes": [202503, 202602, 202601, 202602, 202603, 202601],
            "valor": [1_000.0, 500.0, 300.0, 0.0, -50.0, 9_999.0],
        }
    )


def _competencia() -> pd.DataFrame:
    return pd.DataFrame(
        {
            "ne_ccor": [NE_COMP, NE_COMP, NE_COMP, NE_FORA],
            "ano_mes": [202502, 202512, 202601, 202601],
            "valor": [900.0, 400.0, -100.0, 1.0],
        }
    )


def _contexto(**kwargs) -> ContextoRelatorioLiquidacao:
    padrao = dict(
        data_extracao="28/09/2026",
        hash_manifesto="abcdef12",
        data_emissao="30/09/2026 10:00",
        origem_competencia="Liquidação por Competência.xlsx · modificado em 29/09/2026",
    )
    padrao.update(kwargs)
    return ContextoRelatorioLiquidacao(**padrao)


class MontarRelatorioTest(unittest.TestCase):
    def setUp(self) -> None:
        self.rel = montar_relatorio(_nes(), _lancamento(), _competencia())

    def test_base_escolhida_por_ne(self) -> None:
        bases = dict(zip(self.rel.resumo["ne_ccor"], self.rel.resumo["base"]))
        self.assertEqual(bases, {NE_COMP: BASE_COMPETENCIA, NE_LANC: BASE_LANCAMENTO, NE_VAZIA: BASE_SEM_DADO})
        self.assertEqual(self.rel.nes_por_lancamento, [NE_LANC])
        self.assertEqual(self.rel.nes_sem_dado, [NE_VAZIA])

    def test_ne_nao_marcada_fica_de_fora(self) -> None:
        self.assertNotIn(NE_FORA, set(self.rel.serie["ne_ccor"]))
        self.assertNotIn(NE_FORA, set(self.rel.mensal["ne_ccor"]))

    def test_serie_da_ne_com_competencia_nao_mistura_lancamento(self) -> None:
        serie = self.rel.serie[self.rel.serie["ne_ccor"] == NE_COMP]
        self.assertEqual(serie["ano_mes"].tolist(), [202502, 202512, 202601])
        self.assertEqual(set(serie["base"]), {BASE_COMPETENCIA})

    def test_linha_por_ne_e_ano_com_meses_vazios_nulos(self) -> None:
        mensal = self.rel.mensal
        comp = mensal[mensal["ne_ccor"] == NE_COMP].set_index("ano")
        self.assertEqual(comp.index.tolist(), [2025, 2026])
        self.assertEqual(comp.loc[2025, 2], 900.0)
        self.assertEqual(comp.loc[2025, 12], 400.0)
        self.assertTrue(math.isnan(comp.loc[2025, 1]))  # sem liquidação: nulo, não zero
        self.assertEqual(comp.loc[2025, "total"], 1_300.0)
        self.assertEqual(comp.loc[2026, 1], -100.0)  # estorno com sinal
        self.assertEqual(comp.loc[2026, "ano_ne"], 2025)

    def test_zero_real_preservado_e_estorno_no_fallback(self) -> None:
        linha = self.rel.mensal[self.rel.mensal["ne_ccor"] == NE_LANC].iloc[0]
        self.assertEqual(linha[1], 300.0)
        self.assertEqual(linha[2], 0.0)
        self.assertEqual(linha[3], -50.0)
        self.assertTrue(math.isnan(linha[4]))
        self.assertEqual(linha["total"], 250.0)
        self.assertEqual(linha["base"], BASE_LANCAMENTO)

    def test_ne_sem_dado_aparece_com_valores_nulos(self) -> None:
        linha = self.rel.mensal[self.rel.mensal["ne_ccor"] == NE_VAZIA]
        self.assertEqual(len(linha), 1)
        self.assertTrue(pd.isna(linha.iloc[0]["ano"]))
        self.assertTrue(pd.isna(linha.iloc[0]["total"]))
        self.assertEqual(linha.iloc[0]["base"], BASE_SEM_DADO)

    def test_reconciliacao(self) -> None:
        resumo = self.rel.resumo.set_index("ne_ccor")
        self.assertEqual(resumo.loc[NE_COMP, "total_base"], 1_200.0)
        self.assertEqual(resumo.loc[NE_COMP, "liquidado_execucao_mensal"], 1_500.0)
        self.assertEqual(resumo.loc[NE_COMP, "diferenca"], -300.0)
        self.assertEqual(resumo.loc[NE_LANC, "diferenca"], 0.0)
        self.assertTrue(pd.isna(resumo.loc[NE_VAZIA, "total_base"]))
        self.assertTrue(pd.isna(resumo.loc[NE_VAZIA, "diferenca"]))

    def test_competencia_divergente_do_liquidado(self) -> None:
        # NE_COMP: 1.200 por competência x 1.500 na Execução Mensal; NE_LANC por lançamento
        # (diferença 0 por construção) e NE_VAZIA (nula) não entram.
        self.assertEqual(self.rel.nes_competencia_divergente, [NE_COMP])
        lanc = _lancamento()
        lanc.loc[lanc["ne_ccor"] == NE_COMP, "valor"] = [900.0, 300.0]
        rel = montar_relatorio(_nes(), lanc, _competencia())
        self.assertEqual(rel.nes_competencia_divergente, [])

    def test_ordem_alfabetica_pela_descricao_sem_acento_nem_caixa(self) -> None:
        # "BOLSAS..." < "LIMPEZA..." ; descrição nula (NE_VAZIA) por último
        self.assertEqual(self.rel.resumo["ne_ccor"].tolist(), [NE_LANC, NE_COMP, NE_VAZIA])
        self.assertEqual(
            self.rel.mensal["ne_ccor"].tolist(), [NE_LANC, NE_COMP, NE_COMP, NE_VAZIA]
        )
        nes = _nes()
        nes["ne_descricao"] = ["ágil", "Abacaxi", "acerola"]
        rel = montar_relatorio(nes, _lancamento(), _competencia())
        # "abacaxi" < "acerola" < "agil" (acento e caixa ignorados)
        self.assertEqual(rel.resumo["ne_descricao"].tolist(), ["Abacaxi", "acerola", "ágil"])

    def test_ordem_ignora_pontuacao_e_espacos_repetidos(self) -> None:
        # casos reais da base: espaço duplo e hífen não podem decidir a ordem antes das letras,
        # e "N°"/"Nº" são a mesma coisa.
        nes = _nes()
        nes["ne_descricao"] = [
            "DIARIAS NO PAIS  SERVIDOR AGROECOLOGIA",  # NE_LANC — espaço duplo
            "DIARIAS NO PAIS - DESPACHO Nº 38304",      # NE_COMP
            "DIARIAS NO PAIS - OF. N° 76",              # NE_VAZIA
        ]
        rel = montar_relatorio(nes, _lancamento(), _competencia())
        # despacho < of < servidor
        self.assertEqual(rel.resumo["ne_ccor"].tolist(), [NE_COMP, NE_VAZIA, NE_LANC])

        from src.relatorio_liquidacao_empenhos import _chave_alfabetica

        self.assertEqual(_chave_alfabetica("OF. Nº 17"), _chave_alfabetica("of  n° 17"))
        self.assertEqual(_chave_alfabetica("Ação - Educação"), "acao educacao")

    def test_desempate_por_favorecido_e_ne(self) -> None:
        nes = _nes()
        nes["ne_descricao"] = "MESMA"
        nes["ne_favorecido"] = ["B", "A", "A"]
        rel = montar_relatorio(nes, _lancamento(), _competencia())
        self.assertEqual(rel.resumo["ne_ccor"].tolist(), [NE_COMP, NE_VAZIA, NE_LANC])

    def test_por_exercicio_soma_por_ano_e_expoe_mistura_de_bases(self) -> None:
        ex = self.rel.por_exercicio.set_index("ano")
        self.assertEqual(ex.index.tolist(), [2025, 2026])  # NE "Sem dado" não entra
        # 2025: só NE_COMP (competência) — Fev 900, Dez 400
        self.assertEqual((ex.loc[2025, "nes_competencia"], ex.loc[2025, "nes_lancamento"]), (1, 0))
        self.assertEqual(ex.loc[2025, 2], 900.0)
        self.assertTrue(math.isnan(ex.loc[2025, 1]))  # mês sem valor: nulo, não zero
        self.assertEqual(ex.loc[2025, "total"], 1_300.0)
        # 2026: NE_COMP (competência, Jan −100) + NE_LANC (lançamento, Jan 300, Fev 0, Mar −50)
        self.assertEqual((ex.loc[2026, "nes_competencia"], ex.loc[2026, "nes_lancamento"]), (1, 1))
        self.assertEqual(ex.loc[2026, 1], 200.0)
        self.assertEqual(ex.loc[2026, 2], 0.0)
        self.assertEqual(ex.loc[2026, 3], -50.0)
        self.assertEqual(ex.loc[2026, "total"], 150.0)
        # reconciliação: cada ano = soma das linhas NE × ano daquele ano
        for ano, total in ex["total"].items():
            self.assertAlmostEqual(total, self.rel.mensal.loc[self.rel.mensal["ano"] == ano, "total"].sum())

    def test_sem_aviso_de_bases_quando_todos_os_anos_usam_a_mesma(self) -> None:
        from src.relatorio_liquidacao_empenhos import _aviso_bases_por_exercicio

        so_comp = montar_relatorio(_nes().iloc[[1]], _lancamento(), _competencia())  # NE_COMP: 2025 e 2026
        self.assertIsNone(_aviso_bases_por_exercicio(so_comp.por_exercicio))
        self.assertIsNotNone(_aviso_bases_por_exercicio(self.rel.por_exercicio))

    def test_por_exercicio_vazio_quando_so_ha_ne_sem_dado(self) -> None:
        rel = montar_relatorio(_nes().iloc[[2]], _lancamento(), _competencia())
        self.assertTrue(rel.por_exercicio.empty)

    def test_modo_somente_lancamento_ignora_competencia(self) -> None:
        from src.relatorio_liquidacao_empenhos import MODO_SOMENTE_LANCAMENTO

        rel = montar_relatorio(_nes(), _lancamento(), _competencia(), MODO_SOMENTE_LANCAMENTO)
        self.assertEqual(rel.modo, MODO_SOMENTE_LANCAMENTO)
        bases = dict(zip(rel.resumo["ne_ccor"], rel.resumo["base"]))
        self.assertEqual(bases, {NE_COMP: BASE_LANCAMENTO, NE_LANC: BASE_LANCAMENTO, NE_VAZIA: BASE_SEM_DADO})
        # NE_COMP agora pela Execução Mensal: 2025-03 = 1.000, 2026-02 = 500
        comp = rel.mensal[rel.mensal["ne_ccor"] == NE_COMP].set_index("ano")
        self.assertEqual(comp.loc[2025, 3], 1_000.0)
        self.assertEqual(comp.loc[2026, 2], 500.0)
        self.assertEqual(rel.nes_competencia_divergente, [])
        self.assertTrue((rel.resumo["diferenca"].dropna() == 0).all())
        # comparativo: todos os anos na mesma base -> sem aviso de mistura
        self.assertIsNone(rel.aviso_bases_por_exercicio)
        self.assertEqual(rel.por_exercicio["nes_competencia"].sum(), 0)

    def test_modo_desconhecido_e_rejeitado(self) -> None:
        with self.assertRaises(ValueError):
            montar_relatorio(_nes(), _lancamento(), _competencia(), "qualquer")

    def test_sem_base_de_competencia_tudo_por_lancamento(self) -> None:
        rel = montar_relatorio(_nes(), _lancamento(), None)
        bases = dict(zip(rel.resumo["ne_ccor"], rel.resumo["base"]))
        self.assertEqual(bases[NE_COMP], BASE_LANCAMENTO)
        self.assertEqual(bases[NE_LANC], BASE_LANCAMENTO)
        self.assertEqual(bases[NE_VAZIA], BASE_SEM_DADO)

    def test_so_ne_sem_dado(self) -> None:
        rel = montar_relatorio(_nes().iloc[[2]], _lancamento(), _competencia())
        self.assertEqual(len(rel.mensal), 1)
        self.assertTrue(rel.serie.empty)

    def test_nao_altera_entradas(self) -> None:
        nes, lanc, comp = _nes(), _lancamento(), _competencia()
        copias = (nes.copy(), lanc.copy(), comp.copy())
        montar_relatorio(nes, lanc, comp)
        for original, copia in zip((nes, lanc, comp), copias):
            pd.testing.assert_frame_equal(original, copia)


class LiquidadoLancamentoTest(unittest.TestCase):
    def test_nao_preenche_zero_e_ignora_linha_de_empenho(self) -> None:
        bruto = pd.DataFrame(
            {
                "ne_ccor": ["A", "A", "A", "A", "B"],
                "ano_mes": [202601, 202601, 202602, 202603, 202601],
                "tipo_linha": ["item_execucao", "item_execucao", "item_execucao", "empenho", "item_execucao"],
                "liquidada": [10.0, -3.0, None, 999.0, 0.0],
            }
        )
        resultado = liquidado_por_ne_e_mes(bruto)
        self.assertEqual(
            list(resultado.itertuples(index=False, name=None)),
            [("A", 202601, 7.0), ("B", 202601, 0.0)],
        )


class FixturesReaisTest(unittest.TestCase):
    """Reconciliação com as fixtures congeladas: nenhuma NE marcada some, e o total da série
    de cada NE bate com a base de origem escolhida para ela."""

    @classmethod
    def setUpClass(cls) -> None:
        from src.liquidacao_competencia import ler_liquidacao_competencia
        from src.liquidacao_competencia import liquidado_por_ne_e_mes as competencia_por_ne_e_mes
        from src.tesouro_execucao_mensal import agregar_por_ne, ler_execucao_mensal

        execucao = ler_execucao_mensal("tests/fixtures/execucao_mensal_2026-09-22.xlsx")
        cls.por_ne = agregar_por_ne(execucao).reset_index()
        cls.lancamento = liquidado_por_ne_e_mes(execucao)
        cls.competencia_bruta = ler_liquidacao_competencia("tests/fixtures/liquidacao_competencia_2026-09-11.xlsx")
        cls.competencia = competencia_por_ne_e_mes(cls.competencia_bruta)

    def test_totais_reconciliam_com_as_bases(self) -> None:
        com_comp = set(self.competencia["ne_ccor"])
        nes = self.por_ne[["ne_ccor", "ano", "ne_favorecido", "ne_descricao"]].head(200)
        rel = montar_relatorio(nes, self.lancamento, self.competencia)

        self.assertEqual(set(rel.resumo["ne_ccor"]), set(nes["ne_ccor"]))
        liquidado_exec = self.por_ne.set_index("ne_ccor")["liquidada"]
        comp_total = self.competencia_bruta.groupby("ne_ccor")["valor"].sum()
        for linha in rel.resumo.itertuples():
            with self.subTest(ne=linha.ne_ccor):
                self.assertEqual(linha.base == BASE_COMPETENCIA, linha.ne_ccor in com_comp)
                esperado_exec = liquidado_exec.get(linha.ne_ccor)
                if pd.isna(esperado_exec):
                    self.assertTrue(pd.isna(linha.liquidado_execucao_mensal))
                else:
                    self.assertAlmostEqual(linha.liquidado_execucao_mensal, esperado_exec, places=2)
                if linha.base == BASE_COMPETENCIA:
                    self.assertAlmostEqual(linha.total_base, comp_total[linha.ne_ccor], places=2)
        self.assertAlmostEqual(rel.por_exercicio["total"].sum(), rel.serie["valor"].sum(), places=2)
        # soma da tabela mensal = soma da série longa = soma dos totais do resumo
        self.assertAlmostEqual(rel.mensal["total"].sum(), rel.serie["valor"].sum(), places=2)
        self.assertAlmostEqual(rel.resumo["total_base"].sum(), rel.serie["valor"].sum(), places=2)


class ConsolidarPorDimensaoTest(unittest.TestCase):
    def test_agrupa_ordena_e_preserva_nulos(self) -> None:
        grupo = pd.DataFrame(
            {
                "ne_ccor": ["A", "B", "C", "D"],
                "elemento_cod": ["39", "39", "18", None],
                "elemento_desc": ["OUTROS SERV", "OUTROS SERV", "AUX", None],
                "empenhada": [100.0, 50.0, 500.0, 10.0],
                "liquidada": [40.0, None, None, 5.0],
                "paga": [30.0, None, None, None],
            }
        )
        tabela = consolidar_por_dimensao(grupo, "elemento_cod", "elemento_desc")
        self.assertEqual(tabela["nome"].tolist(), ["AUX", "OUTROS SERV", "(não informado)"])
        aux, outros, nao_inf = (tabela.iloc[i] for i in range(3))
        self.assertTrue(pd.isna(aux["liq"]))  # só nulos soma nulo
        self.assertEqual(aux["saldo"], 500.0)  # liquidado nulo conta como 0 no saldo
        self.assertEqual((outros["qtd"], outros["emp"], outros["liq"], outros["pag"]), (2, 150.0, 40.0, 30.0))
        self.assertEqual(outros["codigo"], "39")
        self.assertEqual(nao_inf["qtd"], 1)  # código nulo não é descartado
        self.assertEqual(tabela["qtd"].sum(), 4)


class GerarArquivosTest(unittest.TestCase):
    def setUp(self) -> None:
        self.rel = montar_relatorio(_nes(), _lancamento(), _competencia())

    def test_xlsx_abas_codigos_como_texto_e_nulos_vazios(self) -> None:
        livro = load_workbook(BytesIO(gerar_xlsx(self.rel, _contexto())))
        self.assertEqual(
            livro.sheetnames, ["Liquidação mensal", "Por exercício", "Resumo por NE", "Série", "Parâmetros"]
        )
        aba_ex = livro["Por exercício"]
        cab_ex = [c.value for c in aba_ex[1]]
        self.assertEqual(cab_ex[:4], ["Ano", "NEs por competência", "NEs por data de liquidação", "Jan"])
        linhas_ex = [dict(zip(cab_ex, [c.value for c in l])) for l in aba_ex.iter_rows(min_row=2)]
        self.assertEqual([l["Ano"] for l in linhas_ex], [2025, 2026])
        self.assertIsNone(linhas_ex[0]["Jan"])
        self.assertEqual(linhas_ex[1]["Jan"], 200.0)
        self.assertEqual(linhas_ex[1]["Total"], 150.0)
        aba = livro["Liquidação mensal"]
        cabecalho = [c.value for c in aba[1]]
        self.assertEqual(
            cabecalho[:7], ["NE", "NE (código completo)", "Exercício da NE", "Favorecido", "Descrição", "Base", "Ano"]
        )
        linhas = [dict(zip(cabecalho, [c.value for c in linha])) for linha in aba.iter_rows(min_row=2)]
        comp_2025 = next(l for l in linhas if l["NE (código completo)"] == NE_COMP and l["Ano"] == 2025)
        self.assertIsInstance(comp_2025["NE (código completo)"], str)
        self.assertEqual(comp_2025["Base"], BASE_COMPETENCIA)
        self.assertEqual(comp_2025["Favorecido"], "EMPRESA A")
        self.assertEqual(comp_2025["Descrição"], "LIMPEZA E CONSERVAÇÃO")
        vazia = next(l for l in linhas if l["NE (código completo)"] == NE_VAZIA)
        self.assertIsNone(vazia["Descrição"])  # nulo continua nulo
        resumo = livro["Resumo por NE"]
        cab_resumo = [c.value for c in resumo[1]]
        self.assertIn("Descrição", cab_resumo)
        descricoes = {l[1].value: l[cab_resumo.index("Descrição")].value for l in resumo.iter_rows(min_row=2)}
        self.assertEqual(descricoes[NE_LANC], "BOLSAS <PIBIC> & AUXÍLIOS")
        self.assertEqual(comp_2025["Fev"], 900.0)
        self.assertIsNone(comp_2025["Jan"])
        lanc = next(l for l in linhas if l["NE (código completo)"] == NE_LANC)
        self.assertEqual(lanc["Fev"], 0)
        self.assertEqual(lanc["Base"], BASE_LANCAMENTO)

        pares = [(linha[0].value, linha[1].value) for linha in livro["Parâmetros"].iter_rows(min_row=2)]
        avisos = [valor for chave, valor in pares if chave == "Avisos"]
        self.assertEqual(len(avisos), 4)
        self.assertIn("NÃO compara a mesma base", avisos[3])
        self.assertIn("2025 só por competência", avisos[3])
        self.assertIn("2026 misto (1 por competência, 1 por data de liquidação)", avisos[3])
        self.assertIn("DATA DE LIQUIDAÇÃO", avisos[0])
        self.assertIn("2026NE000056", avisos[0])
        self.assertIn("difere do total liquidado", avisos[1])
        self.assertIn("2025NE000010", avisos[1])
        self.assertIn("2026NE000099", avisos[2])
        parametros = dict(pares)
        self.assertEqual(parametros["Base do relatório"], "Competência quando houver")
        self.assertEqual(parametros["Por competência"], "1")
        self.assertEqual(parametros["Por data de liquidação"], "1")

    def test_xlsx_aba_consolidacao_com_codigos_texto_e_nulos_vazios(self) -> None:
        grupo = pd.DataFrame(
            {
                "ne_ccor": ["A", "B", "C"],
                "acao_cod": ["0181", "0181", "20RK"],
                "acao_desc": ["APOSENTADORIAS", "APOSENTADORIAS", "FUNCIONAMENTO"],
                "gnd_cod": ["1", "1", "3"],
                "gnd_desc": ["PESSOAL", "PESSOAL", "OUTRAS CORRENTES"],
                "empenhada": [100.0, 50.0, 10.0],
                "liquidada": [60.0, 40.0, None],
                "paga": [60.0, 30.0, None],
            }
        )
        consolidacao = [
            ("Ação de Governo", consolidar_por_dimensao(grupo, "acao_cod", "acao_desc")),
            ("Grupo de Despesa", consolidar_por_dimensao(grupo, "gnd_cod", "gnd_desc")),
        ]
        livro = load_workbook(BytesIO(gerar_xlsx(self.rel, _contexto(), consolidacao)))
        self.assertIn("Consolidação", livro.sheetnames)
        aba = livro["Consolidação"]
        cabecalho = [c.value for c in aba[1]]
        self.assertEqual(
            cabecalho,
            ["Dimensão", "Grupo", "Código", "NEs", "Empenhado", "Liquidado (data de liquidação)", "Pago", "Saldo"],
        )
        linhas = [dict(zip(cabecalho, [c.value for c in l])) for l in aba.iter_rows(min_row=2)]
        self.assertEqual(len(linhas), 4)
        aposentadorias = linhas[0]
        self.assertEqual(aposentadorias["Código"], "0181")  # zero à esquerda preservado
        self.assertEqual(
            (aposentadorias["NEs"], aposentadorias["Empenhado"], aposentadorias["Liquidado (data de liquidação)"]),
            (2, 150.0, 100.0),
        )
        funcionamento = linhas[1]
        self.assertIsNone(funcionamento["Liquidado (data de liquidação)"])  # nulo, não zero
        self.assertEqual(funcionamento["Saldo"], 10.0)
        self.assertEqual({l["Dimensão"] for l in linhas}, {"Ação de Governo", "Grupo de Despesa"})
        parametros = {
            l[0].value: l[1].value for l in livro["Parâmetros"].iter_rows(min_row=2) if l[0].value != "Avisos"
        }
        self.assertIn("não por competência", parametros["Consolidação"])

    def test_xlsx_sem_consolidacao_nao_cria_a_aba(self) -> None:
        livro = load_workbook(BytesIO(gerar_xlsx(self.rel, _contexto())))
        self.assertNotIn("Consolidação", livro.sheetnames)

    def test_xlsx_e_pdf_no_modo_somente_lancamento(self) -> None:
        from src.relatorio_liquidacao_empenhos import MODO_SOMENTE_LANCAMENTO, titulo_relatorio

        rel = montar_relatorio(_nes(), _lancamento(), _competencia(), MODO_SOMENTE_LANCAMENTO)
        livro = load_workbook(BytesIO(gerar_xlsx(rel, _contexto())))
        pares = [(l[0].value, l[1].value) for l in livro["Parâmetros"].iter_rows(min_row=2)]
        parametros = dict(pares)
        self.assertEqual(parametros["Base do relatório"], MODO_SOMENTE_LANCAMENTO)
        self.assertIn("não usada", parametros["Liquidação por Competência"])
        avisos = [v for k, v in pares if k == "Avisos"]
        # só o aviso de NE sem dado — nada sobre competência
        self.assertEqual(len(avisos), 1)
        self.assertIn("sem liquidação na Execução Mensal", avisos[0])
        self.assertTrue(gerar_pdf(rel, _contexto()).startswith(b"%PDF"))
        self.assertIn("DATA DE LIQUIDAÇÃO", titulo_relatorio(MODO_SOMENTE_LANCAMENTO))

    def test_xlsx_aviso_geral_sem_competencia(self) -> None:
        rel = montar_relatorio(_nes(), _lancamento(), None)
        livro = load_workbook(BytesIO(gerar_xlsx(rel, _contexto(origem_competencia=None))))
        avisos = [l[1].value for l in livro["Parâmetros"].iter_rows(min_row=2) if l[0].value == "Avisos"]
        self.assertTrue(any("TODAS as NEs" in a for a in avisos))

    def test_pdf_gera_documento(self) -> None:
        grupo = pd.DataFrame(
            {"ne_ccor": [NE_COMP], "gnd_cod": ["3"], "gnd_desc": ["OUTRAS DESPESAS <CORRENTES>"],
             "empenhada": [10.0], "liquidada": [5.0], "paga": [None]}
        )
        consolidacao = [("Grupo de Despesa", consolidar_por_dimensao(grupo, "gnd_cod", "gnd_desc"))]
        conteudo = gerar_pdf(self.rel, _contexto(), consolidacao)
        self.assertTrue(conteudo.startswith(b"%PDF"))
        self.assertGreater(len(conteudo), 1_000)

    def test_nosplit_cobre_faixa_e_linhas_de_cada_ne(self) -> None:
        from src.relatorio_liquidacao_empenhos import _comandos_nosplit_por_ne

        # cabeçalho na linha 0; NE A: faixa 1 + linhas 2-3; NE B: faixa 4 + linha 5; NE C: 6-7
        self.assertEqual(
            _comandos_nosplit_por_ne([1, 4, 6], 8),
            [("NOSPLIT", (0, 1), (-1, 3)), ("NOSPLIT", (0, 4), (-1, 5)), ("NOSPLIT", (0, 6), (-1, 7))],
        )
        self.assertEqual(_comandos_nosplit_por_ne([], 1), [])

    def test_pdf_com_muitas_nes_quebra_paginas_sem_erro(self) -> None:
        # 60 NEs com 2 anos cada — força várias quebras de página com NOSPLIT em todos os blocos
        nes = pd.DataFrame(
            {
                "ne_ccor": [f"153165152392026NE{n:06d}" for n in range(60)],
                "ano": 2026,
                "ne_favorecido": "EMPRESA",
                "ne_descricao": [f"DESPESA {n:02d}" for n in range(60)],
            }
        )
        lancamento = pd.DataFrame(
            {
                "ne_ccor": [ne for ne in nes["ne_ccor"] for _ in range(2)],
                "ano_mes": [202512, 202601] * 60,
                "valor": [10.0, 20.0] * 60,
            }
        )
        rel = montar_relatorio(nes, lancamento, None)
        self.assertEqual(len(rel.mensal), 120)
        self.assertTrue(gerar_pdf(rel, _contexto(origem_competencia=None)).startswith(b"%PDF"))

    def test_pdf_sem_competencia_e_so_sem_dado(self) -> None:
        rel = montar_relatorio(_nes().iloc[[2]], _lancamento(), None)
        self.assertTrue(gerar_pdf(rel, _contexto(origem_competencia=None)).startswith(b"%PDF"))
        load_workbook(BytesIO(gerar_xlsx(rel, _contexto(origem_competencia=None))))


if __name__ == "__main__":
    unittest.main()
