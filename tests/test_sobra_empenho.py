"""Testes do relatório "Sobra prevista de empenho" (`src/sobra_empenho.py`). Valores calculados à mão."""

from __future__ import annotations

import unittest
from datetime import date
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from src.relatorio_projecao_execucao import (
    CONTRATOS,
    BOLSAS_AUXILIOS,
    ORIGEM_EXECUCAO,
    ORIGEM_SEM_HISTORICO,
    TIPO_PROJETADO,
    RelatorioProjecaoExecucao,
)
from src.sobra_empenho import (
    COL_SOBRA,
    FAIXA_1,
    FAIXA_2,
    FAIXA_3,
    SEM_SOBRA,
    ContextoSobra,
    faixa_da_sobra,
    gerar_xlsx,
    FAIXAS_NA_ANULACAO,
    resumo_da_sobra,
    sobra_elegivel_por_ne,
    sobra_por_ne,
    sugestao_anulacao_por_ne,
)

REF = date(2026, 10, 15)
CONTEXTO = ContextoSobra(2026, "extração 15/10/2026 · sha abcd1234", "Liquidação por Competência.xlsx", "16/10/2026 10:00")


def _linha(ne, *, saldo, projetado, cheio, status="ATIVO", origem=ORIGEM_EXECUCAO, fator=0.9, teto=False,
           vigencia="2027-12-31", ajustada=False, nome="FORNECEDOR", empenhado=1000.0) -> dict:
    return {
        "ne_curta": ne, "fornecedor": nome, "contrato_numero": "1/2026", "status_contrato": status,
        "vigencia_fim": pd.Timestamp(vigencia), "despesa_mensal": 100.0, "valor_empenhado": empenhado, "saldo": saldo,
        "origem_fator": origem, "fator": fator, "meses_usados": "Fev, Mar, Abr", "limitada_ao_maior_mes": teto,
        "projetado_execucao": projetado, "projetado_valor_cheio": cheio, "necessidade_execucao": 0.0,
        "ajustada_manualmente": ajustada,
    }


def _relatorio(linhas: list[dict], rotulos=CONTRATOS) -> RelatorioProjecaoExecucao:
    mensal = []
    for _ in linhas:
        registro = {f"{prefixo}{m}": None for prefixo in "rpte" for m in range(1, 13)}
        registro.update({"r6": 100.0, "p10": 50.0, "t10": TIPO_PROJETADO})
        mensal.append(registro)
    return RelatorioProjecaoExecucao(
        linhas=pd.DataFrame(linhas), mensal=pd.DataFrame(mensal), exercicio=2026, mes_referencia=10, ultimo_fechado=7,
        qtd_sem_ne=0, rotulos=rotulos,
    )


def _contratos() -> RelatorioProjecaoExecucao:
    return _relatorio([
        _linha("2026NE000001", saldo=1000.0, projetado=400.0, cheio=700.0),                        # sobra 600 · cheio 300 · faixa 2
        _linha("2026NE000002", saldo=500.0, projetado=0.0, cheio=0.0, status="VENCIDO"),           # sobra 500 · faixa 1
        _linha("2026NE000003", saldo=800.0, projetado=300.0, cheio=900.0, origem=ORIGEM_SEM_HISTORICO, teto=True),  # 500 · cheio 0 · faixa 3
        _linha("2026NE000004", saldo=200.0, projetado=250.0, cheio=250.0),                         # saldo < projetado: sem sobra
        _linha("2026NE000005", saldo=None, projetado=100.0, cheio=100.0),                          # saldo nulo: sobra nula
        _linha("2026NE000006", saldo=300.0, projetado=None, cheio=None),                           # sem projeção: sobra nula
    ])


class TestSobraPorNe(unittest.TestCase):
    def setUp(self):
        self.tabela = sobra_por_ne({"Contratos Contínuos": _contratos()}, REF).set_index("NE")

    def test_sobra_e_saldo_menos_projetado_nunca_negativa(self):
        self.assertAlmostEqual(self.tabela.at["2026NE000001", COL_SOBRA], 600.0)
        self.assertAlmostEqual(self.tabela.at["2026NE000002", COL_SOBRA], 500.0)
        self.assertAlmostEqual(self.tabela.at["2026NE000003", COL_SOBRA], 500.0)
        self.assertEqual(self.tabela.at["2026NE000004", COL_SOBRA], 0.0)  # projetado acima do saldo: zero, não negativo

    def test_nulo_nao_vira_zero(self):
        self.assertTrue(pd.isna(self.tabela.at["2026NE000005", COL_SOBRA]))  # saldo desconhecido
        self.assertTrue(pd.isna(self.tabela.at["2026NE000006", COL_SOBRA]))  # sem projeção calculável
        self.assertEqual(self.tabela.at["2026NE000005", "Faixa da sobra"], SEM_SOBRA)

    def test_sobra_pelo_valor_cheio(self):
        self.assertAlmostEqual(self.tabela.at["2026NE000001", "Sobra pelo valor cheio"], 300.0)  # 1000 − 700
        self.assertEqual(self.tabela.at["2026NE000003", "Sobra pelo valor cheio"], 0.0)          # 800 − 900 → 0

    def test_faixas(self):
        self.assertEqual(self.tabela.at["2026NE000001", "Faixa da sobra"], FAIXA_2)
        self.assertEqual(self.tabela.at["2026NE000002", "Faixa da sobra"], FAIXA_1)
        self.assertEqual(self.tabela.at["2026NE000003", "Faixa da sobra"], FAIXA_3)
        self.assertEqual(self.tabela.at["2026NE000004", "Faixa da sobra"], SEM_SOBRA)

    def test_ordem_maior_sobra_primeiro_e_nulos_por_ultimo(self):
        ordem = list(sobra_por_ne({"Contratos Contínuos": _contratos()}, REF)["NE"])
        self.assertEqual(ordem[:3], ["2026NE000001", "2026NE000002", "2026NE000003"])
        self.assertEqual(set(ordem[-2:]), {"2026NE000005", "2026NE000006"})

    def test_liquidado_e_empenhado(self):
        self.assertAlmostEqual(self.tabela.at["2026NE000001", "Liquidado (empenhado − saldo)"], 0.0)  # 1000 − 1000
        self.assertTrue(pd.isna(self.tabela.at["2026NE000005", "Liquidado (empenhado − saldo)"]))

    def test_observacoes(self):
        obs = self.tabela["Observações"]
        self.assertIn("conferir faturas pendentes", obs["2026NE000002"])
        self.assertIn("teto", obs["2026NE000003"])
        self.assertIn("pelo valor cheio não há sobra", obs["2026NE000003"])
        self.assertEqual(obs["2026NE000004"], "")

    def test_vigencia_encerrada_com_status_ativo(self):
        rel = _relatorio([_linha("2026NE000010", saldo=100.0, projetado=10.0, cheio=10.0, vigencia="2026-10-03")])
        obs = sobra_por_ne({"Contratos Contínuos": rel}, REF)["Observações"].iloc[0]
        self.assertIn("vigência encerrada em 03/10/2026 com status ATIVO", obs)

    def test_fator_muito_baixo_e_ajuste_manual(self):
        rel = _relatorio([_linha("2026NE000011", saldo=100.0, projetado=10.0, cheio=10.0, fator=0.2, ajustada=True)])
        obs = sobra_por_ne({"Contratos Contínuos": rel}, REF)["Observações"].iloc[0]
        self.assertIn("fator muito baixo (0.20)", obs)
        self.assertIn("ajuste manual", obs)

    def test_nao_altera_o_relatorio_de_entrada(self):
        rel = _contratos()
        copia = rel.linhas.copy()
        sobra_por_ne({"Contratos Contínuos": rel}, REF)
        pd.testing.assert_frame_equal(rel.linhas, copia)

    def test_sem_nenhuma_ne(self):
        vazio = sobra_por_ne({}, REF)
        self.assertTrue(vazio.empty)
        self.assertEqual(int(resumo_da_sobra(vazio).iloc[-1]["Nº de NEs"]), 0)


class TestSobraElegivel(unittest.TestCase):
    def test_as_tres_faixas_entram_por_padrao(self):
        # 15/10/2026: vale para qualquer NE e para empenhos futuros (regra por faixa, nunca por NE)
        tabela = sobra_por_ne({"Contratos Contínuos": _contratos()}, REF)
        self.assertEqual(FAIXAS_NA_ANULACAO, (FAIXA_1, FAIXA_2, FAIXA_3))
        # NE1 (faixa 2): 600; NE2 (faixa 1): 500; NE3 (faixa 3, sem histórico): 500; as sem sobra ficam de fora
        self.assertEqual(
            sobra_elegivel_por_ne(tabela), {"2026NE000001": 600.0, "2026NE000002": 500.0, "2026NE000003": 500.0},
        )

    def test_faixas_configuraveis(self):
        tabela = sobra_por_ne({"Contratos Contínuos": _contratos()}, REF)
        so_1_e_2 = sobra_elegivel_por_ne(tabela, (FAIXA_1, FAIXA_2))
        self.assertEqual(so_1_e_2, {"2026NE000001": 600.0, "2026NE000002": 500.0})

    def test_ne_nas_duas_bases_soma(self):
        outra = _relatorio([_linha("2026NE000001", saldo=200.0, projetado=50.0, cheio=50.0)], BOLSAS_AUXILIOS)
        tabela = sobra_por_ne({"Contratos Contínuos": _contratos(), "Bolsas e Auxílios": outra}, REF)
        self.assertEqual(sobra_elegivel_por_ne(tabela)["2026NE000001"], 600.0 + 150.0)


class TestSugestaoDeAnulacaoPorNe(unittest.TestCase):
    """15/10/2026: o relatório de Anulação traz, por NE do processo, o valor TOTAL sugerido (a sobra prevista), sem rateio."""

    def setUp(self):
        self.tabela = sobra_por_ne({"Contratos Contínuos": _contratos()}, REF)

    def test_valor_total_da_sobra_por_ne_sem_rateio(self):
        sugestao = sugestao_anulacao_por_ne(self.tabela, ["2026NE000001", "2026NE000002"]).set_index("ne_curta")
        self.assertEqual(sugestao.loc["2026NE000001", "sugestao"], 600.0)  # saldo 1.000 − projetado 400
        self.assertEqual(sugestao.loc["2026NE000002", "sugestao"], 500.0)
        self.assertEqual(sugestao["sugestao"].sum(), 1100.0)

    def test_traz_saldo_projetado_faixa_e_observacoes(self):
        linha = sugestao_anulacao_por_ne(self.tabela, ["2026NE000001"]).iloc[0]
        self.assertEqual(linha["saldo"], 1000.0)
        self.assertEqual(linha["projetado"], 400.0)
        self.assertEqual(linha["faixa"], FAIXA_2)
        self.assertEqual(linha["fornecedor"], "FORNECEDOR")
        self.assertIsInstance(linha["observacoes"], str)

    def test_sem_sobra_e_zero_conhecido_nao_nulo(self):
        # saldo 200 < projetado 250: nada a anular — zero (conhecido), nunca vazio
        linha = sugestao_anulacao_por_ne(self.tabela, ["2026NE000004"]).iloc[0]
        self.assertEqual(linha["sugestao"], 0.0)

    def test_sem_projecao_calculavel_fica_nulo_e_nao_zero(self):
        for ne in ("2026NE000005", "2026NE000006"):  # saldo nulo / sem projeção
            with self.subTest(ne=ne):
                self.assertTrue(pd.isna(sugestao_anulacao_por_ne(self.tabela, [ne]).iloc[0]["sugestao"]))

    def test_saldo_com_residuo_de_ponto_flutuante_e_zero_conhecido_e_nao_vazio(self):
        # saldo 5,8e-11 (resíduo de empenho totalmente liquidado) e nada projetado: nada a anular → 0,00, não "—"
        tabela = sobra_por_ne(
            {"Contratos Contínuos": _relatorio([_linha("2026NE000551", saldo=5.8e-11, projetado=0.0, cheio=0.0)])}, REF
        )
        linha = sugestao_anulacao_por_ne(tabela, ["2026NE000551"]).iloc[0]
        self.assertEqual(linha["sugestao"], 0.0)

    def test_ne_que_nao_esta_na_projecao_aparece_vazia_e_avisada(self):
        linha = sugestao_anulacao_por_ne(self.tabela, ["2026NE999999"]).iloc[0]
        self.assertEqual(linha["ne_curta"], "2026NE999999")
        self.assertTrue(pd.isna(linha["sugestao"]))
        self.assertIn("projeção", linha["observacoes"])

    def test_so_as_nes_pedidas_e_sem_repetir(self):
        sugestao = sugestao_anulacao_por_ne(self.tabela, ["2026NE000001", "2026NE000001", "2026NE000002"])
        self.assertEqual(sorted(sugestao["ne_curta"]), ["2026NE000001", "2026NE000002"])

    def test_tabela_nula_ou_vazia(self):
        self.assertTrue(sugestao_anulacao_por_ne(None, ["2026NE000001"]).empty)
        self.assertTrue(sugestao_anulacao_por_ne(self.tabela, []).empty)

    def test_faixa_fora_da_anulacao_nao_sugere_valor(self):
        from unittest import mock
        with mock.patch("src.sobra_empenho.FAIXAS_NA_ANULACAO", (FAIXA_1,)):
            linha = sugestao_anulacao_por_ne(self.tabela, ["2026NE000001"]).iloc[0]  # faixa 2
        self.assertTrue(pd.isna(linha["sugestao"]))
        self.assertEqual(linha["faixa"], FAIXA_2)


class TestFaixa(unittest.TestCase):
    def test_sem_sobra_nao_tem_faixa(self):
        self.assertEqual(faixa_da_sobra("VENCIDO", 0.0, ORIGEM_EXECUCAO, 0.0), SEM_SOBRA)
        self.assertEqual(faixa_da_sobra("VENCIDO", 0.0, ORIGEM_EXECUCAO, None), SEM_SOBRA)

    def test_suspenso_ou_nada_projetado_e_faixa_1(self):
        self.assertEqual(faixa_da_sobra("suspenso", 50.0, ORIGEM_EXECUCAO, 10.0), FAIXA_1)
        self.assertEqual(faixa_da_sobra("ATIVO", 0.0, ORIGEM_EXECUCAO, 10.0), FAIXA_1)


class TestResumo(unittest.TestCase):
    def test_totais_por_faixa_e_base(self):
        bolsas = _relatorio([_linha("2026NE000020", saldo=475.0, projetado=100.0, cheio=475.0, status="ATUALIZADO")], BOLSAS_AUXILIOS)
        tabela = sobra_por_ne({"Contratos Contínuos": _contratos(), "Bolsas e Auxílios": bolsas}, REF)
        resumo = resumo_da_sobra(tabela)
        total = resumo.iloc[-1]
        self.assertEqual(total["Base"], "TOTAL")
        self.assertEqual(int(total["Nº de NEs"]), 4)                      # 3 de contratos + 1 de bolsas
        self.assertAlmostEqual(total["Sobra prevista (R$)"], 600 + 500 + 500 + 375)
        # valor cheio: NE1 1000−700 = 300; NE2 (vencida, nada projetado) 500−0 = 500; NE3 800−900 → 0; bolsas 475−475 → 0
        self.assertAlmostEqual(total["Sobra pelo valor cheio (R$)"], 300 + 500 + 0 + 0)
        faixa2 = resumo[(resumo["Base"] == "Contratos Contínuos") & (resumo["Faixa da sobra"] == FAIXA_2)].iloc[0]
        self.assertAlmostEqual(faixa2["Sobra prevista (R$)"], 600.0)
        self.assertAlmostEqual(resumo.iloc[:-1]["Sobra prevista (R$)"].sum(), total["Sobra prevista (R$)"])


class TestXlsx(unittest.TestCase):
    def setUp(self):
        self.livro = load_workbook(BytesIO(gerar_xlsx({"Contratos Contínuos": _contratos()}, REF, CONTEXTO)))

    def test_abas(self):
        self.assertEqual(self.livro.sheetnames, ["Resumo", "Sobra por NE", "Grade mensal", "Metodologia e fontes"])

    def test_sobra_por_ne_ordenada_e_com_nulo_vazio(self):
        aba = self.livro["Sobra por NE"]
        cabecalho = [c.value for c in aba[1]]
        coluna = cabecalho.index(COL_SOBRA)
        linhas = {r[cabecalho.index("NE")]: r[coluna] for r in aba.iter_rows(min_row=2, values_only=True)}
        self.assertAlmostEqual(linhas["2026NE000001"], 600.0)
        self.assertIsNone(linhas["2026NE000005"])  # nulo é célula vazia, nunca 0
        self.assertEqual(aba.cell(2, 2).value, "2026NE000001")

    def test_grade_mensal_tem_realizado_e_projetado(self):
        aba = self.livro["Grade mensal"]
        cabecalho = [c.value for c in aba[1]]
        primeira = dict(zip(cabecalho, next(aba.iter_rows(min_row=2, max_row=2, values_only=True))))
        self.assertEqual(primeira["Jun"], 100.0)   # realizado
        self.assertEqual(primeira["Out"], 50.0)    # projetado
        self.assertAlmostEqual(primeira["SOBRA PREVISTA"], 600.0)

    def test_metodologia_traz_as_fontes(self):
        aba = self.livro["Metodologia e fontes"]
        itens = {r[0]: r[1] for r in aba.iter_rows(min_row=2, values_only=True)}
        self.assertIn("sha abcd1234", itens["Execução Mensal"])
        self.assertEqual(itens["Emitido em"], "16/10/2026 10:00")
        self.assertIn("NÃO devolve", itens["Sobra prevista"])


if __name__ == "__main__":
    unittest.main()
