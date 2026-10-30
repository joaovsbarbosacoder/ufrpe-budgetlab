"""Testes do relatório (PDF/Excel) de Empenhos com Execução Retardada —
`src/relatorio_empenhos_retardada.py`. Lista sintética, montada à mão, no mesmo formato de
colunas que a página passa (`saldo_por_ne(agregar_por_ne(...))` + `percentual_saldo` +
`passa_corte`)."""

from __future__ import annotations

import base64
import re
import unittest
import zlib
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from src.relatorio_empenhos_retardada import (
    COLUNA_DECISAO,
    COLUNA_GRUPO,
    COLUNA_JUSTIFICATIVA,
    COLUNAS_PDF_PADRAO,
    ROTULOS_COLUNAS_PDF,
    ContextoRelatorio,
    gerar_pdf,
    gerar_pdf_analise,
    gerar_pdf_resumo,
    gerar_xlsx,
    gerar_xlsx_analise,
    gerar_xlsx_resumo,
    montar_lista_relatorio,
    ordenar_lista,
    resumo_por_grupo,
)


def _lista() -> pd.DataFrame:
    return pd.DataFrame(
        {
            "ne_ccor": ["153165152662026NE000056", "153165152662026NE000123", "153165152662026NE000200"],
            "ano": [2026, 2026, 2026],
            "ne_favorecido": ["EMPRESA A", "EMPRESA B & CIA <LTDA>", None],
            "ne_descricao": ["Serviço", None, "Bolsa"],
            "processo_ne": ["23082.000001/2026-01", None, "23082.000003/2026-03"],
            "acao_cod": ["20RK", "20RK", "4002"],
            "acao_desc": ["FUNCIONAMENTO", "FUNCIONAMENTO", "ASSISTENCIA"],
            "ptres": ["012345", "012345", "000789"],
            "fonte_cod": ["1000000000", "1000000000", "1000000000"],
            "fonte_desc": ["RECURSOS", "RECURSOS", "RECURSOS"],
            "gnd_cod": ["3", "3", "3"],
            "natureza_despesa_cod": ["339039", "339039", "339018"],
            "natureza_despesa_desc": ["OUTROS SERV.", "OUTROS SERV.", "AUX. FINANC."],
            "natureza_detalhada_label": ["33903917", "33903917", "33901801"],
            "subitem_resumo": ["17", "17", "01"],
            "pi_cod": ["L20RKG0100N", "L20RKG0100N", "M4002G0101N"],
            "pi_desc": ["PI A", "PI A", "PI B"],
            "ugr_cod": ["153165", "153165", "153165"],
            "ugr_desc": ["UGR", "UGR", "UGR"],
            "ug_executora_cod": ["153165", "153165", "153165"],
            "ne_informacao_complementar": [None, None, None],
            "empenhada": [100_000.0, 50_000.0, None],
            "liquidada": [40_000.0, 0.0, None],
            "saldo": [60_000.0, 50_000.0, None],
            "percentual_saldo": [60.0, 100.0, None],
            "passa_corte": [True, True, False],
        }
    )


def _texto_pdf(conteudo: bytes) -> bytes:
    """Conteúdo das páginas do PDF já descomprimido (o ReportLab grava em ASCII85 + Flate)."""

    partes = []
    for trecho in re.findall(rb"stream\n(.*?)~>endstream", conteudo, flags=re.S):
        partes.append(zlib.decompress(base64.a85decode(trecho)))
    return b"".join(partes)


def _contexto(**kwargs) -> ContextoRelatorio:
    padrao = dict(
        visualizacao="Todos no escopo",
        descricao_corte="saldo ≥ R$ 10.000",
        data_extracao="28/09/2026",
        hash_manifesto="abcdef12",
        data_emissao="28/09/2026 10:00",
        busca="",
        filtros=[("Exercício", ["2026"])],
    )
    padrao.update(kwargs)
    return ContextoRelatorio(**padrao)


class MontarListaRelatorioTests(unittest.TestCase):
    def setUp(self) -> None:
        # ordem de entrada propositalmente fora da ordem de saldo; NE com saldo nulo no meio.
        self.visivel = pd.DataFrame(
            {"ne_ccor": ["A", "B", "C", "D"], "saldo": [5_000.0, None, 80_000.0, 20_000.0]},
            index=[10, 11, 12, 13],
        )
        self.destaque = pd.Series([False, False, True, True], index=self.visivel.index)

    def test_em_destaque_so_as_marcadas_ordenadas_por_saldo(self) -> None:
        lista = montar_lista_relatorio(self.visivel, self.destaque, cortes_ligados=True, todos=False)
        self.assertEqual(lista["ne_ccor"].tolist(), ["C", "D"])
        self.assertEqual(lista["passa_corte"].tolist(), [True, True])
        self.assertEqual(lista.index.tolist(), [0, 1])

    def test_todos_inclui_todas_com_nulo_por_ultimo(self) -> None:
        lista = montar_lista_relatorio(self.visivel, self.destaque, cortes_ligados=True, todos=True)
        self.assertEqual(lista["ne_ccor"].tolist(), ["C", "D", "A", "B"])
        self.assertEqual(lista["passa_corte"].tolist(), [True, True, False, False])
        self.assertTrue(pd.isna(lista["saldo"].iloc[-1]))  # nulo preservado, não virou zero

    def test_sem_corte_ligado_passa_corte_nulo(self) -> None:
        sem_destaque = pd.Series(False, index=self.visivel.index)
        lista = montar_lista_relatorio(self.visivel, sem_destaque, cortes_ligados=False, todos=True)
        self.assertEqual(len(lista), 4)
        self.assertTrue(lista["passa_corte"].isna().all())

    def test_nao_altera_a_entrada(self) -> None:
        original = self.visivel.copy()
        montar_lista_relatorio(self.visivel, self.destaque, cortes_ligados=True, todos=True)
        pd.testing.assert_frame_equal(self.visivel, original)


class OrdenarListaTests(unittest.TestCase):
    """Ordenação e agrupamento (17/10/2026): só reordenam e rotulam — nenhuma linha some nem muda."""

    def setUp(self) -> None:
        self.lista = pd.DataFrame(
            {
                "ne_ccor": ["a", "b", "c", "d", "e", "f"],
                "ne_favorecido": ["zeta", "Alfa", None, "beta", "Alfa", "zeta"],
                "unidade": ["U1", "U2", "U1", None, "U2", "U1"],
                "saldo": [10.0, 500.0, None, -30.0, 0.0, 200.0],
            }
        )

    def test_padrao_saldo_decrescente_com_nulo_por_ultimo(self) -> None:
        resultado = ordenar_lista(self.lista)
        self.assertEqual(resultado["ne_ccor"].tolist(), ["b", "f", "a", "e", "d", "c"])
        self.assertTrue(pd.isna(resultado["saldo"].iloc[-1]))  # nulo ao fim, não virou zero

    def test_saldo_crescente_mantem_nulo_por_ultimo_e_distingue_zero_e_negativo(self) -> None:
        resultado = ordenar_lista(self.lista, "saldo", ascendente=True)
        self.assertEqual(resultado["ne_ccor"].tolist(), ["d", "e", "a", "f", "b", "c"])  # negativo < zero < positivo; nulo no fim

    def test_texto_ordena_sem_diferenciar_maiusculas_e_nulo_por_ultimo(self) -> None:
        resultado = ordenar_lista(self.lista, "ne_favorecido", ascendente=True)
        self.assertEqual(resultado["ne_favorecido"].tolist()[:5], ["Alfa", "Alfa", "beta", "zeta", "zeta"])
        self.assertTrue(pd.isna(resultado["ne_favorecido"].iloc[-1]))

    def test_agrupamento_ordena_grupos_por_saldo_somado_e_nulo_vai_para_o_fim(self) -> None:
        resultado = ordenar_lista(self.lista, "saldo", agrupar_por="unidade")
        # U2 = 500 (b) + 0 (e) = 500; U1 = 10 + 200 (+ nulo) = 210; "—" (sem unidade) por último
        self.assertEqual(resultado[COLUNA_GRUPO].tolist(), ["U2", "U2", "U1", "U1", "U1", "—"])
        self.assertEqual(resultado["ne_ccor"].tolist(), ["b", "e", "f", "a", "c", "d"])

    def test_nao_perde_nem_altera_linhas(self) -> None:
        original = self.lista.copy()
        resultado = ordenar_lista(self.lista, "ne_ccor", agrupar_por="unidade")
        self.assertEqual(sorted(resultado["ne_ccor"]), sorted(original["ne_ccor"]))
        self.assertEqual(resultado["saldo"].sum(min_count=1), original["saldo"].sum(min_count=1))
        pd.testing.assert_frame_equal(self.lista, original)  # entrada intacta
        self.assertNotIn("_grupo_saldo", resultado.columns)

    def test_montar_lista_aceita_ordenacao_e_agrupamento(self) -> None:
        destaque = pd.Series(True, index=self.lista.index)
        resultado = montar_lista_relatorio(
            self.lista, destaque, cortes_ligados=True, todos=True, ordenar_por="saldo", ascendente=True, agrupar_por="unidade"
        )
        self.assertIn(COLUNA_GRUPO, resultado.columns)
        self.assertEqual(len(resultado), 6)


class GerarXlsxTests(unittest.TestCase):
    def setUp(self) -> None:
        self.lista = _lista()
        self.original = self.lista.copy()
        conteudo = gerar_xlsx(self.lista, _contexto())
        self.livro = load_workbook(BytesIO(conteudo))

    def _linhas(self) -> list[dict]:
        planilha = self.livro["Empenhos"]
        cabecalho = [c.value for c in planilha[1]]
        return [dict(zip(cabecalho, [c.value for c in linha])) for linha in planilha.iter_rows(min_row=2)]

    def test_uma_linha_por_empenho_sem_linha_de_total(self) -> None:
        self.assertEqual(len(self._linhas()), 3)

    def test_codigos_ficam_como_texto_com_zeros_iniciais(self) -> None:
        linha = self._linhas()[0]
        self.assertEqual(linha["PTRES"], "012345")
        self.assertEqual(linha["NE"], "2026NE000056")
        self.assertEqual(linha["NE (código completo)"], "153165152662026NE000056")
        self.assertEqual(self._linhas()[2]["Subitem"], "01")

    def test_valores_numericos_e_nulo_distinto_de_zero(self) -> None:
        linhas = self._linhas()
        self.assertEqual(linhas[0]["Saldo"], 60_000.0)
        self.assertEqual(linhas[1]["Liquidado"], 0.0)  # zero continua zero
        self.assertIsNone(linhas[2]["Empenhado"])  # nulo continua vazio
        self.assertIsNone(linhas[2]["% Saldo"])

    def test_passa_do_corte(self) -> None:
        self.assertEqual([l["Passa do corte"] for l in self._linhas()], ["Sim", "Sim", "Não"])

    def test_sem_corte_ligado_passa_do_corte_fica_vazio(self) -> None:
        lista = _lista().assign(passa_corte=pd.NA)
        livro = load_workbook(BytesIO(gerar_xlsx(lista, _contexto(descricao_corte=None))))
        planilha = livro["Empenhos"]
        cabecalho = [c.value for c in planilha[1]]
        coluna = cabecalho.index("Passa do corte")
        self.assertEqual([linha[coluna].value for linha in planilha.iter_rows(min_row=2)], [None, None, None])

    def test_aba_parametros_com_procedencia_filtros_e_totais(self) -> None:
        parametros = {linha[0].value: linha[1].value for linha in self.livro["Parâmetros"].iter_rows(min_row=2)}
        self.assertEqual(parametros["Visualização"], "Todos no escopo")
        self.assertEqual(parametros["Exercício"], "2026")
        self.assertIn("abcdef12", parametros["Procedência"])
        self.assertEqual(parametros["Empenhos listados"], "3")
        self.assertEqual(parametros["Total saldo"], "110.000,00")

    def test_nao_altera_a_lista_recebida(self) -> None:
        pd.testing.assert_frame_equal(self.lista, self.original)


class IdadeNoRelatorioTests(unittest.TestCase):
    def test_xlsx_traz_idade_como_numero_e_vazio_quando_nula(self) -> None:
        lista = _lista().assign(idade_meses=pd.array([6, 0, pd.NA], dtype="Int64"))
        planilha = load_workbook(BytesIO(gerar_xlsx(lista, _contexto())))["Empenhos"]
        cabecalho = [c.value for c in planilha[1]]
        coluna = cabecalho.index("Idade do saldo (meses)")
        self.assertEqual([linha[coluna].value for linha in planilha.iter_rows(min_row=2)], [6, 0, None])

    def test_pdf_com_idade_e_sem_a_coluna(self) -> None:
        com = _lista().assign(idade_meses=pd.array([6, 0, pd.NA], dtype="Int64"))
        self.assertTrue(gerar_pdf(com, _contexto()).startswith(b"%PDF"))
        self.assertTrue(gerar_pdf(_lista(), _contexto()).startswith(b"%PDF"))  # lista antiga, sem idade


class RelatorioAgrupadoTests(unittest.TestCase):
    def test_xlsx_traz_coluna_grupo_e_parametros_de_ordenacao_e_agrupamento(self) -> None:
        lista = ordenar_lista(_lista(), "saldo", agrupar_por="acao_desc")
        contexto = _contexto(ordenacao="Saldo (maior primeiro)", agrupamento="Ação de governo")
        livro = load_workbook(BytesIO(gerar_xlsx(lista, contexto)))
        planilha = livro["Empenhos"]
        cabecalho = [c.value for c in planilha[1]]
        self.assertEqual(cabecalho[0], "Grupo")
        self.assertEqual(len(list(planilha.iter_rows(min_row=2))), 3)  # nenhuma linha extra de grupo
        parametros = {linha[0].value: linha[1].value for linha in livro["Parâmetros"].iter_rows(min_row=2)}
        self.assertEqual(parametros["Agrupamento"], "Ação de governo")
        self.assertEqual(parametros["Ordenação"], "Saldo (maior primeiro)")

    def test_xlsx_sem_agrupamento_nao_tem_coluna_grupo_nem_parametros_extras(self) -> None:
        livro = load_workbook(BytesIO(gerar_xlsx(_lista(), _contexto())))
        self.assertNotIn("Grupo", [c.value for c in livro["Empenhos"][1]])
        parametros = [linha[0].value for linha in livro["Parâmetros"].iter_rows(min_row=2)]
        self.assertNotIn("Agrupamento", parametros)
        self.assertNotIn("Ordenação", parametros)

    def test_pdf_agrupado_e_valido(self) -> None:
        lista = ordenar_lista(_lista(), "saldo", agrupar_por="acao_desc")
        conteudo = gerar_pdf(lista, _contexto(agrupamento="Ação de governo"))
        self.assertTrue(conteudo.startswith(b"%PDF"))


class ModelosDeRelatorioTests(unittest.TestCase):
    """Modelos de 30/10/2026: escolha de colunas do PDF, resumo por grupo e análise de anulação."""

    def setUp(self) -> None:
        self.lista = ordenar_lista(
            _lista().assign(indicacoes_anulacao=["Saldo residual", None, None], indicacao_principal=["Saldo residual", "Sem indicação", "Sem indicação"]),
            "saldo",
            agrupar_por="acao_desc",
        )

    def test_pdf_com_colunas_escolhidas(self) -> None:
        conteudo = gerar_pdf(self.lista, _contexto(), ["ne", "fonte", "ptres", "saldo"])
        self.assertTrue(conteudo.startswith(b"%PDF"))
        self.assertIn(b"(FONTE)", _texto_pdf(conteudo))
        self.assertIn(b"(012345)", _texto_pdf(conteudo))  # PTRES como texto, zero à esquerda
        self.assertNotIn(b"FAVORECIDO", _texto_pdf(conteudo))
        # chave desconhecida é ignorada; nenhuma válida volta ao padrão
        self.assertIn(b"FAVORECIDO", _texto_pdf(gerar_pdf(self.lista, _contexto(), ["inexistente"])))
        self.assertEqual(COLUNAS_PDF_PADRAO[:3], ["ne", "exercicio", "favorecido"])
        self.assertTrue(set(COLUNAS_PDF_PADRAO) <= set(ROTULOS_COLUNAS_PDF))

    def test_resumo_por_grupo_calculado_a_mao(self) -> None:
        resumo = resumo_por_grupo(self.lista).set_index("Grupo")
        # FUNCIONAMENTO: 2 NEs, empenhado 150.000, liquidado 40.000, saldo 110.000 (100% do saldo)
        self.assertEqual(resumo.loc["FUNCIONAMENTO", "Empenhos"], 2)
        self.assertEqual(resumo.loc["FUNCIONAMENTO", "Saldo"], 110_000.0)
        self.assertAlmostEqual(resumo.loc["FUNCIONAMENTO", "% liquidado"], 40_000 / 150_000 * 100)
        self.assertAlmostEqual(resumo.loc["FUNCIONAMENTO", "% do saldo do relatório"], 100.0)
        # ASSISTENCIA: só nulos -> soma nula, não zero; % nulo
        self.assertTrue(pd.isna(resumo.loc["ASSISTENCIA", "Saldo"]))
        self.assertTrue(pd.isna(resumo.loc["ASSISTENCIA", "% liquidado"]))
        self.assertEqual(resumo.loc["TOTAL", "Empenhos"], 3)
        self.assertEqual(list(resumo.index), ["FUNCIONAMENTO", "ASSISTENCIA", "TOTAL"])
        # sem agrupamento: só a linha TOTAL
        self.assertEqual(resumo_por_grupo(_lista())["Grupo"].tolist(), ["TOTAL"])

    def test_resumo_em_pdf_e_excel_e_aba_resumo_no_excel_agrupado(self) -> None:
        self.assertTrue(gerar_pdf_resumo(self.lista, _contexto()).startswith(b"%PDF"))
        livro = load_workbook(BytesIO(gerar_xlsx_resumo(self.lista, _contexto())))
        self.assertEqual(livro.sheetnames, ["Resumo", "Parâmetros"])
        self.assertIn("Resumo", load_workbook(BytesIO(gerar_xlsx(self.lista, _contexto()))).sheetnames)

    def test_analise_de_anulacao(self) -> None:
        conteudo = gerar_pdf_analise(self.lista, _contexto(indicacoes=[("Saldo residual", "regra")]))
        self.assertIn(b"JUSTIFICATIVA", _texto_pdf(conteudo))
        self.assertIn(b"Saldo residual", _texto_pdf(conteudo))
        livro = load_workbook(BytesIO(gerar_xlsx_analise(self.lista, _contexto(indicacoes=[("Saldo residual", "regra")]))))
        planilha = livro["Empenhos"]
        cabecalho = [c.value for c in planilha[1]]
        self.assertEqual(cabecalho[-2:], [COLUNA_DECISAO, COLUNA_JUSTIFICATIVA])
        self.assertIn("Indicações para anulação", cabecalho)
        decisao = cabecalho.index(COLUNA_DECISAO)
        self.assertTrue(all(linha[decisao].value is None for linha in planilha.iter_rows(min_row=2)))
        validacoes = planilha.data_validations.dataValidation
        self.assertEqual(len(validacoes), 1)
        self.assertIn("Manter,Anular,Reforçar", validacoes[0].formula1)
        parametros = {linha[0].value: linha[1].value for linha in livro["Parâmetros"].iter_rows(min_row=2)}
        self.assertEqual(parametros["Indicação: Saldo residual"], "regra")


class GerarPdfTests(unittest.TestCase):
    def test_gera_pdf_valido_com_texto_escapado(self) -> None:
        conteudo = gerar_pdf(_lista(), _contexto(busca="a & b"))
        self.assertTrue(conteudo.startswith(b"%PDF"))

    def test_gera_pdf_com_muitas_linhas_em_varias_paginas(self) -> None:
        lista = pd.concat([_lista()] * 200, ignore_index=True)
        conteudo = gerar_pdf(lista, _contexto(descricao_corte=None, filtros=[]))
        self.assertTrue(conteudo.startswith(b"%PDF"))
        self.assertGreater(conteudo.count(b"/Type /Page\n") + conteudo.count(b"/Type /Page "), 1)


if __name__ == "__main__":
    unittest.main()
