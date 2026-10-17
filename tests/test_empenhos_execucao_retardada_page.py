"""Testes da página Empenhos com Execução Retardada, com fixture congelada da Execução Mensal.

A página lê `Manifesto.atual()`/`carregar_atual()` de `src.importacao_execucao_mensal`; a suíte
mocka essas duas funções contra a MESMA fixture congelada que Consulta de Empenhos já usa
(`tests/fixtures/execucao_mensal_2026-09-22.xlsx`) — não depende de nenhum manifesto real em
`data/manifestos/`.

Desde 17/10/2026 (layout do mockup aprovado, página inteira) a tela abaixo do cabeçalho é um
componente `components.v2` (`src/ui_empenhos_retardada_componente.py`), que o AppTest não
consegue clicar. Por isso estes testes:
  * capturam os DADOS que a página entrega ao componente (`renderizar` substituído por um espião)
    — é exatamente o que a tela desenha;
  * simulam as ações do usuário gravando o estado da tela em `st.session_state` com
    `aplicar_acao` (a mesma função que o callback do componente usa) e reexecutando a página.
As regras em si (cascata de filtros, cortes, lista, paginação, detalhe) são testadas sem
Streamlit em `tests/test_empenhos_retardada_tela.py`; aqui fica a integração da página: estado
inicial, avisos, relatório = tela, procedência, coleta de memória.
"""

from __future__ import annotations

import unittest
from datetime import datetime
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from tests._apptest import TEMPO_LIMITE_APPTEST

from src.empenhos_retardada_tela import AVISO_BUSCA_VAZIA, AVISO_SEM_CORTE, aplicar_acao
from src.importacao_execucao_mensal import gerar_manifesto
from src.relatorio_empenhos_retardada import ContextoRelatorio, gerar_xlsx
from src.tesouro_execucao_mensal import ler_execucao_mensal

PROJECT_ROOT = Path(__file__).resolve().parents[1]
CAMINHO_FIXTURE = PROJECT_ROOT / "tests/fixtures/execucao_mensal_2026-09-22.xlsx"
NOMES_FILTRO = [
    "ano", "acao", "ugr", "gnd", "iduso", "resultado_primario", "categoria_economica", "elemento", "fonte",
    "natureza_despesa", "natureza_detalhada", "subitem", "pi", "ptres", "ug_executora", "ug_responsavel",
    "processo", "ne_informacao_complementar", "unidade_orcamentaria",
]


class EmpenhosExecucaoRetardadaPageTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.dataframe = ler_execucao_mensal(CAMINHO_FIXTURE)
        cls.manifesto = gerar_manifesto(cls.dataframe, CAMINHO_FIXTURE)
        cls.chave_estado = f"empenhos_retardada_estado_{cls.manifesto.sha256[:12]}"

    def setUp(self) -> None:
        self.capturados: list[dict] = []
        patches = [
            patch("src.importacao_execucao_mensal.Manifesto.atual", return_value=self.manifesto),
            patch("src.importacao_execucao_mensal.carregar_atual", return_value=self.dataframe.copy()),
            patch(
                "src.ui_empenhos_retardada_componente.renderizar",
                side_effect=lambda dados, **_kw: self.capturados.append(dados),
            ),
        ]
        for p in patches:
            p.start()
            self.addCleanup(p.stop)

    def _open_page(self) -> AppTest:
        app = AppTest.from_file(str(PROJECT_ROOT / "app.py"), default_timeout=TEMPO_LIMITE_APPTEST)
        app.run()
        app.switch_page("app_pages/empenhos_execucao_retardada.py")
        app.run()
        return app

    def _agir(self, app: AppTest, *acoes: dict) -> dict:
        """Aplica ações como o callback do componente faria e reexecuta a página; devolve os
        dados que a tela recebeu."""

        estado = app.session_state[self.chave_estado]
        for acao in acoes:
            estado = aplicar_acao(estado, acao, NOMES_FILTRO)
        app.session_state[self.chave_estado] = estado
        app.run()
        self.assertEqual(len(app.exception), 0)
        return self.capturados[-1]

    @property
    def dados(self) -> dict:
        return self.capturados[-1]

    # --- página ------------------------------------------------------------------------------
    def test_renders_page_from_current_manifest(self) -> None:
        app = self._open_page()

        self.assertEqual(len(app.exception), 0)
        self.assertEqual(app.title[0].value, "Empenhos com Execução Retardada")
        self.assertGreater(self.dados["resumo"]["empenhos"], 0)
        self.assertIsNotNone(self.dados["lista"])

    def test_coleta_lixo_no_inicio_e_no_fim_da_pagina(self) -> None:
        # cópias da base devolvidas pelo st.cache_data ficavam presas em ciclos (~0,3 GB por interação)
        with patch("gc.collect", return_value=0) as coleta:
            app = self._open_page()
            antes = coleta.call_count
            app.run()
        self.assertEqual(len(app.exception), 0)
        self.assertEqual(coleta.call_count - antes, 2)

    def test_shows_procedencia_footer_with_manifest_hash(self) -> None:
        app = self._open_page()

        self.assertTrue(any(f"hash {self.manifesto.sha256[:8]}" in item.value for item in app.caption))

    def test_exercicio_vigente_pre_selecionado(self) -> None:
        self._open_page()
        ano_extracao = datetime.fromisoformat(self.manifesto.data_extracao).year

        exercicio = next(c for c in self.dados["filtros_rapidos"] if c["rotulo"] == "Exercício")
        self.assertEqual(exercicio["selecionados"], [str(ano_extracao)])

    def test_limpar_filtros_volta_exercicio_para_todos(self) -> None:
        app = self._open_page()
        dados = self._agir(app, {"tipo": "limpar"})

        exercicio = next(c for c in dados["filtros_rapidos"] if c["rotulo"] == "Exercício")
        self.assertEqual(exercicio["selecionados"], [])

    def test_filtros_avancados_incluem_as_2_dimensoes_exclusivas_da_execucao_mensal(self) -> None:
        # Paridade com Consulta de Empenhos (pedido explícito, 22/09/2026): os 13 avançados
        # compartilhados (`src/ui_filtros_execucao.py`) mais estas 2, exclusivas desta base.
        self._open_page()

        rotulos = [c["rotulo"] for c in self.dados["filtros_avancados"]]
        self.assertEqual(len(rotulos), 15)
        self.assertIn("NE - Informação Complementar", rotulos)
        self.assertIn("Unidade Orçamentária", rotulos)
        self.assertEqual([c["rotulo"] for c in self.dados["filtros_rapidos"]], ["Exercício", "Ação de Governo", "UGR - Gestão", "Grupo de Despesa"])

    # --- cortes ------------------------------------------------------------------------------
    def test_corte_rs_ligado_por_padrao_pct_desligado(self) -> None:
        self._open_page()

        self.assertTrue(self.dados["cortes"]["usar_rs"])
        self.assertFalse(self.dados["cortes"]["usar_pct"])
        self.assertIn("saldo ≥", self.dados["resultado_corte"])

    def test_desligar_os_dois_cortes_mostra_aviso_sem_tabela_nem_relatorio(self) -> None:
        app = self._open_page()
        dados = self._agir(app, {"tipo": "usar_rs", "valor": False})

        self.assertIn(AVISO_SEM_CORTE, dados["avisos_lista"])
        self.assertIsNone(dados["lista"])
        self.assertEqual(list(app.get("download_button")), [])

    def test_ligar_os_dois_cortes_combina_com_ou_e_permite_e(self) -> None:
        app = self._open_page()
        dados = self._agir(app, {"tipo": "usar_pct", "valor": True})
        self.assertIn(" ou ", dados["resultado_corte"])  # padrão OU preserva o comportamento validado
        com_ou = dados["lista"]["total"]

        dados = self._agir(app, {"tipo": "modo", "valor": "E"})
        self.assertIn(" e ", dados["resultado_corte"])
        self.assertLessEqual(dados["lista"]["total"], com_ou)  # E restringe

    def test_corte_por_idade_desligado_por_padrao_e_ligavel(self) -> None:
        app = self._open_page()
        self.assertFalse(self.dados["cortes"]["usar_idade"])
        self.assertEqual(self.dados["cortes"]["corte_idade"], 3)
        self.assertTrue(all("idade" in l for l in self.dados["lista"]["linhas"]))

        dados = self._agir(app, {"tipo": "usar_rs", "valor": False}, {"tipo": "usar_idade", "valor": True})
        self.assertIn("empenho há 3 meses ou mais com saldo", dados["resultado_corte"])

    def test_dica_do_corte_por_idade_some_ao_ligar_o_corte(self) -> None:
        # a dica só altera a apresentação: a lista é a mesma com ou sem ela
        app = self._open_page()
        self.assertTrue(self.dados["cortes_padrao"])
        total_padrao = self.dados["lista"]["total"]

        dados = self._agir(app, {"tipo": "usar_idade", "valor": True})  # o atalho da dica envia esta ação
        self.assertFalse(dados["cortes_padrao"])
        self.assertGreaterEqual(dados["lista"]["total"], total_padrao)  # OU amplia, nunca reduz

        dados = self._agir(app, {"tipo": "usar_idade", "valor": False})
        self.assertTrue(dados["cortes_padrao"])  # voltou ao padrão: a dica reaparece
        self.assertEqual(dados["lista"]["total"], total_padrao)

    # --- busca -------------------------------------------------------------------------------
    def test_busca_livre_sem_resultado_avisa(self) -> None:
        app = self._open_page()
        dados = self._agir(app, {"tipo": "busca", "valor": "termo-que-nao-deve-existir-em-nenhuma-ne-xyzxyz"})

        self.assertEqual(dados["aviso"], AVISO_BUSCA_VAZIA)
        self.assertIsNone(dados["lista"])

    def test_busca_livre_restringe_opcoes_dos_filtros_rapidos(self) -> None:
        # "informatica", dentro do exercício vigente pré-selecionado, bate numa única "Ação de
        # Governo" nesta fixture: "FUNCIONAMENTO DE INSTITUICOES FEDERAIS DE ENSINO SUPERIOR".
        app = self._open_page()
        dados = self._agir(app, {"tipo": "busca", "valor": "informatica"})

        acao = next(c for c in dados["filtros_rapidos"] if c["rotulo"] == "Ação de Governo")
        self.assertEqual(len(acao["opcoes"]), 1)
        self.assertIn("FUNCIONAMENTO DE INSTITUICOES FEDERAIS DE ENSINO SUPERIOR", acao["opcoes"][0])

    # --- lista e relatório ---------------------------------------------------------------------
    def test_visualizacao_padrao_em_destaque_com_linha_abaixo_do_corte(self) -> None:
        self._open_page()

        self.assertEqual(self.dados["visao"], "Em destaque")
        self.assertIn("empenhos abaixo do corte", self.dados["lista"]["abaixo"])
        self.assertFalse(self.dados["lista"]["mostra_corte"])

    def test_visualizacao_todos_lista_todo_o_escopo_paginado(self) -> None:
        app = self._open_page()
        dados = self._agir(app, {"tipo": "visao", "valor": "Todos no escopo"})

        lista = dados["lista"]
        self.assertEqual(lista["total"], dados["resumo"]["empenhos"])
        self.assertGreater(lista["total_paginas"], 1)  # a fixture exige mais de uma página
        self.assertEqual(len(lista["linhas"]), 100)
        self.assertIsNone(lista["abaixo"])
        self.assertTrue(lista["mostra_corte"])

    def test_relatorio_bate_com_a_tela(self) -> None:
        # A geração do PDF/Excel só roda no clique do download (callable), que o AppTest não
        # dispara: captura-se a lista que a página entrega aos downloads (`montar_lista_relatorio`,
        # a MESMA que monta a tabela) e passa-se pelo gerador de Excel.
        import src.empenhos_retardada_tela as tela

        capturadas = []
        original = tela.montar_lista_relatorio

        def _espia(*args, **kwargs):
            capturadas.append(original(*args, **kwargs))
            return capturadas[-1]

        with patch.object(tela, "montar_lista_relatorio", side_effect=_espia):
            app = self._open_page()
        lista = capturadas[-1]
        total_tela = self.dados["lista"]["total"]

        self.assertEqual(len(lista), total_tela)
        self.assertEqual({b.label for b in app.get("download_button")}, {"Baixar PDF", "Baixar Excel"})
        self.assertTrue(any(f"Relatório: {total_tela} empenhos" in c.value for c in app.caption))
        # primeira linha da tela = primeira linha do relatório (mesma ordenação)
        self.assertEqual(self.dados["lista"]["linhas"][0]["ne_ccor"], str(lista["ne_ccor"].iloc[0]))

        contexto = ContextoRelatorio("x", None, "22/09/2026", "hash", "29/09/2026 10:00")
        planilha = load_workbook(BytesIO(gerar_xlsx(lista, contexto)))["Empenhos"]
        cabecalho = [c.value for c in planilha[1]]
        saldos = [linha[cabecalho.index("Saldo")].value for linha in planilha.iter_rows(min_row=2)]
        self.assertEqual(len(saldos), total_tela)
        self.assertAlmostEqual(sum(v for v in saldos if v is not None), float(lista["saldo"].sum()), places=2)

    def test_ordenar_pelo_cabecalho_inverte_ao_clicar_de_novo(self) -> None:
        app = self._open_page()
        self.assertEqual((self.dados["ordem"], self.dados["sentido"]), ("Saldo", "desc"))

        dados = self._agir(app, {"tipo": "ordenar", "coluna": "Saldo"})
        self.assertEqual((dados["ordem"], dados["sentido"]), ("Saldo", "asc"))

        dados = self._agir(app, {"tipo": "ordenar", "coluna": "Favorecido"})
        self.assertEqual((dados["ordem"], dados["sentido"]), ("Favorecido", "asc"))
        nomes = [l["favorecido"].lower() for l in dados["lista"]["linhas"] if l["favorecido"] != "—"]
        self.assertEqual(nomes, sorted(nomes))

    def test_agrupar_por_unidade_mantem_todos_os_empenhos(self) -> None:
        app = self._open_page()
        total = self.dados["lista"]["total"]
        dados = self._agir(app, {"tipo": "agrupar", "valor": "Unidade orçamentária"})

        self.assertEqual(dados["lista"]["total"], total)
        self.assertEqual(sum(s["n"] for s in dados["lista"]["subtotais"].values()), total)
        self.assertTrue(all(l["grupo"] for l in dados["lista"]["linhas"]))

    def test_trocar_ordenacao_volta_para_a_primeira_pagina(self) -> None:
        app = self._open_page()
        self._agir(app, {"tipo": "visao", "valor": "Todos no escopo"}, {"tipo": "pagina", "valor": 2})
        self.assertEqual(self.dados["lista"]["pagina"], 2)

        dados = self._agir(app, {"tipo": "ordenar", "coluna": "Empenhado"})
        self.assertEqual(dados["lista"]["pagina"], 1)

    def test_detalhe_do_empenho_traz_linhas_de_origem(self) -> None:
        app = self._open_page()
        primeira = self.dados["lista"]["linhas"][0]
        dados = self._agir(app, {"tipo": "detalhe", "ne": primeira["ne_ccor"]})

        self.assertEqual(dados["detalhe"]["ne_ccor"], primeira["ne_ccor"])
        self.assertGreater(len(dados["detalhe"]["origem"]), 0)

        dados = self._agir(app, {"tipo": "fechar_detalhe"})
        self.assertIsNone(dados["detalhe"])


if __name__ == "__main__":
    unittest.main()
