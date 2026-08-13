"""Testes do reconhecimento e normalização de Dotação Anual (BI CPOC - Por Ano)."""

from __future__ import annotations

from io import BytesIO
import unittest

import pandas as pd
from openpyxl import Workbook, load_workbook

from src.tesouro_dotacao_anual import (
    read_dotacao_anual_sheet,
    recognize_dotacao_anual_base,
    transform_dotacao_anual_sheet,
)


METRIC = "Movim. Líquido - R$ (Item Informação)"


def merge_cells(ws, start_row, start_col, end_row, end_col) -> None:
    ws.merge_cells(
        start_row=start_row,
        start_column=start_col,
        end_row=end_row,
        end_column=end_col,
    )


def write_recognized_sheet(ws) -> None:
    """Assinatura válida, deliberadamente sem o bloco Grupo Despesa.

    Isso exercita a detecção dinâmica: a base deve ser reconhecida mesmo
    faltando um bloco dimensional opcional, com a coluna correspondente
    preenchida com nulo em toda a normalização.
    """

    ws["A1"] = "Iduso"
    merge_cells(ws, 1, 1, 3, 2)
    ws["C1"] = "Resultado Primário Lei"
    merge_cells(ws, 1, 3, 3, 4)
    ws["E1"] = "Ação Governo"
    merge_cells(ws, 1, 5, 3, 6)
    ws["G1"] = "PTRES"
    merge_cells(ws, 1, 7, 3, 7)
    ws["H1"] = "Fonte Recursos Detalhada"
    merge_cells(ws, 1, 8, 3, 9)
    ws["J1"] = "Item Informação"
    merge_cells(ws, 1, 10, 1, 11)
    ws["J2"] = "Ano Lançamento"
    merge_cells(ws, 2, 10, 2, 11)
    ws["J3"] = "Plano Orçamentário"
    merge_cells(ws, 3, 10, 3, 11)

    ws["L1"] = "DOTACAO INICIAL"
    ws["L2"] = 2024
    ws["L3"] = METRIC
    ws["M1"] = "DOTACAO ATUALIZADA"
    ws["M2"] = 2024
    ws["M3"] = METRIC
    ws["N1"] = "ITEM EXPERIMENTAL"
    ws["N2"] = 2024
    ws["N3"] = METRIC

    # Linhas 4-5: Ação Governo mesclada nas duas linhas (código+descrição
    # herdados via merge real); as demais dimensões variam por linha.
    ws["A4"] = 1
    ws["B4"] = "Iduso A"
    ws["C4"] = 10
    ws["D4"] = "RP A"
    ws["E4"] = "ACAO1"
    ws["F4"] = "Acao desc A"
    merge_cells(ws, 4, 5, 5, 5)
    merge_cells(ws, 4, 6, 5, 6)
    ws["G4"] = "PTRES1"
    ws["H4"] = "FONTE1"
    ws["I4"] = "Fonte A"
    ws["J4"] = "PO1"
    ws["K4"] = "PO A"
    ws["L4"] = 1000
    ws["M4"] = 1200
    ws["N4"] = 50

    ws["A5"] = 2
    ws["B5"] = "Iduso B"
    ws["C5"] = 10
    ws["D5"] = "RP A"
    ws["G5"] = "PTRES2"
    ws["H5"] = "FONTE2"
    ws["I5"] = "Fonte B"
    ws["J5"] = "PO2"
    ws["K5"] = "PO B"
    ws["L5"] = None
    ws["M5"] = 0
    ws["N5"] = -30



def write_recognized_sheet_2025(ws) -> None:
    """Variante da base válida, com ano e valores diferentes do fixture padrão."""

    ws["A1"] = "Iduso"
    merge_cells(ws, 1, 1, 3, 2)
    ws["C1"] = "Resultado Primário Lei"
    merge_cells(ws, 1, 3, 3, 4)
    ws["E1"] = "Ação Governo"
    merge_cells(ws, 1, 5, 3, 6)
    ws["G1"] = "PTRES"
    merge_cells(ws, 1, 7, 3, 7)
    ws["H1"] = "Fonte Recursos Detalhada"
    merge_cells(ws, 1, 8, 3, 9)
    ws["J1"] = "Item Informação"
    merge_cells(ws, 1, 10, 1, 11)
    ws["J2"] = "Ano Lançamento"
    merge_cells(ws, 2, 10, 2, 11)
    ws["J3"] = "Plano Orçamentário"
    merge_cells(ws, 3, 10, 3, 11)

    ws["L1"] = "DOTACAO ATUALIZADA"
    ws["L2"] = 2025
    ws["L3"] = METRIC

    ws["A4"] = 1
    ws["B4"] = "Iduso A"
    ws["C4"] = 10
    ws["D4"] = "RP A"
    ws["E4"] = "ACAO9"
    ws["F4"] = "Acao desc Z"
    ws["G4"] = "PTRES9"
    ws["H4"] = "FONTE9"
    ws["I4"] = "Fonte Z"
    ws["J4"] = "PO9"
    ws["K4"] = "PO Z"
    ws["L4"] = 5000


def write_invalid_sheet(ws) -> None:
    """Estrutura reconhecida, mas com valor monetário não numérico."""

    write_recognized_sheet(ws)
    ws["L4"] = "não é número"


def workbook_bytes(builder) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Base"
    builder(worksheet)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def load_sheet(content: bytes, sheet_name: str = "Base"):
    return load_workbook(BytesIO(content), data_only=True)[sheet_name]


class TesouroDotacaoAnualTests(unittest.TestCase):
    def test_recognizes_valid_structure_without_optional_block(self) -> None:
        content = workbook_bytes(write_recognized_sheet)
        worksheet = load_sheet(content)

        recognition = recognize_dotacao_anual_base(worksheet)

        self.assertTrue(recognition.matched)
        self.assertEqual(recognition.validation_failures, [])

    def test_normalizes_resolving_real_merges_and_preserves_value_states(self) -> None:
        content = workbook_bytes(write_recognized_sheet)
        worksheet = load_sheet(content)

        result = transform_dotacao_anual_sheet(worksheet, "base.xlsx", "Base")

        self.assertTrue(result.recognized)
        self.assertEqual(result.raw_row_count, 2)
        self.assertEqual(result.normalized_row_count, 6)
        self.assertEqual(result.anos_encontrados, [2024])

        data = result.normalized_data
        self.assertEqual(sorted(data["linha_origem"].unique()), [4, 5])

        row4 = data[(data["linha_origem"] == 4) & (data["coluna_origem"] == "L")].iloc[0]
        self.assertEqual(row4["acao_codigo"], "ACAO1")
        self.assertEqual(row4["acao_descricao"], "Acao desc A")
        self.assertEqual(row4["item_informacao_codigo"], "dotacao_inicial")
        self.assertEqual(float(row4["valor_movimento_liquido"]), 1000.0)
        self.assertTrue(data.loc[row4.name, ["grupo_despesa_codigo", "grupo_despesa_descricao"]].isna().all())

        row5_inherits_acao = data[(data["linha_origem"] == 5) & (data["coluna_origem"] == "L")].iloc[0]
        self.assertEqual(row5_inherits_acao["acao_codigo"], "ACAO1")
        self.assertEqual(row5_inherits_acao["acao_descricao"], "Acao desc A")
        self.assertEqual(row5_inherits_acao["iduso_codigo"], "2")
        self.assertTrue(pd.isna(row5_inherits_acao["valor_movimento_liquido"]))

        row5_zero = data[(data["linha_origem"] == 5) & (data["coluna_origem"] == "M")].iloc[0]
        self.assertEqual(float(row5_zero["valor_movimento_liquido"]), 0.0)

        row5_negative = data[(data["linha_origem"] == 5) & (data["coluna_origem"] == "N")].iloc[0]
        self.assertEqual(float(row5_negative["valor_movimento_liquido"]), -30.0)
        self.assertEqual(row5_negative["item_informacao_codigo"], "item_nao_mapeado_item_experimental")
        self.assertTrue(
            any("item_nao_mapeado_item_experimental" in warning for warning in result.warnings)
        )

    def test_does_not_recognize_sheet_without_anchor(self) -> None:
        def builder(ws) -> None:
            ws["A1"] = "Coluna qualquer"
            ws["A4"] = "valor"

        content = workbook_bytes(builder)
        worksheet = load_sheet(content)

        recognition = recognize_dotacao_anual_base(worksheet)

        self.assertFalse(recognition.matched)
        self.assertTrue(recognition.validation_failures)

    def test_does_not_recognize_unknown_dimension_block(self) -> None:
        def builder(ws) -> None:
            ws["A1"] = "Bloco Desconhecido"
            merge_cells(ws, 1, 1, 3, 2)
            ws["C1"] = "Item Informação"
            merge_cells(ws, 1, 3, 1, 4)
            ws["C2"] = "Ano Lançamento"
            merge_cells(ws, 2, 3, 2, 4)
            ws["C3"] = "Plano Orçamentário"
            merge_cells(ws, 3, 3, 3, 4)
            ws["E1"] = "DOTACAO INICIAL"
            ws["E2"] = 2024
            ws["E3"] = METRIC
            ws["A4"] = "X"

        content = workbook_bytes(builder)
        worksheet = load_sheet(content)

        recognition = recognize_dotacao_anual_base(worksheet)

        self.assertFalse(recognition.matched)
        self.assertTrue(
            any("desconhecido" in failure for failure in recognition.validation_failures)
        )

    def test_rejects_non_xlsx_extension(self) -> None:
        with self.assertRaises(ValueError):
            read_dotacao_anual_sheet(b"conteudo", "base.xls", "Base")


if __name__ == "__main__":
    unittest.main()
