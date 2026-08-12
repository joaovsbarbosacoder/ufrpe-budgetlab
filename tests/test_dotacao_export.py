"""Testes da exportação auditável da análise de Dotação."""

from datetime import datetime
from io import BytesIO
import unittest

from openpyxl import load_workbook
import pandas as pd

from src.dotacao_export import EXPORT_SHEET_NAMES, build_dotacao_analysis_excel
from src.tesouro_dotacao import transform_dotacao_sheet
from tests.test_tesouro_dotacao import make_dotacao_raw


class DotacaoExportTests(unittest.TestCase):
    def setUp(self) -> None:
        self.normalized = transform_dotacao_sheet(
            make_dotacao_raw(),
            "dotacao.xlsx",
            "2024",
        ).normalized_data
        self.generated_at = datetime(2026, 8, 9, 10, 30, 0)

    def export_workbook(self, active_filters=None):
        original = self.normalized.copy(deep=True)
        content = build_dotacao_analysis_excel(
            self.normalized,
            active_filters=(
                {"periodo": ["MAR/2024"]}
                if active_filters is None
                else active_filters
            ),
            validation_status="Aprovada",
            selected_indicator_code="creditos_adicionais_excesso_arrecadacao",
            generated_at=self.generated_at,
        )
        pd.testing.assert_frame_equal(self.normalized, original)
        return load_workbook(BytesIO(content), data_only=True)

    def test_creates_sheets_in_requested_order(self) -> None:
        workbook = self.export_workbook()

        self.assertEqual(workbook.sheetnames, EXPORT_SHEET_NAMES)

    def test_metadata_contains_audit_context(self) -> None:
        workbook = self.export_workbook()
        sheet = workbook["Metadados"]
        metadata = {
            row[0]: row[1]
            for row in sheet.iter_rows(min_row=2, values_only=True)
            if row[0] is not None
        }

        self.assertEqual(metadata["Status da validação da normalização"], "Aprovada")
        self.assertEqual(metadata["Arquivo(s) de origem"], "dotacao.xlsx")
        self.assertEqual(metadata["Aba(s) de origem"], "2024")
        self.assertIn(
            "creditos_adicionais_excesso_arrecadacao",
            metadata["Indicador selecionado"],
        )
        self.assertEqual(metadata["Data/hora de geração"], self.generated_at)
        self.assertEqual(metadata["Quantidade de registros"], 2)
        self.assertEqual(metadata["Soma do recorte exportado"], -50.0)
        self.assertEqual(metadata["Filtro ativo: Período"], "MAR/2024")
        self.assertIn("fonte original não foi alterada", metadata["Natureza da extração"])

    def test_memory_reconciles_and_keeps_requested_traceability_fields(self) -> None:
        workbook = self.export_workbook()
        sheet = workbook["Memória de Cálculo"]
        headers = [cell.value for cell in sheet[1]]
        rows = list(sheet.iter_rows(min_row=2, values_only=True))
        row_by_line = {row[headers.index("linha_origem")]: row for row in rows}

        for required_header in [
            "item_informacao_origem",
            "item_informacao_codigo",
            "arquivo_origem",
            "aba_origem",
            "linha_origem",
            "coluna_origem",
        ]:
            self.assertIn(required_header, headers)

        self.assertEqual(len(rows), 2)
        self.assertEqual(row_by_line[4][headers.index("valor_movimento_liquido")], -50.0)
        self.assertEqual(row_by_line[5][headers.index("valor_movimento_liquido")], 0.0)
        self.assertEqual(row_by_line[4][headers.index("arquivo_origem")], "dotacao.xlsx")
        self.assertEqual(row_by_line[4][headers.index("aba_origem")], "2024")
        self.assertEqual(row_by_line[4][headers.index("coluna_origem")], "P")

    def test_preserves_codes_and_null_zero_negative_distinction(self) -> None:
        workbook = self.export_workbook(active_filters={})
        sheet = workbook["Detalhamento"]
        headers = [cell.value for cell in sheet[1]]
        header_index = {header: index + 1 for index, header in enumerate(headers)}
        records = {
            (
                row[header_index["linha_origem"] - 1],
                row[header_index["coluna_origem"] - 1],
            ): row
            for row in sheet.iter_rows(min_row=2, values_only=True)
        }

        self.assertEqual(
            records[(4, "N")][header_index["valor_movimento_liquido"] - 1], 0.0
        )
        self.assertIsNone(
            records[(4, "O")][header_index["valor_movimento_liquido"] - 1]
        )
        self.assertEqual(
            records[(4, "P")][header_index["valor_movimento_liquido"] - 1], -50.0
        )

        iduso_cell = sheet.cell(2, header_index["iduso_codigo"])
        action_cell = sheet.cell(2, header_index["acao_codigo"])
        source_cell = sheet.cell(2, header_index["fonte_recursos_detalhada_codigo"])
        self.assertEqual(iduso_cell.value, "0000")
        self.assertEqual(action_cell.value, "09HB")
        self.assertEqual(source_cell.value, "1234567890")
        self.assertEqual(iduso_cell.number_format, "@")
        self.assertEqual(action_cell.number_format, "@")
        self.assertEqual(source_cell.number_format, "@")


if __name__ == "__main__":
    unittest.main()
