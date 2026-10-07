"""Planos de referência da Captação: leitura, validação, carga versionada e busca."""

from __future__ import annotations

import io
import tempfile
import unittest
from pathlib import Path

import openpyxl

from src.captacao import planos, schema
from src.captacao.planos import ErroCarga

CAB_PDI = ("Código do objetivo", "Dimensão / eixo do PDI", "Descrição do objetivo")
CAB_PLS = ("Eixo temático", "Objetivo", "Código da meta", "Descrição da meta")


def _pdi(*linhas):
    return planos.validar_linhas_pdi([CAB_PDI, *linhas])


def _pls(*linhas):
    return planos.validar_linhas_pls([CAB_PLS, *linhas])


def _xlsx(pdi=None, pls=None) -> io.BytesIO:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    if pdi is not None:
        ws = wb.create_sheet("PDI")
        for linha in (CAB_PDI, *pdi):
            ws.append(linha)
    if pls is not None:
        ws = wb.create_sheet("PLS")
        for linha in (CAB_PLS, *pls):
            ws.append(linha)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


class ValidacaoTests(unittest.TestCase):
    def test_linhas_validas_preservam_codigo_como_texto_e_numero_da_linha(self):
        leitura = _pdi(("007", "Ensino", "Graduação"), ("A.1", None, "Pesquisa"))
        self.assertEqual(leitura.erros, [])
        self.assertEqual([l["codigo"] for l in leitura.linhas], ["007", "A.1"])
        self.assertEqual([l["numero_linha"] for l in leitura.linhas], [2, 3])
        self.assertEqual(leitura.linhas[1]["dimensao"], "")

    def test_codigo_numerico_vira_texto_com_aviso(self):
        leitura = _pdi((7, "Ensino", "Graduação"), (1.1, "Ensino", "Pós"))
        self.assertEqual([l["codigo"] for l in leitura.linhas], ["7", "1.1"])
        self.assertEqual(len(leitura.avisos), 2)

    def test_cabecalho_localizado_pelo_texto_e_nao_pela_posicao(self):
        embaralhado = [("Descrição do objetivo *", "Código do objetivo *"), ("Graduação", "O1")]
        leitura = planos.validar_linhas_pdi(embaralhado)
        self.assertEqual(leitura.erros, [])
        self.assertEqual(leitura.linhas[0]["codigo"], "O1")

    def test_coluna_obrigatoria_ausente_bloqueia_a_aba(self):
        leitura = planos.validar_linhas_pdi([("Código do objetivo",), ("O1",)])
        self.assertEqual(len(leitura.erros), 1)
        self.assertIn("Descrição do objetivo", leitura.erros[0])
        self.assertEqual(leitura.linhas, [])

    def test_codigo_e_descricao_ausentes_sao_erros_com_numero_da_linha(self):
        leitura = _pdi((None, "Ensino", "Sem código"), ("O2", "Ensino", None))
        self.assertEqual(len(leitura.erros), 2)
        self.assertIn("linha 2", leitura.erros[0])
        self.assertIn("código ausente", leitura.erros[0])
        self.assertIn("linha 3", leitura.erros[1])
        self.assertIn("descrição ausente", leitura.erros[1])

    def test_codigo_duplicado_e_erro_e_nao_e_descartado_em_silencio(self):
        leitura = _pdi(("O1", "x", "Primeira"), ("O1", "x", "Segunda"))
        self.assertEqual(len(leitura.erros), 1)
        self.assertIn("duplicado", leitura.erros[0])
        self.assertIn("linha 2", leitura.erros[0])

    def test_marcadores_do_modelo_nao_substituidos_sao_recusados(self):
        leitura = _pdi(("[código]", "[dimensão]", "[Substituir pela lista oficial]"))
        self.assertEqual(len(leitura.erros), 1)
        self.assertIn("marcador", leitura.erros[0])
        leitura = _pdi(("O1", "x", "[Substituir pela lista oficial]"))
        self.assertIn("marcador", leitura.erros[0])

    def test_linhas_totalmente_vazias_sao_ignoradas_e_planilha_sem_linhas_e_erro(self):
        self.assertEqual(len(_pdi((None, None, None), ("O1", "x", "ok")).linhas), 1)
        self.assertEqual(_pdi((None, None, None)).erros, ["Aba PDI: nenhuma linha preenchida."])
        self.assertIn("vazia", planos.validar_linhas_pdi([]).erros[0])

    def test_pls_le_eixo_objetivo_codigo_e_descricao(self):
        leitura = _pls(("Eixo 1", "Obj 1", "1.1", "Meta um"))
        self.assertEqual(leitura.erros, [])
        self.assertEqual(
            {k: leitura.linhas[0][k] for k in ("eixo", "objetivo", "codigo", "descricao")},
            {"eixo": "Eixo 1", "objetivo": "Obj 1", "codigo": "1.1", "descricao": "Meta um"},
        )


class LeituraXlsxTests(unittest.TestCase):
    def test_le_as_duas_abas_em_memoria(self):
        pdi, pls = planos.ler_planilha(_xlsx(pdi=[("O1", "Ensino", "Graduação")], pls=[("E1", "Obj", "1.1", "Meta")]))
        self.assertEqual((len(pdi.linhas), len(pls.linhas)), (1, 1))

    def test_aba_ausente_vira_erro_nao_excecao(self):
        pdi, pls = planos.ler_planilha(_xlsx(pdi=[("O1", "Ensino", "Graduação")]))
        self.assertEqual(len(pdi.linhas), 1)
        self.assertEqual(pls.erros, ["A planilha não tem a aba 'PLS'."])

    def test_arquivo_original_nao_e_alterado(self):
        with tempfile.TemporaryDirectory() as pasta:
            caminho = Path(pasta) / "planos.xlsx"
            caminho.write_bytes(_xlsx(pdi=[("O1", "Ensino", "Graduação")], pls=[("E1", "Obj", "1.1", "Meta")]).read())
            antes = caminho.read_bytes()
            planos.ler_planilha(caminho)
            self.assertEqual(caminho.read_bytes(), antes)

    def test_modelo_com_marcadores_e_recusado_por_inteiro(self):
        pdi, pls = planos.ler_planilha(
            _xlsx(pdi=[("[código]", "[dimensão]", "[Substituir]")], pls=[("[eixo]", "[objetivo]", "[código]", "[Substituir]")])
        )
        self.assertTrue(pdi.erros and pls.erros)


class CargaTests(unittest.TestCase):
    def setUp(self):
        self.conexao = schema.conectar(":memory:")

    def _contar(self, tabela, **filtros):
        where = " AND ".join(f"{k} = ?" for k in filtros) or "1 = 1"
        return self.conexao.execute(f"SELECT COUNT(*) FROM {tabela} WHERE {where}", tuple(filtros.values())).fetchone()[0]

    def test_primeira_carga_insere_e_preserva_zeros_a_esquerda(self):
        resultado = planos.carregar_pdi(self.conexao, _pdi(("007", "Ensino", "Graduação"), ("A.1", "Pesquisa", "Ciência")), "PDI 2025")
        self.assertEqual((resultado.inseridos, resultado.atualizados, resultado.inalterados), (2, 0, 0))
        codigos = [r["codigo"] for r in self.conexao.execute("SELECT codigo FROM objetivo_pdi ORDER BY id")]
        self.assertEqual(codigos, ["007", "A.1"])

    def test_recarga_identica_e_idempotente(self):
        leitura = _pdi(("O1", "Ensino", "Graduação"), ("O2", "Pesquisa", "Ciência"))
        planos.carregar_pdi(self.conexao, leitura, "PDI 2025")
        ids_antes = [r["id"] for r in self.conexao.execute("SELECT id FROM objetivo_pdi ORDER BY id")]
        resultado = planos.carregar_pdi(self.conexao, leitura, "PDI 2025")
        self.assertEqual((resultado.inseridos, resultado.atualizados, resultado.inalterados), (0, 0, 2))
        self.assertEqual([r["id"] for r in self.conexao.execute("SELECT id FROM objetivo_pdi ORDER BY id")], ids_antes)

    def test_atualizacao_mantem_o_id_e_informa_a_descricao_alterada(self):
        planos.carregar_pdi(self.conexao, _pdi(("O1", "Ensino", "Antiga")), "PDI 2025")
        id_original = self.conexao.execute("SELECT id FROM objetivo_pdi").fetchone()[0]
        resultado = planos.carregar_pdi(self.conexao, _pdi(("O1", "Ensino", "Nova")), "PDI 2025")
        self.assertEqual(resultado.atualizados, 1)
        self.assertEqual(resultado.descricoes_alteradas, [("O1", "Antiga", "Nova")])
        linha = self.conexao.execute("SELECT id, descricao FROM objetivo_pdi").fetchone()
        self.assertEqual((linha["id"], linha["descricao"]), (id_original, "Nova"))

    def test_item_que_sumiu_e_desativado_e_nao_excluido_e_vinculo_sobrevive(self):
        planos.carregar_pdi(self.conexao, _pdi(("O1", "x", "Um"), ("O2", "x", "Dois")), "PDI 2025")
        self.conexao.execute("INSERT INTO unidade (id, ugr_codigo, nome) VALUES (1, '1', 'U')")
        self.conexao.execute("INSERT INTO ciclo (id, exercicio, criado_em) VALUES (1, 2027, 'x')")
        self.conexao.execute("INSERT INTO demanda (id, ciclo_id, unidade_id, criada_em, atualizada_em) VALUES (1, 1, 1, 'x', 'x')")
        id_o2 = self.conexao.execute("SELECT id FROM objetivo_pdi WHERE codigo = 'O2'").fetchone()[0]
        self.conexao.execute("INSERT INTO demanda_objetivo_pdi VALUES (1, ?)", (id_o2,))
        self.conexao.commit()

        resultado = planos.carregar_pdi(self.conexao, _pdi(("O1", "x", "Um")), "PDI 2025")
        self.assertEqual((resultado.desativados, resultado.inalterados), (1, 1))
        self.assertEqual(self._contar("objetivo_pdi"), 2)
        self.assertEqual(self._contar("objetivo_pdi", codigo="O2", ativo=0), 1)
        self.assertEqual(self._contar("demanda_objetivo_pdi"), 1)

        resultado = planos.carregar_pdi(self.conexao, _pdi(("O1", "x", "Um"), ("O2", "x", "Dois")), "PDI 2025")
        self.assertEqual((resultado.reativados, resultado.atualizados), (1, 1))
        self.assertEqual(self._contar("objetivo_pdi", codigo="O2", ativo=1), 1)

    def test_versoes_diferentes_convivem_com_o_mesmo_codigo(self):
        planos.carregar_pdi(self.conexao, _pdi(("O1", "x", "Antigo")), "PDI 2020")
        planos.carregar_pdi(self.conexao, _pdi(("O1", "x", "Atual")), "PDI 2025")
        self.assertEqual(self._contar("objetivo_pdi"), 2)
        self.assertEqual(self._contar("objetivo_pdi", versao_pdi="PDI 2020", ativo=1), 1)

    def test_erros_de_validacao_recusam_tudo_sem_gravar_nada(self):
        leitura = _pdi(("O1", "x", "Boa"), (None, "x", "Sem código"))
        with self.assertRaises(ErroCarga) as contexto:
            planos.carregar_pdi(self.conexao, leitura, "PDI 2025")
        self.assertIn("código ausente", contexto.exception.erros[0])
        self.assertEqual(self._contar("objetivo_pdi"), 0)

    def test_versao_obrigatoria(self):
        with self.assertRaises(ErroCarga):
            planos.carregar_pdi(self.conexao, _pdi(("O1", "x", "Um")), "  ")
        self.assertEqual(self._contar("objetivo_pdi"), 0)

    def test_carga_do_pls_guarda_eixo_e_objetivo(self):
        planos.carregar_pls(self.conexao, _pls(("Eixo 1", "Obj", "1.1", "Meta um")), "PLS 2025-2027")
        linha = self.conexao.execute("SELECT * FROM meta_pls").fetchone()
        self.assertEqual((linha["eixo"], linha["objetivo"], linha["codigo"], linha["versao_pls"]), ("Eixo 1", "Obj", "1.1", "PLS 2025-2027"))

    def test_pls_idempotente_e_desativa_ausentes(self):
        planos.carregar_pls(self.conexao, _pls(("E1", "O", "1.1", "A"), ("E1", "O", "1.2", "B")), "PLS")
        resultado = planos.carregar_pls(self.conexao, _pls(("E1", "O", "1.1", "A")), "PLS")
        self.assertEqual((resultado.inalterados, resultado.desativados), (1, 1))


class BuscaTests(unittest.TestCase):
    def setUp(self):
        self.conexao = schema.conectar(":memory:")
        planos.carregar_pdi(
            self.conexao,
            _pdi(
                ("O1", "Ensino", "Ampliar a permanência estudantil"),
                ("O10", "Gestão", "Modernizar a governança"),
                ("O2", "Pesquisa", "Fortalecer a produção científica"),
            ),
            "PDI",
        )
        planos.carregar_pls(
            self.conexao,
            _pls(
                ("Energia", "Eficiência", "1.1", "Reduzir o consumo de energia elétrica"),
                ("Resíduos", "Coleta", "2.1", "Ampliar a coleta seletiva"),
                ("Energia", "Eficiência", "1.2", "Instalar usina fotovoltaica"),
            ),
            "PLS",
        )

    def test_busca_por_codigo_poe_codigo_exato_primeiro(self):
        achados = planos.buscar_objetivos_pdi(self.conexao, "o1")
        self.assertEqual([r["codigo"] for r in achados], ["O1", "O10"])

    def test_busca_por_palavra_ignora_acento_e_caixa(self):
        self.assertEqual([r["codigo"] for r in planos.buscar_objetivos_pdi(self.conexao, "PERMANENCIA")], ["O1"])
        self.assertEqual([r["codigo"] for r in planos.buscar_objetivos_pdi(self.conexao, "producao cientifica")], ["O2"])

    def test_todas_as_palavras_precisam_aparecer(self):
        self.assertEqual(planos.buscar_objetivos_pdi(self.conexao, "producao governanca"), [])

    def test_busca_vazia_devolve_tudo_e_busca_nao_encontrada_devolve_vazio(self):
        self.assertEqual(len(planos.buscar_objetivos_pdi(self.conexao, "")), 3)
        self.assertEqual(planos.buscar_objetivos_pdi(self.conexao, "inexistente"), [])

    def test_busca_encontra_pela_dimensao(self):
        self.assertEqual([r["codigo"] for r in planos.buscar_objetivos_pdi(self.conexao, "gestao")], ["O10"])

    def test_metas_filtram_por_eixo_e_texto(self):
        self.assertEqual(len(planos.buscar_metas_pls(self.conexao, eixo="Energia")), 2)
        self.assertEqual([r["codigo"] for r in planos.buscar_metas_pls(self.conexao, "usina", eixo="Energia")], ["1.2"])
        self.assertEqual(planos.buscar_metas_pls(self.conexao, "usina", eixo="Resíduos"), [])

    def test_metas_buscam_por_codigo(self):
        self.assertEqual([r["codigo"] for r in planos.buscar_metas_pls(self.conexao, "2.1")], ["2.1"])

    def test_inativos_ficam_fora_por_padrao(self):
        planos.carregar_pdi(self.conexao, _pdi(("O1", "Ensino", "Ampliar a permanência estudantil")), "PDI")
        self.assertEqual([r["codigo"] for r in planos.buscar_objetivos_pdi(self.conexao, "")], ["O1"])
        self.assertEqual(len(planos.buscar_objetivos_pdi(self.conexao, "", apenas_ativos=False)), 3)

    def test_filtro_por_versao(self):
        planos.carregar_pdi(self.conexao, _pdi(("X1", "x", "Outra versão")), "PDI 2030")
        self.assertEqual([r["codigo"] for r in planos.buscar_objetivos_pdi(self.conexao, "", versao_pdi="PDI 2030")], ["X1"])


if __name__ == "__main__":
    unittest.main()
