"""
Testes da leitura e conferência em memória da planilha de controle de NCs (08/10/2026).

Decisão: a planilha nunca é gravada no banco; os `.xlsx` abaixo são mínimos, criados em memória
com openpyxl (nunca o arquivo real do usuário).
"""

from __future__ import annotations

import io
import unittest
from datetime import date, datetime
from decimal import Decimal

from openpyxl import Workbook, load_workbook

from src.teds_controle_nc import (
    conferir_controle,
    gerar_xlsx_conferencia,
    justificativa_sugerida,
    ler_controle_nc,
)
from src.teds_schema import conectar


def _xlsx(abas: dict[str, list[list[object]]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for nome, linhas in abas.items():
        ws = wb.create_sheet(nome)
        for linha in linhas:
            ws.append(linha)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


CAB_2016 = ["NC", "DATA", "UG EMITENTE", "NOME ÓRGÃO", "OFÍCIO", "FONTE", "PTRES", "ND",
            "TED  N.  TRANSFERENCIA", "VALOR RECEBIDO", "PROCESSO", "OBSERVAÇÃO"]
CAB_2023 = ["NC", "DATA", "UG EMITENTE", "OFÍCIO", "TED", "N.  TRANSFERENCIA", "VALOR",
            "PROCESSO", "OBSERVAÇÃO"]


class LeitorTests(unittest.TestCase):
    def test_layout_2016_com_continuacao(self):
        conteudo = _xlsx({"2016": [
            CAB_2016,
            ["2016NC700145", datetime(2016, 11, 16), 153173, "FNDE", None, 112915153, 108429,
             339018, 688033, 2400000, "23082.1/2016", None],
            [None, None, None, None, None, None, 108430, 339039, None, 200000, None, None],
            [None, None, None, None, None, None, 108431, 339030, None, 50000, None, None],
        ]})
        res = ler_controle_nc(conteudo)
        self.assertEqual(len(res.linhas), 3)
        self.assertEqual({l.nc for l in res.linhas}, {"2016NC700145"})
        self.assertEqual({l.ug_emitente for l in res.linhas}, {"153173"})
        self.assertEqual({l.ted for l in res.linhas}, {"688033"})
        self.assertEqual({l.transferencia for l in res.linhas}, {"688033"})
        self.assertEqual([l.valor for l in res.linhas],
                         [Decimal("2400000"), Decimal("200000"), Decimal("50000")])
        self.assertEqual(res.linhas[1].ptres, "108430")
        self.assertEqual(res.linhas[2].data, date(2016, 11, 16))
        self.assertEqual(res.linhas[0].linha_origem, 2)

    def test_layout_2023(self):
        conteudo = _xlsx({"2023": [
            CAB_2023,
            ["2023NC000076", datetime(2023, 1, 26), 152734, "nº 15/2023", 11926, "1AALOH",
             176562, "23082.001901/2023-56", "BOLSAS"],
        ]})
        (linha,) = ler_controle_nc(conteudo).linhas
        self.assertEqual(linha.ted, "11926")
        self.assertEqual(linha.transferencia, "1AALOH")
        self.assertEqual(linha.valor, Decimal("176562"))
        self.assertEqual(linha.oficio, "nº 15/2023")

    def test_layout_2026_zero_a_esquerda(self):
        conteudo = _xlsx({"2026": [
            ["NC", "DOCUMENTO", "DATA", "UNIDADE ORÇAMENTÁRIA", "OFÍCIO", "FONTE", "PTRES", "ND",
             "PI", "TED", "N.  TRANSFERENCIA", "VALOR RECEBIDO", "PROCESSO", "OBSERVAÇÃO"],
            ["2026NC000081", "2026RO000227", datetime(2026, 2, 11), "FUNDAÇÃO JOAQUIM NABUCO",
             "33/2026", "1000000000", "023541", "339036", "N08RTO940TN", None, "-", 11785.5,
             "23082.004601/2026-71", None],
        ]})
        (linha,) = ler_controle_nc(conteudo).linhas
        self.assertEqual(linha.fonte, "1000000000")
        self.assertEqual(linha.ptres, "023541")
        self.assertEqual(linha.pi, "N08RTO940TN")
        self.assertEqual(linha.orgao, "FUNDAÇÃO JOAQUIM NABUCO")
        self.assertIsNone(linha.transferencia)
        self.assertEqual(linha.valor, Decimal("11785.50"))

    def test_traco_vira_nulo(self):
        conteudo = _xlsx({"2025": [
            CAB_2016,
            ["2025NC000001", datetime(2025, 1, 20), 156687, "UFAPE", "-", "-", 231379, 339037,
             "-", 57101.55, "23082.000697/2025-18", None],
        ]})
        (linha,) = ler_controle_nc(conteudo).linhas
        self.assertIsNone(linha.oficio)
        self.assertIsNone(linha.fonte)
        self.assertIsNone(linha.ted)
        self.assertIsNone(linha.transferencia)

    def test_linha_sem_valor_rejeitada_com_motivo(self):
        conteudo = _xlsx({"2023": [
            CAB_2023,
            ["2023NC000076", datetime(2023, 1, 26), 152734, "x", 11926, "1AALOH", None, None, None],
            ["2023NC000077", datetime(2023, 1, 26), 152734, "x", 11926, "1AALOH", 35313, None, None],
        ]})
        res = ler_controle_nc(conteudo)
        self.assertEqual(len(res.linhas), 1)
        self.assertEqual(len(res.rejeitadas), 1)
        aba, linha, motivo = res.rejeitadas[0]
        self.assertEqual((aba, linha), ("2023", 2))
        self.assertIn("valor", motivo.lower())

    def test_aba_nao_reconhecida_listada(self):
        conteudo = _xlsx({
            "2021": [[None] * 3, [None] * 3, [None, "NC", "DATA", "OFÍCIO"],
                     [None, 62, datetime(2021, 1, 28), "20/2021"]],
            "Resumo": [["qualquer", "coisa"]],
            "2023": [CAB_2023, ["2023NC000076", datetime(2023, 1, 26), 152734, "x", 11926,
                                "1AALOH", 1, None, None]],
        })
        res = ler_controle_nc(conteudo)
        self.assertEqual({a for a, _ in res.abas_ignoradas}, {"2021", "Resumo"})
        self.assertTrue(all(motivo for _, motivo in res.abas_ignoradas))
        self.assertEqual(len(res.linhas), 1)


def _banco():
    conn = conectar(":memory:")
    conn.execute(
        "INSERT INTO import_batch (id, tipo_relatorio, nome_arquivo, hash_arquivo, data_importacao, status) "
        "VALUES (1, 'teste', 'x.xlsx', 'hash', '2026-01-01T00:00:00', 'ok')"
    )
    return conn


def _nc(conn, numero, chave_ted, valor, data="2023-01-26", ug=None, operacao="+"):
    ted = chave_ted.split("|")[0]
    conn.execute(
        "INSERT INTO documento_nc (chave_nc_documento, chave_ted, ted, codigo_siafi, numero_nc, "
        "ug_emitente, data_emissao, operacao, valor_original_total, valor_assinado_total, "
        "quantidade_linhas, status_relacionamento, import_batch_id) "
        "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 1, 'ok', 1)",
        (f"{chave_ted}|{numero}|{operacao}|{data}|{ug}", chave_ted, ted, chave_ted.split("|")[1],
         numero, ug, data, operacao, valor, valor if operacao == "+" else f"-{valor}"),
    )


def _consolidado(conn, chave_ted, ano):
    zero = "0.00"
    conn.execute(
        "INSERT INTO execucao_anual (chave_ted, ano_emissao, total_nc_descentralizacao, "
        "total_nc_devolucao, total_descentralizado, total_pf_repasse, total_pf_devolucao, "
        "total_repassado, import_batch_id, linha_origem) VALUES (?, ?, ?, ?, ?, ?, ?, ?, 1, '{}')",
        (chave_ted, ano, zero, zero, zero, zero, zero, zero),
    )


def _planilha(linhas):
    return ler_controle_nc(_xlsx({"2023": [CAB_2023, *linhas]}))


def _contagens(conn):
    tabelas = [r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")]
    return {t: conn.execute(f'SELECT COUNT(*) FROM "{t}"').fetchone()[0] for t in tabelas}


class ConferenciaTests(unittest.TestCase):
    def setUp(self):
        self.conn = _banco()

    def tearDown(self):
        self.conn.close()

    def test_confere_e_diverge(self):
        _nc(self.conn, "2023NC000076", "11926|1AALOH", "176562.00")
        _nc(self.conn, "2023NC000077", "11926|1AALOH", "46097.73")
        leitura = _planilha([
            ["2023NC000076", datetime(2023, 1, 26), 152734, "x", 11926, "1AALOH", 176562, None, None],
            ["2023NC000077", datetime(2023, 1, 26), 152734, "x", 11926, "1AALOH", 102600.50, None, None],
        ])
        conf = conferir_controle(self.conn, leitura)
        por_nc = {c.numero_nc: c for c in conf.por_nc}
        c1 = por_nc["2023NC000076"]
        self.assertEqual(c1.situacao, "confere")
        self.assertEqual(c1.ug_planilha, "152734")
        self.assertIsNone(c1.ug_simec)
        self.assertEqual(c1.valor_simec, Decimal("176562.00"))
        self.assertEqual(c1.valor_planilha, Decimal("176562.00"))
        c2 = por_nc["2023NC000077"]
        self.assertEqual(c2.situacao, "diverge")
        self.assertEqual(c2.valor_simec, Decimal("46097.73"))
        self.assertEqual(c2.valor_planilha, Decimal("102600.50"))

    def test_sem_planilha_e_soma_de_linhas_da_nc(self):
        _nc(self.conn, "2023NC000001", "11926|1AALOH", "300.00")
        _nc(self.conn, "2023NC000002", "11926|1AALOH", "10.00")
        leitura = _planilha([
            ["2023NC000001", datetime(2023, 1, 26), 152734, "x", 11926, "1AALOH", 100, None, None],
            [None, None, None, None, None, None, 200, None, None],
        ])
        por_nc = {c.numero_nc: c for c in conferir_controle(self.conn, leitura).por_nc}
        self.assertEqual(por_nc["2023NC000001"].situacao, "confere")
        self.assertEqual(por_nc["2023NC000002"].situacao, "sem_planilha")

    def test_ug_diferente_nao_liga(self):
        _nc(self.conn, "2023NC000001", "11926|1AALOH", "100.00", ug="111111")
        leitura = _planilha([
            ["2023NC000001", datetime(2023, 1, 26), 152734, "x", 11926, "1AALOH", 100, None, None],
        ])
        (c,) = conferir_controle(self.conn, leitura).por_nc
        self.assertEqual(c.situacao, "sem_planilha")

    def test_nc_ambigua_nao_e_comparada(self):
        _nc(self.conn, "2024NC000036", "13023|1AAQRN", "500.00")
        leitura = _planilha([
            ["2024NC000036", datetime(2024, 3, 1), 152734, "x", 13023, "1AAQRN", 300, None, None],
            ["2024NC000036", datetime(2024, 3, 1), 153033, "x", 13023, "1AAQRN", 200, None, None],
        ])
        (c,) = conferir_controle(self.conn, leitura).por_nc
        self.assertEqual(c.situacao, "ambigua")
        self.assertIsNone(c.ug_planilha)
        self.assertIsNone(c.valor_planilha)

    def test_mesma_nc_duas_vezes_no_simec_e_ambigua(self):
        _nc(self.conn, "2024NC000036", "13023|1AAQRN", "300.00", data="2024-03-01")
        _nc(self.conn, "2024NC000036", "13023|1AAQRN", "200.00", data="2024-03-02")
        leitura = _planilha([
            ["2024NC000036", datetime(2024, 3, 1), 152734, "x", 13023, "1AAQRN", 500, None, None],
        ])
        self.assertEqual({c.situacao for c in conferir_controle(self.conn, leitura).por_nc},
                         {"ambigua"})

    def test_nc_anterior_ao_consolidado(self):
        _consolidado(self.conn, "11460|1AAKLM", 2023)
        _consolidado(self.conn, "11460|1AAKLM", 2024)
        leitura = _planilha([
            ["2022NC000500", datetime(2022, 12, 20), 152734, "x", 11460, "1AAKLM", 800000, None, None],
            ["2023NC000010", datetime(2023, 2, 1), 152734, "x", 11460, "1AAKLM", 10, None, None],
        ])
        conf = conferir_controle(self.conn, leitura)
        self.assertEqual([l.nc for l in conf.nc_anterior_consolidado["11460|1AAKLM"]],
                         ["2022NC000500"])

    def test_sem_ted(self):
        leitura = _planilha([
            ["2023NC000900", datetime(2023, 1, 26), 152734, "x", None, None, 5, None, None],
        ])
        conf = conferir_controle(self.conn, leitura)
        self.assertEqual([l.nc for l in conf.sem_ted], ["2023NC000900"])

    def test_justificativa_sugerida_pf_maior_que_nc(self):
        _consolidado(self.conn, "11460|1AAKLM", 2023)
        leitura = _planilha([
            ["2022NC000500", datetime(2022, 12, 20), 152734, "x", 11460, "1AAKLM", 800000, None, None],
        ])
        conf = conferir_controle(self.conn, leitura)
        texto = justificativa_sugerida("pf_liquida_maior_que_nc", "11460|1AAKLM", conf)
        self.assertIn("Planilha de controle", texto)
        self.assertIn("2022NC000500", texto)
        self.assertIn("20/12/2022", texto)
        self.assertIn("R$ 800.000,00", texto)
        self.assertIn("152734", texto)
        self.assertIsNone(justificativa_sugerida("pf_liquida_maior_que_nc", "99999|X", conf))
        self.assertIsNone(justificativa_sugerida("tipo_desconhecido", "11460|1AAKLM", conf))

    def test_conferir_nao_grava(self):
        _nc(self.conn, "2023NC000076", "11926|1AALOH", "176562.00")
        _consolidado(self.conn, "11926|1AALOH", 2023)
        self.conn.commit()
        antes = _contagens(self.conn)
        leitura = _planilha([
            ["2023NC000076", datetime(2023, 1, 26), 152734, "x", 11926, "1AALOH", 176562, None, None],
        ])
        conferir_controle(self.conn, leitura)
        self.assertEqual(_contagens(self.conn), antes)
        self.assertFalse(self.conn.in_transaction)

    def test_xlsx_tem_as_abas(self):
        conf = conferir_controle(self.conn, _planilha([
            ["2023NC000076", datetime(2023, 1, 26), 152734, "x", 11926, "1AALOH", 1, None, None],
        ]))
        wb = load_workbook(io.BytesIO(gerar_xlsx_conferencia(conf)))
        self.assertEqual(wb.sheetnames, ["Conferência por NC", "NC anterior ao consolidado",
                                         "Sem TED", "Não lido"])


if __name__ == "__main__":
    unittest.main()
