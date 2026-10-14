"""Testes do relatório separado "Projeção pela Execução" (`src/relatorio_projecao_execucao.py`).
Dados sintéticos; valores esperados calculados à mão em cada caso. O relatório só LÊ a saída de
`necessidade_por_ne` — nenhum outro relatório ou quadro da página muda (ver docstring do módulo)."""

from __future__ import annotations

import unittest
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from src.relatorio_necessidade_empenho import necessidade_por_ne
from src.relatorio_projecao_execucao import (
    BASE_COMPETENCIA,
    LEGENDA_GRADE,
    liquidado_abatido,
    BASE_LANCAMENTO,
    ORIGEM_EXECUCAO,
    ORIGEM_HERDADO,
    ORIGEM_SEM_HISTORICO,
    TIPO_PROJETADO,
    TIPO_RESTANTE_ABERTO,
    COLUNAS_LINHAS,
    ContextoProjecaoExecucao,
    FatorExecucao,
    fator_de_execucao,
    gerar_pdf,
    _LARGURAS_RESUMO_PDF,
    _cabecalho_resumo_pdf,
    gerar_xlsx,
    herdar_fator,
    montar_relatorio,
    projetar_meses,
    ultimo_mes_fechado,
)

CUSTO_1000 = [1000.0] * 12


class TestFatorDeExecucao(unittest.TestCase):
    def test_ultimo_mes_fechado(self):
        self.assertEqual(ultimo_mes_fechado(10), 7)  # extração em outubro: fechados até julho

    def test_execucao_constante_peso_1(self):
        # 0,90 em todos os meses: variância 0 → peso 1, fator = execução observada
        fator = fator_de_execucao({m: 900.0 for m in range(1, 8)}, CUSTO_1000, 7)
        self.assertEqual(fator.origem, ORIGEM_EXECUCAO)
        self.assertEqual(fator.meses_usados, (2, 3, 4, 5, 6, 7))  # janela: últimos 6 fechados
        self.assertAlmostEqual(fator.peso, 1.0)
        self.assertAlmostEqual(fator.fator, 0.9)

    def test_peso_cai_com_a_oscilacao(self):
        # razões 0,6/1,0/0,8/1,0/0,6/0,8 → observada 4800/6000 = 0,8; média 0,8;
        # variância amostral = (0,04+0,04+0+0,04+0,04+0)/5 = 0,032; peso = 6·0,01/(0,06+0,032) = 0,652174;
        # fator = 1 − 0,2 × 0,652174 = 0,869565
        valores = {2: 600.0, 3: 1000.0, 4: 800.0, 5: 1000.0, 6: 600.0, 7: 800.0}
        fator = fator_de_execucao(valores, CUSTO_1000, 7)
        self.assertAlmostEqual(fator.execucao_observada, 0.8)
        self.assertAlmostEqual(fator.peso, 0.06 / 0.092, places=6)
        self.assertAlmostEqual(fator.fator, 1 - 0.2 * 0.06 / 0.092, places=6)

    def test_meses_abertos_e_custo_zero_nao_entram(self):
        # ago/set (abertos) e o mês sem custo ficam de fora; 2 meses fechados → sem histórico
        custo = [0.0] * 4 + [1000.0] * 8
        fator = fator_de_execucao({4: 500.0, 6: 900.0, 7: 900.0, 8: 100.0, 9: 100.0}, custo, 7)
        self.assertEqual(fator.origem, ORIGEM_SEM_HISTORICO)
        self.assertEqual(fator.fator, 1.0)
        self.assertEqual(fator.meses_usados, (6, 7))

    def test_estorno_entra_como_veio(self):
        fator = fator_de_execucao({5: 1000.0, 6: -200.0, 7: 1000.0}, CUSTO_1000, 7)
        self.assertAlmostEqual(fator.execucao_observada, 1800.0 / 3000.0)
        self.assertLess(fator.peso, 0.1)  # oscilação enorme → quase valor cheio

    def test_heranca(self):
        antecessor = FatorExecucao(0.94, 0.97, 0.937, (2, 3, 4, 5, 6, 7), ORIGEM_EXECUCAO)
        sem = FatorExecucao(1.0, 0.0, None, (6, 7), ORIGEM_SEM_HISTORICO)
        herdado = herdar_fator(sem, antecessor, "2026NE000095")
        self.assertEqual((herdado.origem, herdado.fator, herdado.ne_antecessora), (ORIGEM_HERDADO, 0.94, "2026NE000095"))
        proprio = FatorExecucao(0.8, 1.0, 0.8, (5, 6, 7), ORIGEM_EXECUCAO)
        self.assertIs(herdar_fator(proprio, antecessor, "X"), proprio)  # já tem histórico: não herda
        self.assertIs(herdar_fator(sem, sem, "X"), sem)  # antecessor sem histórico: não herda


class TestProjetarMeses(unittest.TestCase):
    def test_mes_aberto_parcial_sem_registro_e_futuros(self):
        # extração em outubro; fator 0,9 → esperado 900/mês. Ago: 100 (< 450) → restante 800;
        # Set: sem registro → 900; Out: 200 já liquidado → 700; Nov/Dez 900. Total 4.200.
        # Valor cheio nos mesmos meses: 900 + 1000 + 800 + 1000 + 1000 = 4.700.
        realizado = {m: 900.0 for m in range(1, 8)} | {8: 100.0, 10: 200.0}
        projecao = projetar_meses(realizado, CUSTO_1000, 0.9, 10)
        self.assertEqual(projecao.projetado[:7], [None] * 7)
        self.assertAlmostEqual(projecao.projetado[7], 800.0)
        self.assertAlmostEqual(projecao.projetado[8], 900.0)
        self.assertAlmostEqual(projecao.projetado[9], 700.0)
        self.assertEqual(projecao.tipos[7], TIPO_RESTANTE_ABERTO)
        self.assertEqual(projecao.tipos[9], TIPO_PROJETADO)
        self.assertAlmostEqual(projecao.total_projetado, 4200.0)
        self.assertAlmostEqual(projecao.total_valor_cheio, 4700.0)
        self.assertEqual(projecao.restante_em_aberto(), [(8, 800.0), (9, 900.0)])

    def test_mes_aberto_com_metade_ou_mais_e_completo(self):
        projecao = projetar_meses({8: 450.0, 9: 900.0}, CUSTO_1000, 0.9, 10)
        self.assertIsNone(projecao.projetado[7])
        self.assertIsNone(projecao.projetado[8])
        self.assertEqual(projecao.tipos[7], "Realizado")

    def test_fora_da_execucao_e_custo_desconhecido_nao_projetam(self):
        custo = [1000.0] * 10 + [0.0, float("nan")]
        projecao = projetar_meses({}, custo, 1.0, 10)
        self.assertAlmostEqual(projecao.projetado[9], 1000.0)
        self.assertIsNone(projecao.projetado[10])
        self.assertIsNone(projecao.projetado[11])

    def test_mes_da_extracao_ja_liquidado_por_inteiro_fica_realizado(self):
        # liquidado 1.000 = esperado 1.000 em out (mês da extração): nada a projetar, o mês aparece como Realizado
        projecao = projetar_meses({10: 1000.0}, CUSTO_1000, 1.0, 10)
        self.assertEqual(projecao.tipos[9], "Realizado")
        self.assertEqual(projecao.realizado[9], 1000.0)
        self.assertIsNone(projecao.projetado[9])
        self.assertEqual(projecao.total_projetado, 4000.0)  # ago e set (abertos) + nov e dez seguem projetados

    def test_mes_da_extracao_liquidado_parcial_continua_projetando_o_restante(self):
        projecao = projetar_meses({10: 400.0}, CUSTO_1000, 1.0, 10)
        self.assertEqual(projecao.tipos[9], TIPO_PROJETADO)
        self.assertEqual(projecao.projetado[9], 600.0)

    def test_teto_limita_o_esperado_do_mes(self):
        # custo 1.000, teto 700: ago e set (abertos, sem registro) 700 cada; out–dez 700 cada → 3.500; cheio sem teto
        projecao = projetar_meses({}, CUSTO_1000, 1.0, 10, teto_mensal=700.0)
        self.assertEqual(projecao.total_projetado, 3500.0)
        self.assertEqual(projecao.total_valor_cheio, 5000.0)
        self.assertTrue(projecao.limitado_pelo_teto)

    def test_teto_acima_do_esperado_nao_muda_nada(self):
        projecao = projetar_meses({}, CUSTO_1000, 0.5, 10, teto_mensal=900.0)
        self.assertEqual(projecao.total_projetado, 2500.0)
        self.assertFalse(projecao.limitado_pelo_teto)

    def test_exercicio_encerrado_nao_projeta(self):
        self.assertEqual(projetar_meses({}, CUSTO_1000, 1.0, 12).total_projetado, 0.0)


def _linha(**campos: object) -> dict:
    base = {
        "ne_curta": None, "fornecedor": "F", "contrato_numero": "1/2026", "despesa_mensal": 1000.0,
        "meses_no_ano": 12.0, "valor_empenhado": 0.0, "valor_empenhado_execucao": None,
        "valor_empenhado_planilha_total_ne": None, "valor_liquidado_execucao": None,
        "liquidado_via_competencia": True, "saldo_execucao": None, "saldo_colado_planilha": None,
        "status_contrato": "ATIVO", "vigencia_fim": pd.NaT,
    }
    base.update(campos)
    return base


NE_EST = "2026NE000001"  # estável a 0,90
NE_NOVA = "2026NE000002"  # sem histórico


def _filtrado() -> pd.DataFrame:
    return pd.DataFrame(
        [
            # empenhado 9.000, liquidado 6.300 (jan–jul a 900) → saldo 2.700
            _linha(ne_curta=NE_EST, fornecedor="ESTAVEL", contrato_numero="10/2025",
                   valor_empenhado_execucao=9000.0, valor_liquidado_execucao=6300.0),
            # empenhado 2.000, nada liquidado → saldo 2.000; início 01/09
            _linha(ne_curta=NE_NOVA, fornecedor="NOVA", contrato_numero="20/2026",
                   valor_empenhado_execucao=2000.0, valor_liquidado_execucao=0.0,
                   inicio_execucao_data=pd.Timestamp("2026-09-01")),
            _linha(ne_curta=None, fornecedor="SEM NE", valor_empenhado=0.0),
        ]
    )


def _liquidacao() -> pd.DataFrame:
    return pd.DataFrame(
        {"ne_curta": [NE_EST] * 7, "ano_mes": [202601 + i for i in range(7)], "valor": [900.0] * 7}
    )


def _relatorio(antecessores=None):
    por_ne, sem_ne = necessidade_por_ne(_filtrado(), None, 2026)
    return montar_relatorio(por_ne, sem_ne, _liquidacao(), 2026, 10, antecessores)


class TestMontarRelatorio(unittest.TestCase):
    def test_estavel(self):
        # fator 0,9; ago e set sem registro → 900 cada (restante em aberto 1.800); out–dez 2.700.
        # projetado 4.500; valor cheio 5.000; necessidade = 4.500 − 2.700 = 1.800.
        # contratual (Resumo) = 12.000 − 9.000 = 3.000.
        relatorio = _relatorio()
        linha = relatorio.linhas.set_index("ne_curta").loc[NE_EST]
        self.assertAlmostEqual(linha["fator"], 0.9)
        self.assertAlmostEqual(linha["restante_em_aberto"], 1800.0)
        self.assertAlmostEqual(linha["projetado_execucao"], 4500.0)
        self.assertAlmostEqual(linha["projetado_valor_cheio"], 5000.0)
        self.assertAlmostEqual(linha["necessidade_execucao"], 1800.0)
        self.assertAlmostEqual(linha["necessidade_contratual"], 3000.0)
        self.assertEqual(relatorio.qtd_sem_ne, 1)

    def test_nova_valor_cheio_e_com_antecessor(self):
        # sem histórico: set (aberto, sem registro) 1.000 + out–dez 3.000 = 4.000; necessidade 4.000 − 2.000
        linha = _relatorio().linhas.set_index("ne_curta").loc[NE_NOVA]
        self.assertEqual(linha["origem_fator"], ORIGEM_SEM_HISTORICO)
        self.assertAlmostEqual(linha["projetado_execucao"], 4000.0)
        self.assertAlmostEqual(linha["necessidade_execucao"], 2000.0)
        # herdando 0,9 do antecessor: 3.600; necessidade 1.600
        herdada = _relatorio({NE_NOVA: NE_EST}).linhas.set_index("ne_curta").loc[NE_NOVA]
        self.assertEqual(herdada["origem_fator"], ORIGEM_HERDADO)
        self.assertEqual(herdada["ne_antecessora"], NE_EST)
        self.assertAlmostEqual(herdada["projetado_execucao"], 3600.0)
        self.assertAlmostEqual(herdada["necessidade_execucao"], 1600.0)

    def test_antecessor_invalido_e_avisado(self):
        relatorio = _relatorio({NE_EST: NE_NOVA})  # NE_EST já tem histórico
        self.assertEqual(relatorio.linhas.set_index("ne_curta").loc[NE_EST, "origem_fator"], ORIGEM_EXECUCAO)
        self.assertTrue(any("antecessor" in a and "ignorado" in a for a in relatorio.avisos))

    def test_ordem_totais_e_avisos(self):
        relatorio = _relatorio()
        self.assertEqual(relatorio.linhas["ne_curta"].tolist(), [NE_NOVA, NE_EST])  # necessidade decrescente
        self.assertAlmostEqual(relatorio.total_necessidade_execucao, 3800.0)
        self.assertAlmostEqual(relatorio.total_saldo, 4700.0)
        self.assertTrue(any("Meses em aberto" in a for a in relatorio.avisos))
        self.assertTrue(any("sem NE" in a for a in relatorio.avisos))
        self.assertTrue(any("valor mensal cheio" in a and NE_NOVA in a for a in relatorio.avisos))

    def test_sem_valor_mensal_necessidade_nao_calculada(self):
        # valor mensal desconhecido → custo nulo nos 12 meses (regra de `custo_mensal_bolsa`, passada em
        # `custos`): a necessidade fica nula ("—"), não 0, sai dos totais e é avisada; a NE com valor
        # continua com os números de sempre (1.800).
        filtrado = pd.DataFrame([
            _linha(ne_curta=NE_EST, fornecedor="ESTAVEL", contrato_numero="10/2025",
                   valor_empenhado_execucao=9000.0, valor_liquidado_execucao=6300.0),
            _linha(ne_curta=NE_NOVA, fornecedor="SEM VALOR", contrato_numero="20/2026",
                   valor_empenhado_execucao=2000.0, valor_liquidado_execucao=0.0),
        ])
        por_ne, sem_ne = necessidade_por_ne(filtrado, None, 2026)
        custos = {NE_EST: CUSTO_1000, NE_NOVA: [float("nan")] * 12}
        relatorio = montar_relatorio(por_ne, sem_ne, _liquidacao(), 2026, 10, custos=custos)
        linhas = relatorio.linhas.set_index("ne_curta")
        for coluna in ("necessidade_execucao", "projetado_execucao", "projetado_valor_cheio", "restante_em_aberto"):
            self.assertTrue(pd.isna(linhas.loc[NE_NOVA, coluna]), coluna)
        self.assertAlmostEqual(linhas.loc[NE_EST, "necessidade_execucao"], 1800.0)
        self.assertAlmostEqual(relatorio.total_necessidade_execucao, 1800.0)
        self.assertTrue(any("Sem valor mensal" in a and NE_NOVA in a for a in relatorio.avisos))
        self.assertEqual(relatorio.linhas["ne_curta"].tolist(), [NE_EST, NE_NOVA])  # nula vai para o fim

    def test_valor_mensal_zero_continua_zero(self):
        filtrado = pd.DataFrame([_linha(ne_curta=NE_NOVA, despesa_mensal=0.0, valor_empenhado_execucao=0.0,
                                        valor_liquidado_execucao=0.0)])
        por_ne, sem_ne = necessidade_por_ne(filtrado, None, 2026)
        linha = montar_relatorio(por_ne, sem_ne, None, 2026, 10).linhas.iloc[0]
        self.assertEqual(linha["necessidade_execucao"], 0.0)

    def test_aviso_abaixo_e_acima_do_cadastro(self):
        filtrado = pd.DataFrame([
            _linha(ne_curta=NE_EST, valor_empenhado_execucao=1000.0, valor_liquidado_execucao=700.0),
            _linha(ne_curta=NE_NOVA, valor_empenhado_execucao=1000.0, valor_liquidado_execucao=700.0),
        ])
        liquidacao = pd.DataFrame({
            "ne_curta": [NE_EST] * 7 + [NE_NOVA] * 7,
            "ano_mes": [202601 + i for i in range(7)] * 2,
            "valor": [100.0] * 7 + [1100.0] * 7,
        })
        por_ne, sem_ne = necessidade_por_ne(filtrado, None, 2026)
        avisos = montar_relatorio(por_ne, sem_ne, liquidacao, 2026, 10).avisos
        self.assertTrue(any("muito abaixo do cadastrado" in a and NE_EST in a for a in avisos))
        self.assertTrue(any("acima do contratado" in a and NE_NOVA in a for a in avisos))


class TestArquivos(unittest.TestCase):
    def setUp(self):
        self.relatorio = _relatorio()
        self.contexto = ContextoProjecaoExecucao(2026, "06/10/2026", "abcd1234", "06/10/2026 12:00", "Liquidação por Competência.xlsx")

    def test_pdf(self):
        self.assertTrue(gerar_pdf(self.relatorio, self.contexto).startswith(b"%PDF"))

    def test_xlsx(self):
        livro = load_workbook(BytesIO(gerar_xlsx(self.relatorio, self.contexto)))
        self.assertEqual(livro.sheetnames, ["Resumo por NE", "Grade mensal", "Detalhe mensal", "Parâmetros"])
        resumo = livro["Resumo por NE"]
        cabecalho = [c.value for c in resumo[1]]
        linha_est = next(r for r in resumo.iter_rows(min_row=2, values_only=True) if r[0] == NE_EST)
        self.assertAlmostEqual(linha_est[cabecalho.index("Necessidade pela execução")], 1800.0)
        self.assertAlmostEqual(linha_est[cabecalho.index("Fator de execução")], 0.9)
        grade = livro["Grade mensal"]
        cab_grade = [c.value for c in grade[1]]
        linha_grade = next(r for r in grade.iter_rows(min_row=2, values_only=True) if r[0] == NE_EST)
        self.assertEqual(linha_grade[cab_grade.index("Jan")], 900.0)  # realizado
        self.assertAlmostEqual(linha_grade[cab_grade.index("Ago")], 900.0)  # restante em aberto
        self.assertAlmostEqual(linha_grade[cab_grade.index("Out")], 900.0)  # projetado


class TestCabecalhoResumoPdf(unittest.TestCase):
    """Cabeçalho do resumo no PDF (13/10/2026): rótulos longos invadiam a coluna vizinha. Agora quebram de linha,
    então basta que a MAIOR PALAVRA de cada rótulo caiba na coluna (descontado o espaço lateral de 3 pt de cada lado)."""

    def test_maior_palavra_de_cada_rotulo_cabe_na_coluna(self):
        from reportlab.pdfbase.pdfmetrics import stringWidth

        rotulos = _relatorio().rotulos
        cabecalho = _cabecalho_resumo_pdf(rotulos)
        self.assertEqual(len(cabecalho), len(_LARGURAS_RESUMO_PDF))
        for rotulo, largura in zip(cabecalho, _LARGURAS_RESUMO_PDF):
            maior = max(stringWidth(palavra, "Helvetica-Bold", 6.5) for palavra in rotulo.split())
            self.assertLessEqual(maior, largura - 6, f"{rotulo!r} estoura a coluna de {largura} pt")

    def test_tabela_cabe_na_pagina_a4_paisagem(self):
        self.assertLessEqual(sum(_LARGURAS_RESUMO_PDF), 841.89 - 2 * 10 * 72 / 25.4)


class TestRelatorioSemNenhumaNE(unittest.TestCase):
    """Regressão (13/10/2026): busca em Contratos Contínuos que isola só contratos sem NE gerava `linhas` sem coluna
    alguma e quebrava `aplicar_projecao_na_necessidade` com KeyError: 'ne_curta'."""

    def _so_sem_ne(self):
        filtrado = _filtrado().iloc[[2]]  # só o contrato SEM NE
        por_ne, sem_ne = necessidade_por_ne(filtrado, None, 2026)
        return por_ne, montar_relatorio(por_ne, sem_ne, _liquidacao(), 2026, 10)

    def test_linhas_vazias_mantem_as_colunas(self):
        _, relatorio = self._so_sem_ne()
        self.assertTrue(relatorio.linhas.empty)
        self.assertEqual(list(relatorio.linhas.columns), COLUNAS_LINHAS)
        self.assertEqual(relatorio.qtd_sem_ne, 1)

    def test_aplicar_projecao_na_necessidade_aceita_relatorio_vazio(self):
        from src.relatorio_projecao_execucao import aplicar_projecao_na_necessidade

        por_ne, relatorio = self._so_sem_ne()
        saida = aplicar_projecao_na_necessidade(por_ne, relatorio)
        self.assertEqual(len(saida), len(por_ne))

    def test_colunas_fixas_batem_com_o_relatorio_cheio(self):
        self.assertEqual(list(_relatorio().linhas.columns), COLUNAS_LINHAS)


if __name__ == "__main__":
    unittest.main()


class TestBaseDaLiquidacao(unittest.TestCase):
    """13/10/2026: cada NE informa de onde vem o liquidado (competência ou mês do lançamento)."""

    def _montar(self, **extra):
        por_ne, sem_ne = necessidade_por_ne(_filtrado(), None, 2026)
        return montar_relatorio(por_ne, sem_ne, _liquidacao(), 2026, 10, **extra)

    def test_padrao_e_competencia(self):
        self.assertEqual(set(self._montar().linhas["base_liquidacao"]), {BASE_COMPETENCIA})

    def test_ne_informada_vira_lancamento_sem_mudar_a_conta(self):
        antes = self._montar().linhas.set_index("ne_curta")
        depois_rel = self._montar(nes_pelo_lancamento={NE_EST})
        depois = depois_rel.linhas.set_index("ne_curta")
        self.assertEqual(depois.loc[NE_EST, "base_liquidacao"], BASE_LANCAMENTO)
        self.assertEqual(depois.loc[NE_NOVA, "base_liquidacao"], BASE_COMPETENCIA)
        self.assertAlmostEqual(depois.loc[NE_EST, "necessidade_execucao"], antes.loc[NE_EST, "necessidade_execucao"])

    def test_xlsx_e_pdf_trazem_a_base(self):
        relatorio = self._montar(nes_pelo_lancamento={NE_EST})
        contexto = ContextoProjecaoExecucao(2026, "01/10/2026", "abcd1234", "13/10/2026 10:00", "arquivo")
        planilha = load_workbook(BytesIO(gerar_xlsx(relatorio, contexto)))
        textos = [str(c.value) for aba in planilha for linha in aba.iter_rows() for c in linha if c.value]
        self.assertIn("Base da liquidação", textos)
        self.assertIn(BASE_LANCAMENTO, textos)
        self.assertTrue(gerar_pdf(relatorio, contexto).startswith(b"%PDF"))


class TestTetoSemHistorico(unittest.TestCase):
    """14/10/2026: NE sem histórico fechado não projeta acima do maior mês já liquidado (só com a opção ligada)."""

    def _montar(self, liquidacao, **extra):
        por_ne, sem_ne = necessidade_por_ne(_filtrado(), None, 2026)
        return montar_relatorio(por_ne, sem_ne, liquidacao, 2026, 10, **extra)

    def _liquidacao_nova(self):
        # NE_NOVA liquidou só em ago/2026: 600 (< custo 1.000); sem 3 meses fechados → sem histórico
        extra = pd.DataFrame({"ne_curta": [NE_NOVA], "ano_mes": [202608], "valor": [600.0]})
        return pd.concat([_liquidacao(), extra], ignore_index=True)

    def test_desligado_por_padrao(self):
        linha = self._montar(self._liquidacao_nova()).linhas.set_index("ne_curta").loc[NE_NOVA]
        self.assertFalse(linha["limitada_ao_maior_mes"])

    def test_ligado_limita_ao_maior_mes_liquidado(self):
        sem = self._montar(self._liquidacao_nova()).linhas.set_index("ne_curta").loc[NE_NOVA]
        com_rel = self._montar(self._liquidacao_nova(), teto_sem_historico=True)
        com = com_rel.linhas.set_index("ne_curta").loc[NE_NOVA]
        self.assertTrue(com["limitada_ao_maior_mes"])
        self.assertLess(com["projetado_execucao"], sem["projetado_execucao"])
        # set (aberto, sem registro) + out–dez a 600 = 2.400 (ago já liquidado 600 = 100% do teto: completo)
        self.assertAlmostEqual(com["projetado_execucao"], 2400.0)
        self.assertAlmostEqual(com["projetado_valor_cheio"], sem["projetado_valor_cheio"])  # comparação não leva teto
        self.assertTrue(any(NE_NOVA in aviso and "maior mês" in aviso for aviso in com_rel.avisos))

    def test_sem_nenhum_mes_liquidado_nao_ha_teto(self):
        linha = self._montar(_liquidacao(), teto_sem_historico=True).linhas.set_index("ne_curta").loc[NE_NOVA]
        self.assertFalse(linha["limitada_ao_maior_mes"])

    def test_ne_com_historico_nao_tem_teto(self):
        linha = self._montar(_liquidacao(), teto_sem_historico=True).linhas.set_index("ne_curta").loc[NE_EST]
        self.assertFalse(linha["limitada_ao_maior_mes"])


class TestLegendaDaGrade(unittest.TestCase):
    def test_legenda_no_excel_e_no_pdf_diz_que_o_saldo_nao_e_descontado(self):
        por_ne, sem_ne = necessidade_por_ne(_filtrado(), None, 2026)
        relatorio = montar_relatorio(por_ne, sem_ne, _liquidacao(), 2026, 10)
        contexto = ContextoProjecaoExecucao(2026, "01/10/2026", "abcd1234", "14/10/2026 10:00", "arquivo")
        planilha = load_workbook(BytesIO(gerar_xlsx(relatorio, contexto)))
        textos_grade = [c.value for linha in planilha["Grade mensal"].iter_rows() for c in linha if c.value]
        textos_parametros = [c.value for linha in planilha["Parâmetros"].iter_rows() for c in linha if c.value]
        for texto in LEGENDA_GRADE:
            self.assertIn(texto, textos_grade)
            self.assertIn(texto, textos_parametros)
        self.assertTrue(any("SALDO DO EMPENHO NÃO" in texto for texto in LEGENDA_GRADE))
        self.assertTrue(gerar_pdf(relatorio, contexto).startswith(b"%PDF"))


class TestLiquidadoAbatido(unittest.TestCase):
    """14/10/2026: o mês projetado que já descontou liquidado do próprio mês é marcado (*) e detalhado."""

    def _relatorio(self):
        por_ne, sem_ne = necessidade_por_ne(_filtrado(), None, 2026)
        # NE_EST: jan–jul a 900 (fator 0,9) e out parcial de 400 → out projeta 900 − 400 = 500
        extra = pd.DataFrame({"ne_curta": [NE_EST], "ano_mes": [202610], "valor": [400.0]})
        return montar_relatorio(por_ne, sem_ne, pd.concat([_liquidacao(), extra], ignore_index=True), 2026, 10)

    def test_marca_so_o_mes_com_liquidado_no_proprio_mes(self):
        relatorio = self._relatorio()
        indice = relatorio.linhas.index[relatorio.linhas["ne_curta"] == NE_EST][0]
        self.assertTrue(liquidado_abatido(relatorio, indice, 10))  # out: projetado 500 depois de abater 400
        self.assertFalse(liquidado_abatido(relatorio, indice, 11))  # nov: sem liquidação no mês
        self.assertFalse(liquidado_abatido(relatorio, indice, 3))  # mar: realizado, não projetado
        self.assertAlmostEqual(relatorio.mensal.at[indice, "p10"], 500.0)
        self.assertAlmostEqual(relatorio.mensal.at[indice, "e10"], 900.0)

    def test_excel_traz_formato_com_asterisco_e_detalhe_com_esperado(self):
        relatorio = self._relatorio()
        contexto = ContextoProjecaoExecucao(2026, "01/10/2026", "abcd1234", "14/10/2026 10:00", "arquivo")
        planilha = load_workbook(BytesIO(gerar_xlsx(relatorio, contexto)))
        grade = planilha["Grade mensal"]
        linha_est = next(i for i, l in enumerate(grade.iter_rows(values_only=True), start=1) if l[0] == NE_EST)
        coluna_out = [c.value for c in grade[1]].index("Out") + 1
        self.assertIn('"*"', grade.cell(row=linha_est, column=coluna_out).number_format)
        self.assertNotIn('"*"', grade.cell(row=linha_est, column=coluna_out + 1).number_format)
        detalhe = list(planilha["Detalhe mensal"].iter_rows(values_only=True))
        cabecalho = list(detalhe[0])
        self.assertIn("Esperado do mês (custo × fator)", cabecalho)
        out = next(l for l in detalhe[1:] if l[0] == NE_EST and l[cabecalho.index("Mês")] == "Out")
        self.assertAlmostEqual(out[cabecalho.index("Esperado do mês (custo × fator)")], 900.0)
        self.assertEqual(out[cabecalho.index("Liquidado abatido")], "Sim")
        self.assertTrue(gerar_pdf(relatorio, contexto).startswith(b"%PDF"))


class TestAjustesDoMes(unittest.TestCase):
    """14/10/2026: o ajuste do usuário substitui o ESPERADO do mês (custo × fator, com teto); o liquidado do mês
    continua sendo descontado. Mês de referência 10 → fechados até jul, abertos ago e set."""

    def test_ajuste_substitui_esperado_e_teto(self):
        projecao = projetar_meses({}, CUSTO_1000, 0.9, 10, teto_mensal=600.0, ajustes={11: 700.0})
        self.assertEqual(projecao.projetado[10], 700.0)  # nov: ajuste acima do teto
        self.assertEqual(projecao.projetado[11], 600.0)  # dez: sem ajuste, teto vale
        self.assertTrue(projecao.ajustado[10])
        self.assertFalse(projecao.ajustado[11])

    def test_ajuste_abate_o_liquidado_do_mes(self):
        projecao = projetar_meses({10: 400.0}, CUSTO_1000, 0.9, 10, ajustes={10: 900.0})
        self.assertEqual(projecao.projetado[9], 500.0)
        self.assertEqual(projecao.esperado[9], 900.0)
        self.assertEqual(projecao.tipos[9], TIPO_PROJETADO)

    def test_ajuste_zero_e_valido(self):
        projecao = projetar_meses({}, CUSTO_1000, 0.9, 10, ajustes={11: 0.0})
        self.assertEqual(projecao.projetado[10], 0.0)
        self.assertTrue(projecao.ajustado[10])
        self.assertEqual(projecao.tipos[10], TIPO_PROJETADO)

    def test_mes_vazio_da_metodologia_e_editavel(self):
        projecao = projetar_meses({}, [0.0] * 12, 1.0, 10, ajustes={11: 300.0})
        self.assertEqual(projecao.projetado[10], 300.0)
        self.assertIsNone(projecao.metodologia[10])
        self.assertIsNone(projecao.projetado_valor_cheio[10])  # a comparação "valor cheio" não leva ajuste

    def test_mes_fechado_e_superado(self):
        projecao = projetar_meses({}, CUSTO_1000, 0.9, 10, ajustes={5: 123.0})
        self.assertEqual(projecao.superados, [5])
        self.assertIsNone(projecao.projetado[4])
        self.assertFalse(projecao.ajustado[4])

    def test_exercicio_encerrado_ignora_ajustes(self):
        projecao = projetar_meses({}, CUSTO_1000, 0.9, 12, ajustes={12: 500.0})
        self.assertEqual(projecao.total_projetado, 0.0)
        self.assertEqual(projecao.superados, [12])

    def test_metodologia_e_guardada_mesmo_com_ajuste(self):
        sem = projetar_meses({}, CUSTO_1000, 0.9, 10)
        com = projetar_meses({}, CUSTO_1000, 0.9, 10, ajustes={11: 100.0})
        self.assertAlmostEqual(sem.metodologia[10], 900.0)
        self.assertAlmostEqual(com.metodologia[10], 900.0)
        self.assertIsNone(com.metodologia[4])  # mês fechado não é editável

    def test_mes_aberto_com_ajuste_segue_a_regra_dos_50_por_cento(self):
        completo = projetar_meses({8: 500.0}, CUSTO_1000, 0.9, 10, ajustes={8: 900.0})  # 500 ≥ 450
        self.assertEqual(completo.tipos[7], "Realizado")
        self.assertIsNone(completo.projetado[7])
        self.assertTrue(completo.ajustado[7])
        aberto = projetar_meses({8: 500.0}, CUSTO_1000, 0.9, 10, ajustes={8: 2000.0})  # 500 < 1000
        self.assertEqual(aberto.tipos[7], TIPO_RESTANTE_ABERTO)
        self.assertEqual(aberto.projetado[7], 1500.0)

    def test_valor_cheio_nao_leva_ajuste(self):
        projecao = projetar_meses({}, CUSTO_1000, 0.9, 10, ajustes={11: 5000.0})
        self.assertEqual(projecao.projetado_valor_cheio[10], 1000.0)

    def test_sem_ajustes_nada_muda(self):
        self.assertEqual(projetar_meses({}, CUSTO_1000, 0.9, 10).projetado, projetar_meses({}, CUSTO_1000, 0.9, 10, ajustes={}).projetado)


class TestRelatorioComAjustes(unittest.TestCase):
    def _montar(self, ajustes):
        por_ne, sem_ne = necessidade_por_ne(_filtrado(), None, 2026)
        return montar_relatorio(por_ne, sem_ne, _liquidacao(), 2026, 10, ajustes=ajustes)

    def test_ajuste_muda_projetado_e_necessidade_so_da_ne_ajustada(self):
        # NE_EST sem ajuste: projetado 4.500, necessidade 1.800. Com dez = 5.000: 4.500 − 900 + 5.000 = 8.600 → 5.900.
        linhas = self._montar({NE_EST: {12: 5000.0}}).linhas.set_index("ne_curta")
        self.assertAlmostEqual(linhas.loc[NE_EST, "projetado_execucao"], 8600.0)
        self.assertAlmostEqual(linhas.loc[NE_EST, "necessidade_execucao"], 5900.0)
        self.assertTrue(linhas.loc[NE_EST, "ajustada_manualmente"])
        self.assertFalse(linhas.loc[NE_NOVA, "ajustada_manualmente"])

    def test_sem_ajustes_igual_ao_de_sempre(self):
        base = self._montar(None).linhas
        com_vazio = self._montar({}).linhas
        pd.testing.assert_frame_equal(base, com_vazio)
        self.assertFalse(base["ajustada_manualmente"].any())

    def test_ajuste_de_ne_inexistente_e_ignorado_e_avisado(self):
        relatorio = self._montar({"2026NE999999": {12: 10.0}})
        self.assertEqual(relatorio.ajustes_sem_ne, ["2026NE999999"])
        self.assertTrue(any("2026NE999999" in aviso for aviso in relatorio.avisos))
        self.assertFalse(relatorio.linhas["ajustada_manualmente"].any())

    def test_ajuste_superado_e_avisado_com_titulo_do_contrato(self):
        relatorio = self._montar({NE_EST: {5: 10.0}})
        self.assertEqual(relatorio.superados, {NE_EST: [5]})
        aviso = next(a for a in relatorio.avisos if "superado" in a.lower())
        self.assertIn(NE_EST, aviso)
        self.assertIn("10/2025", aviso)
        self.assertIn("ESTAVEL", aviso)

    def test_aviso_de_ajustadas_traz_numero_do_contrato_e_fornecedor(self):
        relatorio = self._montar({NE_EST: {12: 5000.0}})
        aviso = next(a for a in relatorio.avisos if "manualmente" in a)
        self.assertIn("10/2025", aviso)
        self.assertIn("ESTAVEL", aviso)
        self.assertIn("Dez", aviso)

    def test_mensal_traz_marca_e_metodologia(self):
        relatorio = self._montar({NE_EST: {12: 5000.0}})
        indice = relatorio.linhas.index[relatorio.linhas["ne_curta"] == NE_EST][0]
        self.assertTrue(relatorio.mensal.at[indice, "a12"])
        self.assertFalse(relatorio.mensal.at[indice, "a11"])
        self.assertAlmostEqual(relatorio.mensal.at[indice, "m12"], 900.0)
        self.assertAlmostEqual(relatorio.mensal.at[indice, "e12"], 5000.0)

    def test_excel_traz_coluna_ajustado_e_formato_com_marca(self):
        relatorio = self._montar({NE_EST: {12: 5000.0}})
        contexto = ContextoProjecaoExecucao(2026, "01/10/2026", "abcd1234", "14/10/2026 10:00", "arquivo")
        planilha = load_workbook(BytesIO(gerar_xlsx(relatorio, contexto)))
        resumo = list(planilha["Resumo por NE"].iter_rows(values_only=True))
        cab = list(resumo[0])
        self.assertIn("Ajustado manualmente", cab)
        linha = next(l for l in resumo[1:] if l[0] == NE_EST)
        self.assertEqual(linha[cab.index("Ajustado manualmente")], "Sim")
        grade = planilha["Grade mensal"]
        linha_est = next(i for i, l in enumerate(grade.iter_rows(values_only=True), start=1) if l[0] == NE_EST)
        coluna_dez = [c.value for c in grade[1]].index("Dez") + 1
        self.assertIn("†", grade.cell(row=linha_est, column=coluna_dez).number_format)
        self.assertNotIn("†", grade.cell(row=linha_est, column=coluna_dez - 1).number_format)
        self.assertTrue(gerar_pdf(relatorio, contexto).startswith(b"%PDF"))

    def test_legenda_cita_o_ajuste(self):
        self.assertTrue(any("†" in texto for texto in LEGENDA_GRADE))
