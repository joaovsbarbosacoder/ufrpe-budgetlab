"""Testes da reimportação pela interface — página "Atualizar Planilhas", card "Dotação
Orçamentária (Dotação Anual)" (movida de `app_pages/dotacao_orcamentaria.py`, ver
`src/reimportacao_especificacoes.py`).

Espelha `tests/test_execucao_orcamentaria_reimportacao.py` (mesmo componente reutilizado,
`src/ui_reimportacao.py`), com fixtures sintéticas de duas colunas-ano em vez de variantes
do arquivo real — a Dotação Anual guarda um indicador por coluna, uma coluna por ano, então
duas colunas do mesmo indicador com anos diferentes já reproduzem exercício retroativo,
removido e em avanço sem precisar editar um arquivo de milhares de linhas.

"Atualizar Planilhas" tem mais de um `st.file_uploader` na mesma página — os testes pegam o
SEGUNDO uploader com rótulo "Nova extração (.xlsx)" (o primeiro é da Execução Anual, card
renderizado antes do de Dotação — ver `app_pages/atualizar_planilhas.py`).
"""

from __future__ import annotations

import io
import shutil
import unittest
from pathlib import Path

from openpyxl import Workbook
from streamlit.testing.v1 import AppTest

from src.importacao_dotacao import Manifesto
from tests.test_tesouro_dotacao_anual import merge_cells

PROJECT_ROOT = Path(__file__).resolve().parents[1]
DIRETORIO_RAW = PROJECT_ROOT / "data/raw"
DIRETORIO_MANIFESTOS = PROJECT_ROOT / "data/manifestos"
MANIFESTO_ATUAL = DIRETORIO_MANIFESTOS / "dotacao_anual_atual.json"

METRIC = "Movim. Líquido - R$ (Item Informação)"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _header(ws) -> None:
    ws["A1"] = "Iduso"
    merge_cells(ws, 1, 1, 3, 2)
    ws["C1"] = "Resultado Primário Lei"
    merge_cells(ws, 1, 3, 3, 4)
    ws["E1"] = "Ação Governo"
    merge_cells(ws, 1, 5, 3, 6)
    ws["G1"] = "PTRES"
    merge_cells(ws, 1, 7, 3, 7)
    ws["H1"] = "Fonte Recursos Detalhada"
    merge_cells(ws, 1, 8, 3, 9)
    ws["J1"] = "Item Informação"
    merge_cells(ws, 1, 10, 1, 11)
    ws["J2"] = "Ano Lançamento"
    merge_cells(ws, 2, 10, 2, 11)
    ws["J3"] = "Plano Orçamentário"
    merge_cells(ws, 3, 10, 3, 11)

    ws["A4"] = 1
    ws["B4"] = "Iduso A"
    ws["C4"] = 10
    ws["D4"] = "RP A"
    ws["E4"] = "ACAO1"
    ws["F4"] = "Acao desc A"
    ws["G4"] = "PTRES1"
    ws["H4"] = "FONTE1"
    ws["I4"] = "Fonte A"
    ws["J4"] = "PO1"
    ws["K4"] = "PO A"


def _two_years(valor_2023: float, valor_2024: float):
    def builder(ws) -> None:
        _header(ws)
        ws["L1"] = "DOTACAO ATUALIZADA"
        ws["L2"] = 2023
        ws["L3"] = METRIC
        ws["M1"] = "DOTACAO ATUALIZADA"
        ws["M2"] = 2024
        ws["M3"] = METRIC
        ws["L4"] = valor_2023
        ws["M4"] = valor_2024

    return builder


def _only_2024(valor_2024: float):
    def builder(ws) -> None:
        _header(ws)
        ws["L1"] = "DOTACAO ATUALIZADA"
        ws["L2"] = 2024
        ws["L3"] = METRIC
        ws["L4"] = valor_2024

    return builder


def _invalid(ws) -> None:
    _two_years(1000, 2000)(ws)
    ws["L4"] = "não é número"


def _workbook_bytes(builder) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Base"
    builder(worksheet)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@unittest.skipUnless(MANIFESTO_ATUAL.exists(), f"Manifesto ausente em {MANIFESTO_ATUAL}")
class DotacaoReimportacaoPageTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.baseline = _workbook_bytes(_two_years(1000, 2000))
        cls.retroativo = _workbook_bytes(_two_years(1500, 2000))
        cls.removido = _workbook_bytes(_only_2024(2000))
        cls.avanco = _workbook_bytes(_two_years(1000, 2500))
        cls.invalido = _workbook_bytes(_invalid)

    def setUp(self) -> None:
        self._manifesto_backup = MANIFESTO_ATUAL.read_bytes()
        self._raw_antes = set(DIRETORIO_RAW.glob("*.xlsx"))
        self._manifestos_antes = set(DIRETORIO_MANIFESTOS.glob("dotacao_anual_*.json"))

    def tearDown(self) -> None:
        MANIFESTO_ATUAL.write_bytes(self._manifesto_backup)
        for novo in set(DIRETORIO_RAW.glob("*.xlsx")) - self._raw_antes:
            novo.unlink(missing_ok=True)
        for novo in set(DIRETORIO_MANIFESTOS.glob("dotacao_anual_*.json")) - self._manifestos_antes:
            novo.unlink(missing_ok=True)
        staging = DIRETORIO_RAW / "_staging"
        if staging.exists():
            shutil.rmtree(staging)

    def _open_page(self) -> AppTest:
        app = AppTest.from_file(str(PROJECT_ROOT / "app.py"))
        app.run()
        app.switch_page("app_pages/atualizar_planilhas.py")
        app.run(timeout=60)
        return app

    def _upload(self, app: AppTest, content: bytes, filename: str) -> None:
        uploaders = [u for u in app.file_uploader if u.label == "Nova extração (.xlsx)"]
        uploaders[1].set_value((filename, content, XLSX_MIME))
        app.run(timeout=60)

    def _importar_baseline(self) -> None:
        from src.importacao_dotacao import importar

        caminho = DIRETORIO_RAW / "dotacao_baseline.xlsx"
        caminho.write_bytes(self.baseline)
        resultado = importar(caminho)
        assert resultado.ok, resultado.validacao.erros

    def test_retroactive_change_shows_gate_and_blocks_commit_by_default(self) -> None:
        self._importar_baseline()
        app = self._open_page()
        self._upload(app, self.retroativo, "dotacao_retroativo.xlsx")

        self.assertEqual(len(app.exception), 0)
        textos = [item.value for item in app.markdown]
        self.assertTrue(any("Exercícios com valores alterados retroativamente" in t for t in textos))
        self.assertTrue(any("2023" in t and "→" in t for t in textos))

        confirmar = next(b for b in app.button if b.label == "Confirmar substituição")
        self.assertTrue(confirmar.disabled)
        self.assertFalse(any(b.label == "Confirmar importação" for b in app.button))

        checkbox = next(c for c in app.checkbox if "confirmo a atualização" in c.label)
        self.assertFalse(checkbox.value)

    def test_retroactive_change_commits_once_explicitly_confirmed(self) -> None:
        self._importar_baseline()
        app = self._open_page()
        self._upload(app, self.retroativo, "dotacao_retroativo.xlsx")

        checkbox = next(c for c in app.checkbox if "confirmo a atualização" in c.label)
        checkbox.check()
        app.run(timeout=60)

        confirmar = next(b for b in app.button if b.label == "Confirmar substituição")
        self.assertFalse(confirmar.disabled)
        confirmar.click()
        app.run(timeout=60)

        self.assertEqual(len(app.exception), 0)
        manifesto_novo = Manifesto.atual()
        self.assertEqual(manifesto_novo.totais_por_ano["2023"]["dotacao_atualizada"], 1500.0)
        self.assertEqual(manifesto_novo.arquivo, "dotacao_retroativo.xlsx")

    def test_removed_exercise_is_informational_and_does_not_block(self) -> None:
        # mesmo critério de test_execucao_orcamentaria_reimportacao.py: composição por ano não
        # bloqueia mais por exercício ausente, só informa.
        self._importar_baseline()
        app = self._open_page()
        self._upload(app, self.removido, "dotacao_removido.xlsx")

        self.assertEqual(len(app.exception), 0)
        textos = [item.value for item in app.markdown]
        self.assertTrue(any("não incluídos nesta extração" in t and "2023" in t for t in textos))
        self.assertFalse(any(b.label == "Confirmar substituição" for b in app.button))

        confirmar = next(b for b in app.button if b.label == "Confirmar importação")
        self.assertFalse(confirmar.disabled)
        confirmar.click()
        app.run(timeout=60)

        self.assertEqual(len(app.exception), 0)
        manifesto_novo = Manifesto.atual()
        self.assertNotIn(2023, manifesto_novo.anos)

    def test_ongoing_exercise_advance_commits_without_gate(self) -> None:
        self._importar_baseline()
        app = self._open_page()
        self._upload(app, self.avanco, "dotacao_avanco.xlsx")

        self.assertEqual(len(app.exception), 0)
        self.assertFalse(any(b.label == "Confirmar substituição" for b in app.button))
        confirmar = next(b for b in app.button if b.label == "Confirmar importação")
        self.assertFalse(confirmar.disabled)

        confirmar.click()
        app.run(timeout=60)

        self.assertEqual(len(app.exception), 0)
        manifesto_novo = Manifesto.atual()
        self.assertEqual(manifesto_novo.arquivo, "dotacao_avanco.xlsx")
        self.assertEqual(manifesto_novo.totais_por_ano["2024"]["dotacao_atualizada"], 2500.0)

    def test_uploading_identical_file_reports_nothing_to_save(self) -> None:
        self._importar_baseline()
        conteudo_identico = (DIRETORIO_RAW / "dotacao_baseline.xlsx").read_bytes()

        app = self._open_page()
        self._upload(app, conteudo_identico, "dotacao_baseline.xlsx")

        self.assertEqual(len(app.exception), 0)
        self.assertTrue(
            any("Arquivo idêntico à extração atual" in item.value for item in app.info)
        )
        self.assertFalse(
            any(b.label in ("Confirmar importação", "Confirmar substituição") for b in app.button)
        )

    def test_invalid_extraction_is_blocked_and_reports_errors(self) -> None:
        app = self._open_page()
        self._upload(app, self.invalido, "dotacao_invalido.xlsx")

        self.assertEqual(len(app.exception), 0)
        textos_erro = [item.value for item in app.error]
        self.assertTrue(any("bloqueada" in t for t in textos_erro))
        self.assertFalse(
            any(b.label in ("Confirmar importação", "Confirmar substituição") for b in app.button)
        )
        self.assertEqual(MANIFESTO_ATUAL.read_bytes(), self._manifesto_backup)


@unittest.skipUnless(MANIFESTO_ATUAL.exists(), f"Manifesto ausente em {MANIFESTO_ATUAL}")
class PastaDeEntradaDotacaoTests(unittest.TestCase):
    """Confirma que a Dotação Anual também reconhece a pasta de entrada compartilhada
    (`data/raw/_entrada/`, mesma pasta da Execução Anual — ver
    `tests/test_execucao_orcamentaria_reimportacao.py::PastaDeEntradaTests`, que já cobre o
    mecanismo em detalhe). Só um caso aqui: o essencial é confirmar a fiação, não repetir
    a bateria inteira de gate/ambiguidade já testada do outro lado."""

    DIRETORIO_ENTRADA = DIRETORIO_RAW / "_entrada"

    @classmethod
    def setUpClass(cls) -> None:
        cls.baseline = _workbook_bytes(_two_years(1000, 2000))

    def _importar_baseline(self) -> None:
        from src.importacao_dotacao import importar

        caminho = DIRETORIO_RAW / "dotacao_baseline.xlsx"
        caminho.write_bytes(self.baseline)
        resultado = importar(caminho)
        assert resultado.ok, resultado.validacao.erros

    def setUp(self) -> None:
        self._manifesto_backup = MANIFESTO_ATUAL.read_bytes()
        self._raw_antes = set(DIRETORIO_RAW.glob("*.xlsx"))
        self._manifestos_antes = set(DIRETORIO_MANIFESTOS.glob("dotacao_anual_*.json"))
        if self.DIRETORIO_ENTRADA.exists():
            shutil.rmtree(self.DIRETORIO_ENTRADA)

    def tearDown(self) -> None:
        MANIFESTO_ATUAL.write_bytes(self._manifesto_backup)
        for novo in set(DIRETORIO_RAW.glob("*.xlsx")) - self._raw_antes:
            novo.unlink(missing_ok=True)
        for novo in set(DIRETORIO_MANIFESTOS.glob("dotacao_anual_*.json")) - self._manifestos_antes:
            novo.unlink(missing_ok=True)
        for pasta in (DIRETORIO_RAW / "_staging", self.DIRETORIO_ENTRADA):
            if pasta.exists():
                shutil.rmtree(pasta)

    def _open_page(self) -> AppTest:
        app = AppTest.from_file(str(PROJECT_ROOT / "app.py"))
        app.run()
        app.switch_page("app_pages/atualizar_planilhas.py")
        app.run(timeout=60)
        return app

    def test_arquivo_seguro_na_pasta_compartilhada_e_aplicado_automaticamente(self) -> None:
        self._importar_baseline()
        self.DIRETORIO_ENTRADA.mkdir(parents=True, exist_ok=True)
        (self.DIRETORIO_ENTRADA / "dotacao_avanco.xlsx").write_bytes(_workbook_bytes(_two_years(1000, 2500)))

        app = self._open_page()

        self.assertEqual(len(app.exception), 0)
        self.assertFalse(any(b.label == "Confirmar importação" for b in app.button))
        self.assertFalse(any(b.label == "Confirmar substituição" for b in app.button))
        self.assertNotEqual(MANIFESTO_ATUAL.read_bytes(), self._manifesto_backup)
        manifesto_novo = Manifesto.atual()
        self.assertEqual(manifesto_novo.arquivo, "dotacao_avanco.xlsx")
        self.assertEqual(manifesto_novo.totais_por_ano["2024"]["dotacao_atualizada"], 2500.0)
        self.assertFalse((self.DIRETORIO_ENTRADA / "dotacao_avanco.xlsx").exists())


if __name__ == "__main__":
    unittest.main()
