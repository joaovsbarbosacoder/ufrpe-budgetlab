"""Testes de `src.ddp_siape_relatorio` (relatório em Excel da aba Folha (DDP), 27/10/2026).

Monta o relatório com a base sintética de `tests.test_ddp_siape` (folha + RPPS) e confere a
organização: as três visões (somente reforço, sugestões de anulação, anulação + reforço),
cabeçalhos em português, notas de empenho completas, totais que fecham com as tabelas da aba e
o resumo com os dois cenários.

Valores à mão (base sintética): faltas do mês 200 (NE1, 31901101) + 200 (NE2) + 20 (NE4+5) +
19 (NE6, 31900306) + 7 (sem empenho) = 446. Com anulação, a NE1 cobre os seus 200 com o 13º ⇒
reforço do mês 246. Até dezembro, sem anular: 1.200 + 800 + 20 + 100 + 19 + 7 = 2.146 (o 13º da
NE1, 2.000, e a sobra da NE3, 200, ficam parados); com anulação: 946."""
from __future__ import annotations

import unittest
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from src.ddp_siape import (
    conciliacao_por_acao, conferencia, confronto_liquidado, ddp_por_natureza, plano_operacoes, por_empenho_somente_reforco,
    quadro_reforco, quadro_sem_ddp, reforco_por_acao_po, variacao_por_natureza,
)
from src.ddp_siape_relatorio import DadosRelatorio, LINHA_CABECALHO, relatorio_excel
from src.despesas_pessoal import GRUPO_RPPS
from tests.test_ddp_siape import RESERVA, _base_sintetica, _ddp_sintetica, _item, _relatorios_set

ABAS = ["Resumo", "Somente reforço", "Sugestões de anulação", "Anulação + reforço", "Por empenho (com anulação)",
        "Dotação por ação e PO", "Conciliação com o painel", "Detalhe por item", "Conferência da DDP"]


def _dados(com_variacao=False):
    base = pd.concat([_base_sintetica(), pd.DataFrame(
        _item(10, "31911303", 1000.0, {202608: 300.0, 202609: 300.0}, acao="09HB", favorecido="FUNDO RPPS"))],
        ignore_index=True)
    grade = pd.DataFrame([{"grupo": GRUPO_RPPS, "natureza_detalhada_cod": "31911303", "natureza_detalhada_desc": "RPPS",
                           "meses": [None] * 7 + [300.0, 300.0, 350.0, 700.0, 350.0]}])
    natureza = _ddp_sintetica()
    itens, por_empenho = quadro_reforco(natureza, base, 202610, RESERVA)
    itens_rpps, por_empenho_rpps = quadro_sem_ddp(base, 202610, grade, GRUPO_RPPS)
    dotacao = pd.Series({("20TP", "X"): 7000.0, ("2004", "X"): 10500.0, ("09HB", "X"): 1900.0})
    relatorios = _relatorios_set()
    return DadosRelatorio(
        competencia=202610, procedencia="Procedência de teste.", conferencia=conferencia(relatorios), natureza=natureza,
        por_empenho=por_empenho, itens=itens, plano=plano_operacoes(itens), confronto=confronto_liquidado(natureza, base, 202610),
        rpps_itens=itens_rpps, rpps_por_empenho=por_empenho_rpps,
        acao_po=reforco_por_acao_po([(itens, por_empenho), (itens_rpps, por_empenho_rpps)], base, 2026, dotacao),
        acao_po_somente_reforco=reforco_por_acao_po(
            [(itens, por_empenho_somente_reforco(itens)), (itens_rpps, por_empenho_somente_reforco(itens_rpps))], base, 2026, dotacao),
        conciliacao=conciliacao_por_acao([(itens, por_empenho), (itens_rpps, por_empenho_rpps)], base, 202610, grade, dotacao),
        ano_dotacao=2026,
        variacao=variacao_por_natureza(natureza, ddp_por_natureza(relatorios)) if com_variacao else None,
        competencia_anterior=202609 if com_variacao else None,
    )


def _linhas(aba, inicio=LINHA_CABECALHO + 1):
    """Linhas de dados de uma tabela, até a linha "Total" (exclusive)."""
    linhas = []
    for r in aba.iter_rows(min_row=inicio, values_only=True):
        if r[0] == "Total" or all(v is None for v in r):
            break
        linhas.append(r)
    return linhas


def _cabecalho(aba, linha=LINHA_CABECALHO):
    return [c.value for c in aba[linha]]


class TestRelatorioExcel(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.dados = _dados(com_variacao=True)
        cls.livro = load_workbook(BytesIO(relatorio_excel(cls.dados)))

    def test_abas_na_ordem_de_leitura(self):
        self.assertEqual(self.livro.sheetnames, ABAS)

    def test_cabecalhos_em_portugues_sem_nomes_tecnicos(self):
        for nome in ABAS:
            cabecalho = [v for v in _cabecalho(self.livro[nome]) if v]
            self.assertTrue(cabecalho, nome)
            for titulo in cabecalho:
                with self.subTest(aba=nome, titulo=titulo):
                    self.assertNotIn("_", str(titulo))

    def test_resumo_compara_os_dois_cenarios(self):
        aba = self.livro["Resumo"]
        self.assertEqual(_cabecalho(aba)[:3], ["Indicador", "Somente reforço (sem anular agora)", "Anulação + reforço"])
        cenario = {r[0]: (r[1], r[2]) for r in _linhas(aba)}
        self.assertEqual(cenario["Reforço necessário para out/2026 — empenhos da folha"], (446.0, 246.0))
        self.assertEqual(cenario["Reforço necessário até dez/2026 — empenhos da folha"], (2146.0, 946.0))
        self.assertEqual(cenario["Reforço necessário até dez/2026 — RPPS patronal"], (1000.0, 1000.0))
        self.assertEqual(cenario["Anulações (remanejamento entre itens do mesmo empenho)"], (0.0, 200.0))
        self.assertEqual(cenario["Sobra que ficaria parada nos empenhos em dezembro"], (2200.0, 1000.0))
        textos = [r[0] for r in aba.iter_rows(values_only=True) if r[0]]
        self.assertIn("Onde o reforço não cabe na dotação — somente reforço", textos)
        self.assertIn("Onde o reforço não cabe na dotação — anulação + reforço", textos)
        # 2004/X: dotação não empenhada 500; reforço até dez 800 nos dois cenários ⇒ falta 300
        self.assertGreaterEqual([r for r in aba.iter_rows(values_only=True) if r[0] == "2004 / X"].count(("2004 / X", 300.0, None)), 1)

    def test_somente_reforco_reforca_a_falta_inteira_de_cada_item(self):
        aba = self.livro["Somente reforço"]
        cab = _cabecalho(aba)
        linhas = [dict(zip(cab, r)) for r in _linhas(aba)]
        ne1 = next(l for l in linhas if l["Nota de empenho"] == "2026NE000001")
        self.assertEqual((ne1["Natureza"], ne1["Reforço necessário para out/2026"], ne1["Reforço necessário até dez/2026"]),
                         ("31901101", 200.0, 1200.0))
        folha = [l for l in linhas if l["Origem"] == "Folha (DDP)"]
        self.assertEqual(round(sum(l["Reforço necessário para out/2026"] for l in folha), 2), 446.0)
        self.assertTrue(any(l["Origem"] == "RPPS patronal (projeção)" for l in linhas))
        textos = [r[0] for r in aba.iter_rows(values_only=True) if isinstance(r[0], str)]
        self.assertIn("Por nota de empenho", textos)
        self.assertTrue(any(t.startswith("O reforço cabe na dotação?") for t in textos))

    def test_sugestoes_de_anulacao(self):
        aba = self.livro["Sugestões de anulação"]
        linhas = [dict(zip(_cabecalho(aba), r)) for r in _linhas(aba)]
        decimo = next(l for l in linhas if l["Natureza"] == "31901143")
        self.assertEqual((decimo["Nota de empenho"], decimo["Pode ser anulado"], decimo["Usado no plano Anulação + reforço"],
                          decimo["Restante após o plano"]), ("2026NE000001", 2000.0, 200.0, 1800.0))
        per_capita = next(l for l in linhas if l["Nota de empenho"] == "2026NE000003")
        self.assertEqual((per_capita["Pode ser anulado"], per_capita["Usado no plano Anulação + reforço"]), (200.0, 0.0))

    def test_anulacao_mais_reforco_com_operacoes_legiveis(self):
        aba = self.livro["Anulação + reforço"]
        linhas = [dict(zip(_cabecalho(aba), r)) for r in _linhas(aba)]
        remanejo = next(l for l in linhas if l["Nota de empenho"] == "2026NE000001")
        self.assertEqual(remanejo["Operação"], "Remanejar entre itens do empenho")
        self.assertEqual((remanejo["Item que cede (anular)"], remanejo["Item que recebe (reforçar)"], remanejo["Valor (R$)"]),
                         ("31901143", "31901101", 200.0))
        self.assertEqual(remanejo["Descrição do item que recebe"], "DESC 31901101")
        self.assertIn("Reforçar o empenho (valor novo)", {l["Operação"] for l in linhas})
        self.assertEqual(round(sum(l["Valor (R$)"] for l in linhas), 2), 446.0)

    def test_por_empenho_com_nota_completa_e_total_por_formula(self):
        aba = self.livro["Por empenho (com anulação)"]
        cab = _cabecalho(aba)
        linhas = [dict(zip(cab, r)) for r in _linhas(aba)]
        notas = {l["Nota de empenho"] for l in linhas}
        self.assertTrue({"2026NE000001", "2026NE000004, 2026NE000005", "(sem empenho)"} <= notas)
        self.assertEqual(round(sum(l["Reforço necessário para out/2026"] for l in linhas), 2), 246.0)
        total = LINHA_CABECALHO + 1 + len(linhas)
        coluna = cab.index("Reforço necessário para out/2026") + 1
        self.assertEqual(aba.cell(row=total, column=1).value, "Total")
        self.assertTrue(str(aba.cell(row=total, column=coluna).value).startswith("=SUM("))

    def test_dotacao_com_situacao_e_conciliacao(self):
        aba = self.livro["Dotação por ação e PO"]
        linhas = {r[0]: dict(zip(_cabecalho(aba), r)) for r in _linhas(aba)}
        self.assertEqual(linhas["2004"]["Situação"], "Não cabe")
        self.assertEqual(linhas["20TP"]["Situação"], "Cabe")
        self.assertEqual(linhas["0181"]["Situação"], "Sem dotação na base")
        conc = self.livro["Conciliação com o painel"]
        self.assertIn("Diferença (deve ser zero)", _cabecalho(conc))

    def test_conferencia_com_variacao_e_confronto(self):
        textos = [c.value for c in self.livro["Conferência da DDP"]["A"] if isinstance(c.value, str)]
        self.assertIn("O que mudou de set/2026 para out/2026", textos)
        self.assertIn("DDP × liquidado na base em out/2026", textos)

    def test_sem_dotacao_nem_variacao_omite_as_abas_correspondentes(self):
        dados = _dados()
        dados.acao_po = dados.acao_po_somente_reforco = dados.conciliacao = None
        livro = load_workbook(BytesIO(relatorio_excel(dados)))
        self.assertEqual(livro.sheetnames, [a for a in ABAS if a not in ("Dotação por ação e PO", "Conciliação com o painel")])
        textos = [c.value for c in livro["Conferência da DDP"]["A"] if isinstance(c.value, str)]
        self.assertFalse(any(t.startswith("O que mudou") for t in textos))


if __name__ == "__main__":
    unittest.main()
