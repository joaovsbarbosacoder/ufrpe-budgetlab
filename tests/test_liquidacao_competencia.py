"""
Testes da base de Liquidação por Competência.

Usa uma fixture congelada em `tests/fixtures/liquidacao_competencia_2026-09-11.xlsx` (não um
arquivo de `data/raw/` — esta base não tem importação versionada, mesmo padrão de
`tests/test_tesouro_execucao_mensal.py`) — pulado se o arquivo não existir.

Dois níveis, mesmo espírito de `test_tesouro_execucao_mensal.py`:
  * INVARIANTES — valem para qualquer extração desta base.
  * REFERÊNCIA — totais absolutos da extração identificada por hash (`REFERENCIAS`). Ao
    trocar a fixture, NÃO edite os números existentes: acrescente uma entrada nova.

IMPORTANTE: ao atualizar a planilha de trabalho, NÃO sobrescreva esta fixture automaticamente
— mesma regra das demais fixtures deste diretório (ver AGENTS.md).
"""

from __future__ import annotations

import hashlib
import tempfile
import unittest
from pathlib import Path

import openpyxl

from src.liquidacao_competencia import (
    ErroLayoutBase,
    _rotulo_mes,
    dimensoes_por_ne,
    ler_liquidacao_competencia,
    liquidado_por_ne_e_mes,
    reconciliar,
    validar,
)

CAMINHO_BASE = Path("tests/fixtures/liquidacao_competencia_2026-09-11.xlsx")

# Extração repassada pelo usuário via Downloads em 11/09/2026 (hash 255a8cd9...). Ao trocar a
# fixture, NÃO edite estes números: acrescente uma entrada nova (o hash decide qual
# referência se aplica).
REFERENCIAS = {
    "255a8cd9": {
        "descricao": "Extração recebida em 11/09/2026 ('Liquidação por competência - JVSB (2).xlsx')",
        "linhas": 3696,
        "documentos_habeis_distintos": 3219,
        "notas_empenho_distintas": 436,
        "meses_referencia": [
            202510, 202511, 202512, 202601, 202602, 202603,
            202604, 202605, 202606, 202607, 202608, 202609,
        ],
        "total": 580183775.26,
        "total_estornos": -1733191.18,
    },
}


def _hash_curto(caminho: Path) -> str:
    return hashlib.sha256(caminho.read_bytes()).hexdigest()[:8]


class TestRotuloMes(unittest.TestCase):
    def test_rotulo_valido(self):
        self.assertEqual(_rotulo_mes("MAI/2026"), (2026, 5))
        self.assertEqual(_rotulo_mes("dez/2025"), (2025, 12))

    def test_rotulo_invalido_e_erro(self):
        with self.assertRaises(ErroLayoutBase):
            _rotulo_mes("XXX/2026")
        with self.assertRaises(ErroLayoutBase):
            _rotulo_mes(None)


@unittest.skipUnless(CAMINHO_BASE.exists(), f"Planilha ausente em {CAMINHO_BASE}")
class TestLeituraEValidacaoAssinatura(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.df = ler_liquidacao_competencia(CAMINHO_BASE)

    def test_colunas_essenciais_presentes(self):
        for coluna in (
            "ne_ccor", "documento_habil", "doc_contabil", "ano_referencia",
            "mes_referencia", "ano_mes_referencia", "valor", "linha_origem", "arquivo_origem",
        ):
            self.assertIn(coluna, self.df.columns)

    def test_chave_ne_dh_doccontabil_mes_e_unica(self):
        chave = ["ne_ccor", "documento_habil", "doc_contabil", "ano_mes_referencia"]
        self.assertEqual(self.df.duplicated(subset=chave).sum(), 0)

    def test_layout_diferente_e_rejeitado(self):
        with tempfile.TemporaryDirectory() as tmp:
            wb = openpyxl.load_workbook(CAMINHO_BASE)
            aba = wb.active
            aba.cell(row=1, column=1).value = "Outra Coisa"  # não mais "NE"
            destino = Path(tmp) / CAMINHO_BASE.name
            wb.save(destino)
            with self.assertRaises(ErroLayoutBase):
                ler_liquidacao_competencia(destino)


@unittest.skipUnless(CAMINHO_BASE.exists(), f"Planilha ausente em {CAMINHO_BASE}")
class TestEstornosMantemSinal(unittest.TestCase):
    """Regra central da base (ver docstring de src/liquidacao_competencia.py): estornos
    entram com sinal negativo, nunca abs()."""

    @classmethod
    def setUpClass(cls):
        cls.df = ler_liquidacao_competencia(CAMINHO_BASE)

    def test_existem_linhas_negativas(self):
        self.assertGreater((self.df["valor"] < 0).sum(), 0)

    def test_soma_por_ne_e_mes_preserva_estornos(self):
        agregado = liquidado_por_ne_e_mes(self.df)
        self.assertAlmostEqual(float(agregado["valor"].sum()), float(self.df["valor"].sum()), places=2)
        # o caso citado na especificação: mesma NE, mesmo valor absoluto em meses diferentes,
        # sinais opostos — soma zero, sem cancelar as duas linhas antes de somar.
        ne_tf = "153165152392026NE000365"
        linhas = agregado[agregado["ne_ccor"] == ne_tf]
        self.assertAlmostEqual(float(linhas["valor"].sum()), 0.0, places=2)


@unittest.skipUnless(CAMINHO_BASE.exists(), f"Planilha ausente em {CAMINHO_BASE}")
class TestLiquidadoPorNeEMes(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.df = ler_liquidacao_competencia(CAMINHO_BASE)
        cls.agregado = liquidado_por_ne_e_mes(cls.df)

    def test_uma_linha_por_ne_e_mes_com_dado(self):
        chave = ["ne_ccor", "ano_mes"]
        self.assertEqual(self.agregado.duplicated(subset=chave).sum(), 0)

    def test_ne_sem_nenhuma_linha_de_competencia_fica_de_fora(self):
        # NE só existe no resultado se tiver ao menos uma linha de origem — ausência de
        # competência não deve virar uma linha com valor 0 implícita.
        nes_no_resultado = set(self.agregado["ne_ccor"])
        nes_na_origem = set(self.df["ne_ccor"])
        self.assertEqual(nes_no_resultado, nes_na_origem)


@unittest.skipUnless(CAMINHO_BASE.exists(), f"Planilha ausente em {CAMINHO_BASE}")
class TestReferenciaContraExtracaoReal(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.hash_curto = _hash_curto(CAMINHO_BASE)
        cls.referencia = REFERENCIAS.get(cls.hash_curto)
        if cls.referencia is not None:
            cls.df = ler_liquidacao_competencia(CAMINHO_BASE)

    def test_hash_conhecido(self):
        if self.referencia is None:
            self.skipTest(
                f"Fixture com hash {self.hash_curto} sem entrada em REFERENCIAS — "
                "acrescente uma nova entrada (não edite as existentes)."
            )

    def test_totais_batem_com_a_referencia(self):
        if self.referencia is None:
            self.skipTest("sem referência para este hash — ver test_hash_conhecido")
        resumo = reconciliar(self.df)
        self.assertEqual(resumo["linhas"], self.referencia["linhas"])
        self.assertEqual(resumo["documentos_habeis_distintos"], self.referencia["documentos_habeis_distintos"])
        self.assertEqual(resumo["notas_empenho_distintas"], self.referencia["notas_empenho_distintas"])
        self.assertEqual(resumo["meses_referencia"], self.referencia["meses_referencia"])
        self.assertAlmostEqual(resumo["total"], self.referencia["total"], places=2)
        self.assertAlmostEqual(resumo["total_estornos"], self.referencia["total_estornos"], places=2)

    def test_validacao_estrutural_aprova(self):
        if self.referencia is None:
            self.skipTest("sem referência para este hash — ver test_hash_conhecido")
        relatorio = validar(self.df)
        self.assertTrue(relatorio.ok, msg=f"erros: {relatorio.erros}")
        self.assertEqual(relatorio.alertas, [])


if __name__ == "__main__":
    unittest.main()


# ---------------------------------------------------------------------------------------
# Layout com dimensões (relatório "Liquidação por competência - JVSB", 13/10/2026): 12
# colunas — Ação, PTRES e Credor antes das 7 do layout antigo. Planilha sintética, montada
# em memória (a extração real tem ~24 mil linhas; os valores abaixo são calculados à mão).
# ---------------------------------------------------------------------------------------

_CABECALHO_COM_DIMENSOES = [
    "Ação Governo", None, "PTRES", "DH - Credor", None,
    "NE", "Documento Hábil", "DH - Doc. Contábil",
    "Ano Referência ACC", "Mês Referência ACC", "Métrica", None,
]
_NE_A = "153165152392022NE000484"
_NE_B = "153165152392023NE000133"
_LINHAS_COM_DIMENSOES = [
    ["00P1", "CONCESSAO DE BOLSAS", "169150", "153165", "UFRPE", _NE_A,
     "153165152392023FL000011", "153165152392023NS000231", 2022, "DEZ/2022", "DetaCusto DH - R$", 35312.37],
    ["00P1", "CONCESSAO DE BOLSAS", "169150", "153165", "UFRPE", _NE_B,
     "153165152392023FL000013", "'-8", 2023, "JAN/2023", "DetaCusto DH - R$", 100.00],
    # estorno: sinal preservado
    ["00P1", "CONCESSAO DE BOLSAS", "169150", "153165", "UFRPE", _NE_B,
     "153165152392023FL000025", "'-8", 2023, "FEV/2023", "DetaCusto DH - R$", -40.50],
]


def _gravar_com_dimensoes(destino: Path, linhas=None, cabecalho=None) -> Path:
    wb = openpyxl.Workbook()
    aba = wb.active
    aba.append(cabecalho or _CABECALHO_COM_DIMENSOES)
    for linha in (_LINHAS_COM_DIMENSOES if linhas is None else linhas):
        aba.append(linha)
    wb.save(destino)
    return destino


class TestLayoutComDimensoes(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.caminho = _gravar_com_dimensoes(Path(self._tmp.name) / "liq.xlsx")

    def test_le_o_layout_com_dimensoes(self):
        df = ler_liquidacao_competencia(self.caminho)
        self.assertEqual(len(df), 3)
        self.assertEqual(round(float(df["valor"].sum()), 2), 35371.87)
        self.assertEqual(round(float(df.loc[df["valor"] < 0, "valor"].sum()), 2), -40.50)
        self.assertEqual(sorted(df["ano_mes_referencia"]), [202212, 202301, 202302])

    def test_dimensoes_novas_preservadas_como_identificador(self):
        df = ler_liquidacao_competencia(self.caminho)
        self.assertEqual(set(df["acao_cod"]), {"00P1"})  # alfanumérico, zeros à esquerda
        self.assertEqual(set(df["ptres"]), {"169150"})
        self.assertEqual(set(df["credor_cod"]), {"153165"})
        self.assertEqual(set(df["credor_desc"]), {"UFRPE"})

    def test_colunas_originais_continuam_as_mesmas(self):
        df = ler_liquidacao_competencia(self.caminho)
        for coluna in ("ne_ccor", "documento_habil", "doc_contabil", "ano_referencia", "valor",
                       "mes_referencia", "ano_mes_referencia", "linha_origem", "arquivo_origem"):
            self.assertIn(coluna, df.columns)
        # mesmos consumidores do layout antigo funcionam sem alteração
        serie = liquidado_por_ne_e_mes(df)
        self.assertEqual(len(serie), 3)

    def test_validacao_aprova(self):
        self.assertTrue(validar(ler_liquidacao_competencia(self.caminho)).ok)

    def test_cabecalho_com_dimensoes_alterado_e_rejeitado(self):
        cab = list(_CABECALHO_COM_DIMENSOES)
        cab[2] = "Outra Coisa"  # não mais "PTRES"
        destino = _gravar_com_dimensoes(Path(self._tmp.name) / "ruim.xlsx", cabecalho=cab)
        with self.assertRaises(ErroLayoutBase):
            ler_liquidacao_competencia(destino)

    def test_numero_de_colunas_diferente_e_rejeitado(self):
        cab = _CABECALHO_COM_DIMENSOES + ["Extra"]
        linhas = [linha + [1] for linha in _LINHAS_COM_DIMENSOES]
        destino = _gravar_com_dimensoes(Path(self._tmp.name) / "larga.xlsx", linhas=linhas, cabecalho=cab)
        with self.assertRaises(ErroLayoutBase):
            ler_liquidacao_competencia(destino)


@unittest.skipUnless(CAMINHO_BASE.exists(), f"Planilha ausente em {CAMINHO_BASE}")
class TestLayoutAntigoTemDimensoesNulas(unittest.TestCase):
    def test_dimensoes_novas_nulas_no_layout_antigo(self):
        df = ler_liquidacao_competencia(CAMINHO_BASE)
        for coluna in ("acao_cod", "acao_desc", "ptres", "credor_cod", "credor_desc"):
            self.assertIn(coluna, df.columns)
            self.assertTrue(df[coluna].isna().all())


class TestDimensoesPorNe(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)

    def test_lista_todas_as_acoes_e_ptres_da_ne_sem_escolher(self):
        linhas = [list(l) for l in _LINHAS_COM_DIMENSOES]
        linhas[2][0], linhas[2][2] = "4002", "169999"  # mesma NE_B, outra ação/PTRES
        df = ler_liquidacao_competencia(_gravar_com_dimensoes(Path(self._tmp.name) / "d.xlsx", linhas=linhas))
        dims = dimensoes_por_ne(df)
        self.assertEqual(dims.loc[_NE_A, "acoes"], "00P1")
        self.assertEqual(dims.loc[_NE_B, "acoes"], "00P1 / 4002")
        self.assertEqual(dims.loc[_NE_B, "ptres"], "169150 / 169999")

    @unittest.skipUnless(CAMINHO_BASE.exists(), f"Planilha ausente em {CAMINHO_BASE}")
    def test_layout_antigo_devolve_nulo(self):
        dims = dimensoes_por_ne(ler_liquidacao_competencia(CAMINHO_BASE))
        self.assertTrue(dims["acoes"].isna().all())
        self.assertTrue(dims["ptres"].isna().all())
